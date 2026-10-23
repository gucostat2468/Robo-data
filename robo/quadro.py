"""Quadro de pessoal: todos os andares lotados, do gerente ao estagiário, cada um com trabalho REAL.

Regra da casa: ninguém é decorativo. Cada cargo abaixo tem uma conferência própria, executada em todo
arquivo, com prova técnica no RH (um caso bom que deve passar e um caso ruim que ela precisa pegar).

Escopo por senioridade (vale para todos os andares):
  Gerente      responde pelo resultado do andar: metas, prazos (SLA), prioridades e a manchete para o dono.
               Não executa tarefa operacional e não aprova o próprio trabalho.
  Coordenador  portão de qualidade: só libera com o checklist completo; pode BLOQUEAR.
  Especialista/
  Sênior       executa as etapas críticas, define o método, revisa plenos e juniores, mentora.
  Pleno        segunda conferência independente, de média complexidade. Achado dele é invariante:
               vai para o coordenador e BLOQUEIA a produção (no 1º andar) ou vira ressalva (2º e 3º).
  Júnior       checagens padronizadas de baixo risco, com lista fechada de critérios. Também bloqueia
               no 1º andar, porque o que ele confere são regras objetivas (não opinião).
  Estagiário   MODO OBSERVAÇÃO: refaz uma amostra e relata. Nunca altera nada e nunca bloqueia.

Fases em que cada conferência roda:
  engenharia   antes do checklist da coordenadora do 1º andar (pode bloquear a publicação)
  publicacao   logo depois de gravar os arquivos (releitura do que foi gravado)
  analise      depois dos analistas do 2º andar (ressalvas da análise)
  excel        depois da entrega da Sala Excel (ressalvas da entrega)
  governanca   depois de tudo (dados pessoais e catálogo de dados)
"""

from __future__ import annotations

import json
import random
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import pandas as pd

from robo.limpeza import normalizar_nome_coluna

ESCOPO_NIVEL = {
    "head": {"nome": "Head", "decide": "a arquitetura e o método da área (custos, segurança, plano de volta)",
             "bloqueia": True, "executa": "revisa o desenho técnico e as verificações da área",
             "nao_faz": "não aprova a própria mudança nem altera permissões sem autorização do dono"},
    "gerente": {"nome": "Gerente", "decide": "metas, prioridades e prazos do andar", "bloqueia": False,
                "executa": "acompanha indicadores do andar e escreve a manchete para o dono",
                "nao_faz": "não executa tarefa operacional nem aprova o próprio trabalho"},
    "coordenador": {"nome": "Coordenador(a)", "decide": "se o trabalho entra em produção", "bloqueia": True,
                    "executa": "checklist completo de produção", "nao_faz": "não refaz o trabalho da equipe; devolve com o motivo"},
    "especialista": {"nome": "Especialista", "decide": "o método técnico da área", "bloqueia": True,
                     "executa": "etapas críticas e casos difíceis", "nao_faz": "não aprova a própria entrega"},
    "senior": {"nome": "Sênior", "decide": "como executar a etapa", "bloqueia": True,
               "executa": "etapas críticas; revisa plenos e juniores", "nao_faz": "não aprova a própria entrega"},
    "pleno": {"nome": "Pleno", "decide": "se a sua conferência passou", "bloqueia": True,
              "executa": "segunda conferência independente (média complexidade)", "nao_faz": "não muda regra nem dado"},
    "junior": {"nome": "Júnior", "decide": "nada fora da lista de critérios", "bloqueia": True,
               "executa": "checagens padronizadas de baixo risco", "nao_faz": "não interpreta nem muda dado"},
    "estagiario": {"nome": "Estagiário", "decide": "nada (aprendizado)", "bloqueia": False,
                   "executa": "refaz uma amostra em modo observação e relata", "nao_faz": "não altera nem bloqueia nada"},
}

AUDITORIA = {"_arquivo_origem", "_processado_em", "_aba_origem", "_linha_origem"}


@dataclass
class Contexto:
    arquivo: str
    cfg: object = None
    regra: dict = field(default_factory=dict)
    brutas: dict = field(default_factory=dict)
    analisaveis: dict = field(default_factory=dict)
    tabelas: dict = field(default_factory=dict)
    saidas: list = field(default_factory=list)
    analise: dict | None = None
    excel: dict | None = None
    duracao: float = 0.0


# =============================================================================== 1º andar · Engenharia
def conf_tipos(c: Contexto) -> list[str]:
    """Júlia (pleno): cada coluna com tipo definido na regra saiu com esse tipo de verdade."""
    problemas = []
    tipos = {normalizar_nome_coluna(k): str(v).lower() for k, v in (c.regra.get("tipos") or {}).items()}
    for nome, df in c.analisaveis.items():
        for col, tipo in tipos.items():
            if col not in df.columns:
                continue
            s = df[col]
            ok = {"inteiro": pd.api.types.is_integer_dtype(s) or s.dropna().map(lambda x: float(x).is_integer()).all()
                  if pd.api.types.is_numeric_dtype(s) else False,
                  "decimal": pd.api.types.is_numeric_dtype(s), "numero": pd.api.types.is_numeric_dtype(s),
                  "data": pd.api.types.is_datetime64_any_dtype(s),
                  "texto": not pd.api.types.is_numeric_dtype(s) or s.isna().all()}.get(tipo, True)
            if not ok:
                problemas.append(f"[{nome}.{col}] deveria ser {tipo}, saiu como {s.dtype}")
    return problemas


def conf_espacos(c: Contexto) -> list[str]:
    """Bruno (júnior): nenhum texto com espaço sobrando nas pontas e nenhuma linha 100% vazia."""
    problemas = []
    for nome, df in c.analisaveis.items():
        dados = df[[x for x in df.columns if x not in AUDITORIA]]
        vazias = int(dados.isna().all(axis=1).sum()) if len(dados.columns) else 0
        if vazias:
            problemas.append(f"[{nome}] {vazias} linha(s) totalmente vazia(s) publicada(s)")
        for col in dados.columns:
            s = dados[col]
            if s.dtype == object or pd.api.types.is_string_dtype(s):
                txt = s.dropna().astype(str)
                sujos = int((txt != txt.str.strip()).sum())
                if sujos:
                    problemas.append(f"[{nome}.{col}] {sujos} texto(s) com espaço sobrando nas pontas")
    return problemas


def conf_auditoria(c: Contexto) -> list[str]:
    """Camila (júnior): toda linha publicada sabe de onde veio e quando foi processada."""
    if c.cfg is not None and not c.cfg.saida.get("colunas_auditoria", True):
        return []
    problemas = []
    for nome, df in c.tabelas.items():
        if nome.startswith("dim_"):
            continue
        if "_arquivo_origem" in df.columns and (df["_arquivo_origem"].astype(str) != c.arquivo).any():
            problemas.append(f"[{nome}] _arquivo_origem diferente de {c.arquivo}")
        if "_processado_em" in df.columns and df["_processado_em"].isna().any():
            problemas.append(f"[{nome}] _processado_em vazio em {int(df['_processado_em'].isna().sum())} linha(s)")
    return problemas


def conf_amostra(c: Contexto) -> list[str]:
    """Théo (estagiário, observação): sorteia até 5 chaves publicadas e confirma que existem na planilha original."""
    chaves = [normalizar_nome_coluna(k) for k in (c.regra.get("chave_unica") or [])]
    if not chaves:
        return []
    obs = []
    rnd = random.Random(c.arquivo)
    for aba, bruto in c.brutas.items():
        b = bruto.rename(columns=lambda x: normalizar_nome_coluna(str(x)))
        for nome, limpo in c.analisaveis.items():
            if not all(k in limpo.columns and k in b.columns for k in chaves) or limpo.empty:
                continue
            origem = set(map(tuple, b[chaves].astype(str).apply(lambda s: s.str.strip()).values.tolist()))
            amostra = limpo[chaves].dropna()
            amostra = amostra.iloc[rnd.sample(range(len(amostra)), min(5, len(amostra)))]
            for linha in amostra.astype(str).values.tolist():
                valor = tuple(v.strip().removesuffix(".0") for v in linha)
                if valor not in origem and tuple(linha) not in origem:
                    obs.append(f"chave {', '.join(valor)} não encontrada como texto igual na aba {aba} (pode ser só formatação)")
    return obs


def conf_releitura(c: Contexto) -> list[str]:
    """Rafael (pleno): relê cada arquivo gravado e confere linhas e colunas com o que deveria ter sido gravado."""
    problemas = []
    for caminho in c.saidas:
        p = Path(caminho)
        if not p.exists():
            problemas.append(f"{p.name} não existe depois de gravado")
            continue
        if p.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(p, read_only=True)
            try:
                for nome, df in c.tabelas.items():
                    aba = next((ws for ws in wb.worksheets if normalizar_nome_coluna(ws.title) == normalizar_nome_coluna(nome[:31])), None)
                    if aba is None:
                        continue
                    linhas = sum(1 for _ in aba.iter_rows(values_only=True)) - 1
                    if linhas != len(df):
                        problemas.append(f"{p.name}/{aba.title}: {linhas} linha(s) gravada(s), esperado {len(df)}")
            finally:
                wb.close()
        elif p.suffix.lower() == ".csv":
            nome = next((n for n in c.tabelas if p.stem.endswith(n) or n in p.stem), None)
            if nome is None and len(c.tabelas) == 1:
                nome = next(iter(c.tabelas))
            if nome is not None:
                lido = pd.read_csv(p, sep=None, engine="python", encoding="utf-8-sig", dtype=str)
                if len(lido) != len(c.tabelas[nome]):
                    problemas.append(f"{p.name}: {len(lido)} linha(s) gravada(s), esperado {len(c.tabelas[nome])}")
    return problemas


def conf_sla(c: Contexto) -> list[str]:
    """Marcelo (gerente): prazo do andar. Não bloqueia; vira aviso e entra na sprint."""
    limite = float(((c.cfg.dados.get("quadro") or {}) if c.cfg is not None else {}).get("sla_segundos", 120))
    return [f"processamento levou {c.duracao:.0f}s (meta do andar: até {limite:.0f}s)"] if c.duracao > limite else []


# =============================================================================== 2º andar · Análise
def _df_da_coluna(c: Contexto, coluna: str) -> pd.Series | None:
    for df in c.analisaveis.values():
        if coluna in df.columns:
            return df[coluna]
    return None


def conf_recalculo(c: Contexto) -> list[str]:
    """Sérgio (pleno): recalcula o NPS geral por contagem simples, sem reaproveitar o código do Caio."""
    problemas = []
    for i in (c.analise or {}).get("_indicadores") or []:
        if i.get("indicador") != "NPS" or i.get("segmento") != "Total":
            continue
        s = _df_da_coluna(c, i.get("coluna"))
        if s is None:
            continue
        notas = [float(x) for x in s if pd.notna(x)]
        if not notas:
            continue
        prom = sum(1 for x in notas if x >= 9)
        det = sum(1 for x in notas if x <= 6)
        esperado = (prom - det) * 100.0 / len(notas)
        if abs(float(i.get("valor")) - esperado) > 0.051 or int(i.get("n")) != len(notas):
            problemas.append(f"NPS de {i.get('coluna')}: relatório {i.get('valor')} (n={i.get('n')}), "
                             f"recálculo {esperado:.1f} (n={len(notas)})")
    return problemas


def conf_base_minima(c: Contexto) -> list[str]:
    """Larissa (pleno): todo número com base abaixo do mínimo está sinalizado como base pequena."""
    minimo = int(((c.cfg.dados.get("analise") or {}) if c.cfg is not None else {}).get("base_minima", 30))
    erradas = [i for i in (c.analise or {}).get("_indicadores") or []
               if int(i.get("n") or 0) < minimo and not i.get("base_pequena")]
    return [f"{len(erradas)} indicador(es) com base < {minimo} sem o aviso de base pequena"] if erradas else []


def conf_faixas(c: Contexto) -> list[str]:
    """Igor (júnior): NPS entre -100 e 100, percentuais entre 0 e 100, margem não negativa."""
    problemas = []
    for i in (c.analise or {}).get("_indicadores") or []:
        v = float(i.get("valor") or 0)
        faixa = (-100, 100) if i.get("indicador") == "NPS" else (0, 100)
        if not faixa[0] <= v <= faixa[1]:
            problemas.append(f"{i.get('indicador')} {i.get('segmento')}: valor {v} fora de {faixa}")
        if float(i.get("margem") or 0) < 0:
            problemas.append(f"{i.get('indicador')} {i.get('segmento')}: margem negativa")
        for k in ("pct_promotores", "pct_neutros", "pct_detratores"):
            if i.get(k) is not None and not 0 <= float(i[k]) <= 100:
                problemas.append(f"{i.get('indicador')} {i.get('segmento')}: {k} fora de 0–100")
    return problemas


def conf_recontagem(c: Contexto) -> list[str]:
    """Manu (estagiária, observação): reconta a base (n) de cada indicador geral direto na tabela."""
    obs = []
    for i in (c.analise or {}).get("_indicadores") or []:
        if i.get("segmento") != "Total":
            continue
        s = _df_da_coluna(c, i.get("coluna"))
        if s is not None and int(s.notna().sum()) != int(i.get("n") or 0):
            obs.append(f"{i.get('indicador')}: contei {int(s.notna().sum())} respostas, o relatório diz {i.get('n')}")
    return obs


def manchete(analise: dict | None) -> str:
    """Patrícia (gerente): uma frase para o dono — o que mudou e onde olhar primeiro."""
    if not analise:
        return ""
    ind = [i for i in analise.get("_indicadores") or [] if i.get("segmento") == "Total"]
    dest = [d for d in analise.get("_destaques") or [] if d.get("significativo")]
    partes = []
    if ind:
        i = ind[0]
        partes.append(f"{i['indicador']} geral {i['valor']} (±{i.get('margem')}, n={i.get('n')})")
    if dest:
        d = dest[0]
        partes.append(f"diferença real em {d.get('segmento_coluna') or d.get('coluna') or 'um segmento'}"
                      + (f" ({d.get('grupo')})" if d.get("grupo") else ""))
    return "; ".join(partes) if partes else "base sem indicador reconhecido — só perfil e qualidade"


# =============================================================================== 3º andar · Sala Excel
def _abrir_entrega(c: Contexto):
    if not c.excel or not Path(c.excel.get("arquivo", "")).exists():
        return None
    from openpyxl import load_workbook
    return load_workbook(c.excel["arquivo"])


def conf_integridade_excel(c: Contexto) -> list[str]:
    """Sílvia (coordenadora): abas obrigatórias, nenhuma fórmula com #REF! e nenhum vínculo externo."""
    wb = _abrir_entrega(c)
    if wb is None:
        return []
    problemas = [f"aba obrigatória ausente: {a}" for a in ("Resumo", "Dados", "Conferência") if a not in wb.sheetnames]
    if getattr(wb, "_external_links", None):
        problemas.append("a entrega aponta para outro arquivo (vínculo externo)")
    for ws in wb.worksheets:
        if ws.title == "Dados":
            continue
        for linha in ws.iter_rows():
            for cel in linha:
                if isinstance(cel.value, str) and cel.value.startswith("="):
                    if "#REF!" in cel.value or re.search(r"\[\d+\]", cel.value):
                        problemas.append(f"{ws.title}!{cel.coordinate}: fórmula quebrada ou externa")
    return problemas


def conf_formulas_vivas(c: Contexto) -> list[str]:
    """Hugo (pleno): na Conferência, valor da fórmula e status são fórmulas vivas; o esperado é número."""
    wb = _abrir_entrega(c)
    if wb is None or "Conferência" not in wb.sheetnames:
        return []
    ws = wb["Conferência"]
    problemas = []
    for r in range(5, ws.max_row + 1):
        if ws.cell(row=r, column=1).value is None:
            continue
        for col in (3, 5):
            v = ws.cell(row=r, column=col).value
            if not (isinstance(v, str) and v.startswith("=")):
                problemas.append(f"Conferência!{ws.cell(row=r, column=col).coordinate} não é fórmula (valor digitado)")
        if not isinstance(ws.cell(row=r, column=4).value, (int, float)):
            problemas.append(f"Conferência!D{r}: valor esperado não é número")
    return problemas


def conf_formatos(c: Contexto) -> list[str]:
    """Paula (júnior): todo número em fórmula tem formato definido (milhar, %, decimal) e as abas grandes têm cabeçalho congelado."""
    wb = _abrir_entrega(c)
    if wb is None:
        return []
    problemas = []
    colunas = {"Plano tabular": range(2, 7), "Indicadores": range(2, 3), "Trade": range(2, 4)}
    for aba, cols in colunas.items():
        if aba not in wb.sheetnames:
            continue
        ws = wb[aba]
        sem = [ws.cell(row=r, column=col).coordinate for r in range(1, ws.max_row + 1) for col in cols
               if isinstance(ws.cell(row=r, column=col).value, str) and ws.cell(row=r, column=col).value.startswith("=")
               and ws.cell(row=r, column=col).number_format == "General"]
        if sem:
            problemas.append(f"{aba}: {len(sem)} número(s) sem formato (ex.: {', '.join(sem[:3])})")
    for aba in ("Dados", "Conferência"):
        if aba in wb.sheetnames and not wb[aba].freeze_panes:
            problemas.append(f"{aba}: cabeçalho não está congelado")
    return problemas


def conf_amostra_conferencias(c: Contexto) -> list[str]:
    """Enzo (estagiário, observação): sorteia 3 conferências e confirma que a célula conferida existe e é fórmula."""
    wb = _abrir_entrega(c)
    if wb is None or "Conferência" not in wb.sheetnames:
        return []
    ws = wb["Conferência"]
    linhas = [r for r in range(5, ws.max_row + 1) if ws.cell(row=r, column=3).value]
    obs = []
    for r in random.Random(c.arquivo).sample(linhas, min(3, len(linhas))):
        m = re.search(r"'([^']+)'!([A-Z]+\d+)", str(ws.cell(row=r, column=3).value))
        if not m or m.group(1) not in wb.sheetnames:
            obs.append(f"Conferência linha {r}: não achei a aba referenciada")
            continue
        alvo = wb[m.group(1)][m.group(2)].value
        if not (isinstance(alvo, str) and alvo.startswith("=")):
            obs.append(f"{m.group(1)}!{m.group(2)} conferida mas não é fórmula")
    return obs


def conf_compatibilidade_excel(c: Contexto) -> list[str]:
    """Ricardo (gerente da Sala Excel, coordena a limpeza): barra o que o Excel estragaria ao abrir a base."""
    problemas = []
    for nome, df in c.tabelas.items():
        if len(df) > 1_048_575:
            problemas.append(f"[{nome}] {len(df):,} linhas: passa do limite de uma aba do Excel (1.048.576)".replace(",", "."))
        if len(df.columns) > 16_384:
            problemas.append(f"[{nome}] {len(df.columns)} colunas: passa do limite do Excel (16.384)")
        for col in df.columns:
            s = df[col]
            if pd.api.types.is_integer_dtype(s) or pd.api.types.is_float_dtype(s):
                grandes = s.dropna().abs() >= 10 ** 15
                if grandes.any():
                    problemas.append(f"[{nome}.{col}] {int(grandes.sum())} número(s) com 16+ dígitos: o Excel arredonda "
                                     f"(ex.: {s.dropna()[grandes].iloc[0]:.0f}). Se é código/ID, deve ser texto")
            elif pd.api.types.is_datetime64_any_dtype(s):
                antigas = s.dropna() < pd.Timestamp("1900-01-01")
                if antigas.any():
                    problemas.append(f"[{nome}.{col}] {int(antigas.sum())} data(s) antes de 1900: o Excel não mostra")
            else:
                longos = s.dropna().astype(str).str.len() > 32_767
                if longos.any():
                    problemas.append(f"[{nome}.{col}] {int(longos.sum())} texto(s) maiores que 32.767 caracteres: "
                                     f"o Excel corta")
    return problemas


# =============================================================================== 4º andar · Governança e canal
RE_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$", re.IGNORECASE)
RE_CPF = re.compile(r"^\d{3}\.?\d{3}\.?\d{3}-?\d{2}$")
RE_TEL = re.compile(r"^(\+?55\s?)?\(?\d{2}\)?\s?9?\d{4}[-\s]?\d{4}$")  # com ou sem +55 (achado do treino)
RE_EMAIL_NO_TEXTO = re.compile(r"[^@\s]+@[^@\s]+\.[a-z]{2,}", re.IGNORECASE)


def cpf_valido(texto: str) -> bool:
    """CPF de verdade passa no dígito verificador; número de pedido com a mesma cara não passa (achado do treino)."""
    d = re.sub(r"\D", "", str(texto))
    if len(d) != 11 or len(set(d)) == 1:
        return False
    n = [int(x) for x in d]
    for k in (9, 10):
        v = (sum(a * b for a, b in zip(n[:k], range(k + 1, 1, -1))) * 10) % 11
        if (0 if v == 10 else v) != n[k]:
            return False
    return True


def conf_dados_pessoais(c: Contexto) -> list[str]:
    """Vitor (pleno, LGPD): aponta colunas publicadas que parecem dado pessoal sem anonimização."""
    achados = []
    for nome, df in c.analisaveis.items():
        for col in df.columns:
            if col in AUDITORIA:
                continue
            amostra = df[col].dropna().astype(str).str.strip().head(200)
            if amostra.empty:
                continue
            testes = (("e-mail", lambda s: s.str.match(RE_EMAIL)),
                      ("CPF", lambda s: s.str.match(RE_CPF) & s.map(cpf_valido)),
                      ("telefone", lambda s: s.str.match(RE_TEL)),
                      ("e-mail dentro do texto", lambda s: s.str.contains(RE_EMAIL_NO_TEXTO)))
            for tipo, teste in testes:
                if teste(amostra).mean() >= 0.6:
                    achados.append(f"[{nome}.{col}] parece {tipo} sem anonimização — confirme se pode ser publicado "
                                   f"(habilidade 'anonimizar' resolve)")
                    break
    return achados


def registrar_catalogo(c: Contexto) -> list[str]:
    """Lara (estagiária): mantém o catálogo de dados (logs/catalogo_dados.json) — o que existe e com que tipo."""
    if c.cfg is None:
        return []
    arq = Path(c.cfg.logs) / "catalogo_dados.json"
    try:
        catalogo = json.loads(arq.read_text(encoding="utf-8")) if arq.exists() else {}
    except (OSError, json.JSONDecodeError):
        catalogo = {}
    from robo.analise import chave_base
    base = chave_base(c.arquivo, c.regra or {})
    for nome, df in c.tabelas.items():
        catalogo[f"{base}.{nome}"] = {"base": base, "tabela": nome, "ultimo_arquivo": c.arquivo, "atualizado_em": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                          "linhas": int(len(df)), "colunas": {col: str(df[col].dtype) for col in df.columns}}
    arq.parent.mkdir(parents=True, exist_ok=True)
    tmp = arq.with_suffix(".tmp")
    tmp.write_text(json.dumps(catalogo, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(arq)
    return []


# =============================================================================== o quadro
def _p(id, nome, cargo, nivel, andar, area, gestor, cor, fase, funcao, escopo, criterios, metodologias, nao_faz=None):
    return {"id": id, "nome": nome, "cargo": cargo, "papel": cargo.split(" — ")[0], "nivel": nivel, "andar": andar,
            "area": area, "gestor": gestor, "cor": cor, "fase": fase, "funcao": funcao, "escopo": escopo,
            "criterios": criterios, "metodologias": metodologias,
            "nao_faz": [_frase(x) for x in (nao_faz or [ESCOPO_NIVEL[nivel]["nao_faz"]])],
            "descricao": _frase((funcao.__doc__ or "").split(":", 1)[-1]) if funcao else ""}


def _frase(texto: str) -> str:
    texto = " ".join(texto.split())
    return texto[:1].upper() + texto[1:]


QUADRO = [
    # ---- 1º andar · Engenharia de dados
    _p("ger_engenharia", "Marcelo", "Gerente de Engenharia de Dados", "gerente", 1, "Engenharia de Dados", "dono", "#0F172A",
       "fim", conf_sla, ["Metas do andar: taxa de aprovação, prazo (SLA) e retrabalho", "Prioriza o backlog da engenharia",
                         "Conta na sprint o que travou e por quê"],
       ["SLA por arquivo dentro da meta", "Nenhum arquivo esquecido na fila"], ["Gestão por indicadores (OKR/KPI)", "SLA e capacidade"]),
    _p("eng_pleno_tipos", "Júlia", "Engenheira de Dados Pleno — Conferência de tipos", "pleno", 1, "Engenharia de Dados",
       "coord_engenharia", "#2563EB", "engenharia", conf_tipos,
       ["Confere que cada coluna com tipo na regra saiu com esse tipo", "Devolve para a coordenadora com o exemplo do erro"],
       ["Coluna 'inteiro' sem decimais", "Coluna 'data' em formato de data de verdade"], ["Contratos de dados (data contracts)", "Tipagem forte"]),
    _p("eng_pleno_releitura", "Rafael", "Engenheiro de Dados Pleno — Releitura do que foi gravado", "pleno", 1,
       "Engenharia de Dados", "coord_engenharia", "#1D4ED8", "publicacao", conf_releitura,
       ["Relê cada arquivo publicado e confere a contagem de linhas", "Se o disco gravou errado, o arquivo volta para Bases-Erros"],
       ["Linhas gravadas = linhas aprovadas"], ["Verificação pós-escrita (read-after-write)"]),
    _p("eng_junior_espacos", "Bruno", "Analista de Dados Júnior — Higiene de texto", "junior", 1, "Qualidade de Dados",
       "eng_pleno_tipos", "#3B82F6", "engenharia", conf_espacos,
       ["Confere: nenhum texto com espaço sobrando nas pontas", "Confere: nenhuma linha 100% vazia publicada"],
       ["Zero espaços nas pontas", "Zero linhas vazias"], ["Checklist de higiene de dados"]),
    _p("eng_junior_auditoria", "Camila", "Analista de Dados Júnior — Rastreabilidade", "junior", 1, "Governança e Privacidade",
       "eng_pleno_tipos", "#60A5FA", "engenharia", conf_auditoria,
       ["Confere que toda linha publicada tem arquivo de origem e data de processamento"],
       ["_arquivo_origem correto em 100% das linhas"], ["Linhagem de dados (data lineage)"]),
    _p("eng_estagiario", "Théo", "Estagiário de Engenharia de Dados", "estagiario", 1, "Engenharia de Dados",
       "eng_pleno_releitura", "#93C5FD", "engenharia", conf_amostra,
       ["Sorteia até 5 chaves publicadas e procura cada uma na planilha original", "Relata o que achou; não bloqueia"],
       ["Amostra rastreada até a origem"], ["Amostragem aleatória reprodutível"]),
    # ---- 2º andar · Análise e qualidade
    _p("ger_analise", "Patrícia", "Gerente de Analytics e Customer Insights", "gerente", 2, "Customer Insights e Pesquisa", "dono",
       "#4C1D95", "fim", None, ["Escreve a manchete de cada base para o dono", "Prioriza quais análises viram ação",
                                "Conta na sprint o que as análises mostraram"],
       ["Manchete com número, base e margem"], ["Storytelling executivo (pirâmide de Minto)", "Priorização por impacto"]),
    _p("ana_pleno_recalculo", "Sérgio", "Analista de Dados Pleno — Recálculo independente", "pleno", 2, "Estatística",
       "coord_analise", "#7C3AED", "analise", conf_recalculo,
       ["Recalcula o NPS geral por contagem simples, sem usar o código do Caio", "Diverge → ressalva no checklist"],
       ["Diferença ≤ 0,05 ponto e mesma base (n)"], ["Dupla checagem independente (four-eyes)"]),
    _p("ana_pleno_base", "Larissa", "Analista de Customer Insights Pleno — Base mínima", "pleno", 2, "Customer Insights e Pesquisa",
       "coord_analise", "#8B5CF6", "analise", conf_base_minima,
       ["Confere que todo número com base abaixo do mínimo está sinalizado"], ["100% das bases pequenas sinalizadas"],
       ["Leitura responsável de pesquisa (base mínima 30)"]),
    _p("ana_junior_faixas", "Igor", "Analista de Dados Júnior — Faixas válidas", "junior", 2, "Análise de Dados",
       "ana_pleno_recalculo", "#A78BFA", "analise", conf_faixas,
       ["Confere NPS entre −100 e 100, percentuais entre 0 e 100 e margem não negativa"], ["Zero valores impossíveis"],
       ["Checagem de domínio (range checks)"]),
    _p("ana_estagiario", "Manu", "Estagiária de Análise de Dados", "estagiario", 2, "Análise de Dados", "ana_pleno_base",
       "#C4B5FD", "analise", conf_recontagem, ["Reconta a base de cada indicador direto na tabela e relata"],
       ["Base (n) recontada"], ["Recontagem manual"]),
    # ---- 3º andar · Sala Excel
    _p("coord_excel", "Sílvia", "Coordenadora da Sala Excel — Integridade da entrega", "coordenador", 3, "Engenharia de Dados",
       "excel_supervisor", "#065F46", "excel", conf_integridade_excel,
       ["Confere abas obrigatórias, fórmulas sem #REF! e nenhum vínculo para outro arquivo", "Só então a entrega vale"],
       ["Zero fórmulas quebradas", "Zero vínculos externos"], ["Auditoria de planilhas (ICAEW/FAST)"]),
    _p("xl_pleno_formulas", "Hugo", "Analista de Excel Pleno — Fórmulas vivas", "pleno", 3, "Engenharia de Dados", "coord_excel",
       "#15803D", "excel", conf_formulas_vivas,
       ["Confere que nenhum número da Conferência foi digitado por cima da fórmula"], ["100% das conferências em fórmula"],
       ["Rastreabilidade célula a célula"]),
    _p("xl_junior_formatos", "Paula", "Analista de Excel Júnior — Formatos e leitura", "junior", 3, "Storytelling e Negócio",
       "xl_pleno_formulas", "#22C55E", "excel", conf_formatos,
       ["Confere que todo número tem formato (milhar, %, decimal)", "Confere cabeçalhos congelados nas abas grandes"],
       ["Zero números em formato Geral"], ["Padrões de formatação para leitura executiva"]),
    _p("xl_estagiario", "Enzo", "Estagiário da Sala Excel", "estagiario", 3, "Engenharia de Dados", "xl_pleno_formulas",
       "#86EFAC", "excel", conf_amostra_conferencias,
       ["Sorteia 3 conferências e confirma que a célula conferida existe e é fórmula"], ["Amostra conferida"],
       ["Amostragem aleatória reprodutível"]),
    # ---- 4º andar · Diretoria, governança e canal
    _p("ger_operacoes", "Renata", "Gerente de Operações e Governança de Dados", "gerente", 4, "Governança e Privacidade", "dono",
       "#7C2D12", "fim", None, ["Publica o boletim diário no canal geral", "Cuida da ordem da casa: prazos, sprints e prioridades",
                                "Responde pelo 4º andar (canal, RH, ideias, LGPD)"],
       ["Boletim diário publicado", "Sprint publicada a cada 15 dias"], ["Gestão de operações", "Rituais ágeis (sprint de 15 dias)"]),
    _p("gov_pleno_lgpd", "Vitor", "Analista de Governança Pleno — LGPD", "pleno", 4, "Governança e Privacidade", "canal_helena",
       "#B45309", "governanca", conf_dados_pessoais,
       ["Aponta colunas publicadas que parecem e-mail, CPF ou telefone", "Recomenda anonimizar quando não houver necessidade"],
       ["Todo dado pessoal publicado é intencional"], ["LGPD: necessidade e minimização"]),
    _p("gov_estagiaria", "Lara", "Estagiária de Documentação de Dados", "estagiario", 4, "Governança e Privacidade",
       "gov_pleno_lgpd", "#F59E0B", "governanca", registrar_catalogo,
       ["Mantém o catálogo de dados: cada tabela, colunas, tipos e último arquivo"], ["Catálogo atualizado a cada publicação"],
       ["Catálogo de dados (data catalog)"]),
]

POR_ID = {p["id"]: p for p in QUADRO}

# quem já existia ganha gestor (organograma completo)
GESTORES_EXISTENTES = {
    "coord_engenharia": "ger_engenharia", "supervisor": "coord_engenharia", "dba": "coord_engenharia",
    "recepcao": "coord_engenharia", "leitura": "coord_engenharia", "limpeza": "coord_engenharia", "validacao": "coord_engenharia",
    "modelagem": "coord_engenharia", "banco": "dba", "publicacao": "coord_engenharia",
    "coord_analise": "ger_analise", "analise_perfil": "coord_analise", "analise_indicadores": "coord_analise",
    "analise_estatistica": "coord_analise", "analise_monitoramento": "coord_analise", "analise_relator": "coord_analise",
    "excel_supervisor": "dono", "excel_pesquisa": "coord_excel", "excel_formulas": "coord_excel", "excel_trade": "coord_excel",
    "excel_dinamicas": "coord_excel",
    "canal_helena": "ger_operacoes", "canal_iara": "ger_operacoes", "canal_nina": "canal_helena",
}


# Conferências feitas por quem já existe em outro módulo (não duplicam no organograma).
# Ricardo é gerente, mas coordena as equipes de limpeza: o que ele barra é limite objetivo do Excel.
CONFERENCIAS_EXTRAS = [
    {"id": "excel_supervisor", "nome": "Ricardo", "cargo": "Gerente da Sala Excel — coordenação da limpeza",
     "nivel": "gerente", "andar": 3, "fase": "engenharia", "funcao": conf_compatibilidade_excel, "bloqueia": True},
]


def lotacao() -> dict:
    """Andar → níveis presentes. Usado para provar que todo andar está completo (registro único)."""
    from robo import organizacao
    return organizacao.lotacao()


def conferir(fase: str, ctx: Contexto) -> list[dict]:
    """Roda as conferências da fase. Falha de código vira achado (nunca derruba o robô)."""
    resultado = []
    for p in QUADRO + CONFERENCIAS_EXTRAS:
        if p["fase"] != fase or p["funcao"] is None:
            continue
        try:
            achados = list(p["funcao"](ctx) or [])
        except Exception as e:  # conferência que quebra não pode virar "aprovado"
            achados = [f"não consegui concluir a conferência ({type(e).__name__}: {e})"]
        resultado.append({"id": p["id"], "nome": p["nome"], "cargo": p["cargo"], "nivel": p["nivel"], "andar": p["andar"],
                          "ok": not achados, "achados": achados[:20],
                          "bloqueia": p.get("bloqueia", ESCOPO_NIVEL[p["nivel"]]["bloqueia"]) and p["nivel"] != "estagiario"})
    return resultado


# =============================================================================== provas técnicas (RH)
def _ctx_prova(**kw) -> Contexto:
    return Contexto(arquivo="prova.xlsx", **kw)


def provas_quadro() -> dict:
    """Cada cargo prova que (1) não acusa nada num caso correto e (2) pega o erro plantado num caso ruim."""
    import tempfile
    bom = pd.DataFrame({"id_resposta": pd.array([1, 2, 3], dtype="Int64"), "nota_nps": pd.array([10, 7, 2], dtype="Int64"),
                        "curso": ["A", "B", "A"], "_arquivo_origem": ["prova.xlsx"] * 3,
                        "_processado_em": [pd.Timestamp("2026-01-01")] * 3})
    regra = {"tipos": {"nota_nps": "inteiro"}, "chave_unica": ["id_resposta"]}
    ruim_tipo = bom.assign(nota_nps=bom["nota_nps"].astype(str))
    ruim_txt = bom.assign(curso=["A ", "B", "A"])
    ruim_aud = bom.assign(_arquivo_origem=["outro.xlsx"] * 3)
    bruto = pd.DataFrame({"ID Resposta": ["1", "2", "3"]})
    bruto_ruim = pd.DataFrame({"ID Resposta": ["7", "8", "9"]})
    c = lambda df, **kw: _ctx_prova(regra=regra, analisaveis={"base": df}, tabelas={"base": df}, **kw)
    r: dict[str, list] = {}
    r["eng_pleno_tipos"] = [("Aceita tipos corretos", not conf_tipos(c(bom))), ("Pega nota em texto", bool(conf_tipos(c(ruim_tipo))))]
    r["eng_junior_espacos"] = [("Aceita texto limpo", not conf_espacos(c(bom))), ("Pega espaço sobrando", bool(conf_espacos(c(ruim_txt))))]
    r["eng_junior_auditoria"] = [("Aceita origem correta", not conf_auditoria(c(bom))),
                                 ("Pega origem trocada", bool(conf_auditoria(c(ruim_aud))))]
    r["eng_estagiario"] = [("Rastreia amostra existente", not conf_amostra(c(bom, brutas={"aba": bruto}))),
                           ("Relata chave que não existe na origem", bool(conf_amostra(c(bom, brutas={"aba": bruto_ruim}))))]
    with tempfile.TemporaryDirectory() as d:
        arq = Path(d) / "saida.xlsx"
        with pd.ExcelWriter(arq) as w:
            bom.to_excel(w, sheet_name="base", index=False)
        r["eng_pleno_releitura"] = [("Confere arquivo íntegro", not conf_releitura(c(bom, saidas=[arq]))),
                                    ("Pega arquivo com linha a menos", bool(conf_releitura(c(pd.concat([bom, bom.head(1)]),
                                                                                             saidas=[arq]))))]
        # análise: indicador certo e indicador adulterado
        ind = {"indicador": "NPS", "coluna": "nota_nps", "segmento": "Total", "n": 3, "valor": 0.0, "margem": 80.0,
               "base_pequena": True, "pct_promotores": 33.3, "pct_neutros": 33.3, "pct_detratores": 33.3}
        ok_a, ruim_a = {"_indicadores": [ind]}, {"_indicadores": [{**ind, "valor": 12.0, "base_pequena": False}]}
        r["ana_pleno_recalculo"] = [("Confirma NPS correto", not conf_recalculo(c(bom, analise=ok_a))),
                                    ("Pega NPS adulterado", bool(conf_recalculo(c(bom, analise=ruim_a))))]
        r["ana_pleno_base"] = [("Aceita base pequena sinalizada", not conf_base_minima(c(bom, analise=ok_a))),
                               ("Pega base pequena sem aviso", bool(conf_base_minima(c(bom, analise=ruim_a))))]
        r["ana_junior_faixas"] = [("Aceita valores possíveis", not conf_faixas(c(bom, analise=ok_a))),
                                  ("Pega NPS impossível", bool(conf_faixas(c(bom, analise={"_indicadores": [{**ind, "valor": 140}]}))))]
        r["ana_estagiario"] = [("Recontagem bate", not conf_recontagem(c(bom, analise=ok_a))),
                               ("Relata base diferente", bool(conf_recontagem(c(bom, analise={"_indicadores": [{**ind, "n": 9}]}))))]
        r["ger_analise"] = [("Manchete com número, margem e base", "n=3" in manchete(ok_a) and "±" in manchete(ok_a)),
                            ("Sem indicador, diz isso com honestidade", "sem indicador" in manchete({"_indicadores": []}))]
        # Sala Excel: entrega real e entrega sabotada
        from openpyxl import load_workbook
        from robo.excel import Entrega
        df_x = bom.drop(columns=["_arquivo_origem", "_processado_em"])
        e = Entrega(df_x, "p", "p.xlsx", {"id_resposta": "id", "nota_nps": "numero", "curso": "categoria"}, {"nps": "nota_nps"}, {})
        boa = Path(d) / "boa.xlsx"
        e.gerar(boa)
        wb = load_workbook(boa)
        wb["Conferência"]["C5"] = 2          # número digitado por cima da fórmula
        wb["Conferência"].freeze_panes = None
        wb["Indicadores"]["B5"] = "=#REF!+1"
        wb["Indicadores"]["B6"].number_format = "General"
        sabotada = Path(d) / "sabotada.xlsx"
        wb.save(sabotada)
        cx = lambda p: c(bom, excel={"arquivo": str(p)})
        r["coord_excel"] = [("Aprova entrega íntegra", not conf_integridade_excel(cx(boa))),
                            ("Pega fórmula quebrada", bool(conf_integridade_excel(cx(sabotada))))]
        r["xl_pleno_formulas"] = [("Aprova conferência em fórmula", not conf_formulas_vivas(cx(boa))),
                                  ("Pega número digitado", bool(conf_formulas_vivas(cx(sabotada))))]
        r["xl_junior_formatos"] = [("Aprova formatos", not conf_formatos(cx(boa))),
                                   ("Pega número sem formato e cabeçalho solto", len(conf_formatos(cx(sabotada))) >= 2)]
        r["xl_estagiario"] = [("Amostra de conferências OK", not conf_amostra_conferencias(cx(boa))),
                              ("Executa sem quebrar na entrega sabotada", isinstance(conf_amostra_conferencias(cx(sabotada)), list))]
        # governança
        pessoal = bom.assign(contato=["a@b.com", "c@d.com.br", "e@f.org"])
        r["gov_pleno_lgpd"] = [("Não acusa base sem dado pessoal", not conf_dados_pessoais(c(bom))),
                               ("Pega coluna de e-mail", bool(conf_dados_pessoais(c(pessoal))))]

        class _Cfg:
            logs = Path(d) / "logs"
            dados: dict = {}
            saida: dict = {}
        registrar_catalogo(c(bom, cfg=_Cfg()))
        cat = json.loads((_Cfg.logs / "catalogo_dados.json").read_text(encoding="utf-8"))
        r["gov_estagiaria"] = [("Registra a tabela no catálogo", "prova.base" in cat),
                               ("Registra colunas e tipos", cat["prova.base"]["colunas"].get("nota_nps") == "Int64")]

    id_grande = bom.assign(id_resposta=pd.array([1234567890123456, 2, 3], dtype="Int64"))
    r["excel_supervisor"] = [("Base compatível com o Excel passa", not conf_compatibilidade_excel(c(bom))),
                             ("Barra ID de 16 dígitos guardado como número", bool(conf_compatibilidade_excel(c(id_grande))))]

    class _CfgSla:
        dados = {"quadro": {"sla_segundos": 60}}
    r["ger_engenharia"] = [("Dentro do prazo não reclama", not conf_sla(_ctx_prova(cfg=_CfgSla(), duracao=10))),
                           ("Fora do prazo aponta", bool(conf_sla(_ctx_prova(cfg=_CfgSla(), duracao=90))))]
    return r
