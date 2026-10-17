"""Centro de treinamento: planilhas difíceis, geradas na hora, com resposta conhecida.

Objetivo do dono: a equipe não pode ser engessada. Cada caso abaixo imita um problema real
(cabeçalho fora do lugar, cabeçalho em duas linhas, células mescladas, várias abas, datas e números
em formatos misturados, NPS com nome de pergunta, CSV de outro sistema, arquivo corrompido...).

O treinamento roda o ROBÔ DE VERDADE (o mesmo pipeline da produção) numa pasta temporária, compara o
resultado com o esperado e registra quem passou. Caso que falha = defeito real a corrigir, não "nota baixa".
Ninguém treina em produção: nada aqui toca as suas pastas, o banco ou o histórico.
"""

from __future__ import annotations

import csv
import json
import tempfile
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable

import pandas as pd
from openpyxl import Workbook


@dataclass
class Caso:
    id: str
    titulo: str
    problema: str            # o que acontece na vida real
    responsaveis: list[str]  # quem precisa dominar
    gerar: Callable[[Path], Path]
    verificar: Callable[[object], list[str]]
    regra: dict = field(default_factory=dict)


def _xlsx(pasta: Path, nome: str, linhas: list[list], titulo_aba: str = "Dados", mesclar: list[str] = (),
          extras: dict[str, list[list]] | None = None, ocultas: set[str] = frozenset()) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = titulo_aba
    for l in linhas:
        ws.append(l)
    for faixa in mesclar:
        ws.merge_cells(faixa)
    for aba, conteudo in (extras or {}).items():
        w = wb.create_sheet(aba)
        for l in conteudo:
            w.append(l)
        if aba in ocultas:
            w.sheet_state = "hidden"
    caminho = pasta / nome
    wb.save(caminho)
    return caminho


def _tabela(r, nome: str | None = None) -> pd.DataFrame:
    tabs = getattr(r, "tabelas", {}) or {}
    if not tabs:
        return pd.DataFrame()
    if nome and nome in tabs:
        return tabs[nome]
    return next(iter(tabs.values()))


def _aprovado(r) -> list[str]:
    if r.status != "sucesso":
        erros = [o.mensagem for o in r.relatorio.erros][:2]
        return [f"deveria ser aprovado, foi {r.status}: {'; '.join(erros)}"]
    return []


# =============================================================================== casos
def _c_cabecalho_baixo(p):
    linhas = [["Relatório mensal de satisfação"], ["Gerado pelo sistema X em 03/10/2026"], [], ["Filtro: todas as lojas"], [],
              [None, "Observação: dados preliminares"],
              ["ID", "Loja", "Nota NPS", "Comentário"]]
    linhas += [[i, f"Loja {i % 4}", i % 11, "ok"] for i in range(1, 41)]
    return _xlsx(p, "cabecalho_baixo.xlsx", linhas)


def _v_cabecalho_baixo(r):
    df = _tabela(r)
    f = _aprovado(r)
    if not f:
        if {"id", "loja", "nota_nps"} - set(df.columns):
            f.append(f"colunas esperadas id, loja, nota_nps; vieram {list(df.columns)[:6]}")
        if len(df) != 40:
            f.append(f"esperava 40 linhas, vieram {len(df)}")
    return f


def _c_cabecalho_duas_linhas(p):
    linhas = [["Respondente", None, "Atendimento", None, "Produto", None],
              ["ID", "Região", "Nota", "Comentário", "Nota", "Comentário"]]
    linhas += [[i, ["Sul", "Norte"][i % 2], i % 5 + 1, "x", (i + 2) % 5 + 1, "y"] for i in range(1, 31)]
    return _xlsx(p, "duas_linhas.xlsx", linhas, mesclar=["A1:B1", "C1:D1", "E1:F1"])


def _v_cabecalho_duas_linhas(r):
    df = _tabela(r)
    f = _aprovado(r)
    if not f:
        esperadas = {"atendimento_nota", "produto_nota"}
        if esperadas - set(df.columns):
            f.append(f"as duas colunas 'Nota' deveriam virar atendimento_nota e produto_nota; vieram {list(df.columns)}")
    return f


def _c_mesclado_vertical(p):
    linhas = [["Região", "Loja", "Vendas"]]
    for reg in ("Sul", "Norte"):
        for i in range(3):
            linhas.append([reg if i == 0 else None, f"{reg}-{i}", 100 + i])
    return _xlsx(p, "mesclado.xlsx", linhas, mesclar=["A2:A4", "A5:A7"])


def _v_mesclado_vertical(r):
    df = _tabela(r)
    f = _aprovado(r)
    if not f:
        if "regiao" not in df.columns or df["regiao"].isna().any():
            f.append("célula mesclada de Região deveria valer para todas as linhas do bloco")
        elif sorted(df["regiao"].value_counts().tolist()) != [3, 3]:
            f.append(f"esperava 3 lojas por região, veio {df['regiao'].value_counts().to_dict()}")
    return f


def _c_varias_abas(p):
    cab = ["ID", "Nota NPS", "Canal"]
    abas = {m: [cab] + [[i + k * 100, (i * 3) % 11, "App"] for i in range(1, 21)] for k, m in enumerate(["Fev", "Mar"], 1)}
    abas["Instruções"] = [["Preencha uma linha por respondente."], ["Não altere o cabeçalho."]]
    abas["Base antiga"] = [cab] + [[999, 10, "Loja"]]
    return _xlsx(p, "varias_abas.xlsx", [cab] + [[i, i % 11, "Site"] for i in range(1, 21)], titulo_aba="Jan",
                 extras=abas, ocultas={"Base antiga"})


def _v_varias_abas(r):
    df = _tabela(r)
    f = _aprovado(r)
    if not f:
        if len(df) != 60:
            f.append(f"as 3 abas mensais (mesmo cabeçalho) deveriam ser empilhadas: esperava 60 linhas, vieram {len(df)}")
        if "_aba_origem" not in df.columns:
            f.append("cada linha deveria dizer de qual aba veio (_aba_origem)")
        if 999 in set(pd.to_numeric(df.get("id", pd.Series(dtype=float)), errors="coerce")):
            f.append("aba oculta não deveria entrar")
    return f


def _c_numeros_misturados(p):
    linhas = [["Produto", "Preço", "Desconto", "Margem", "Devolução"]]
    vals = [("R$ 1.234,56", "12,5%", "0,35", "(150,00)"), ("R$ 99,90", "5%", "0,4", "0"), ("R$ 10.000,00", "0%", "0,25", "(1.000,00)")]
    for i in range(30):
        a, b, c, d = vals[i % 3]
        linhas.append([f"P{i}", a, b, c, d])
    return _xlsx(p, "numeros.xlsx", linhas)


def _v_numeros_misturados(r):
    df = _tabela(r)
    f = _aprovado(r)
    if not f:
        esperado = {"preco": 1234.56, "desconto": 0.125, "margem": 0.35, "devolucao": -150.0}
        for col, v in esperado.items():
            if col not in df.columns or not pd.api.types.is_numeric_dtype(df[col]):
                f.append(f"{col} deveria ser número")
            elif abs(float(df[col].iloc[0]) - v) > 1e-9:
                f.append(f"{col}: esperava {v}, veio {df[col].iloc[0]}")
    return f


def _c_numeros_us(p):
    caminho = p / "export_sistema_us.csv"
    linhas = [["order_id", "amount", "qty"]] + [[f"A{i}", "1,234.50" if i % 2 else "99.90", str(i)] for i in range(1, 31)]
    with open(caminho, "w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(linhas)
    return caminho


def _v_numeros_us(r):
    df = _tabela(r)
    f = _aprovado(r)
    if not f and (not pd.api.types.is_numeric_dtype(df.get("amount", pd.Series(dtype=object)))
                  or abs(float(df["amount"].iloc[0]) - 1234.5) > 1e-9):
        f.append("valor no formato americano (1,234.50) deveria virar 1234,5")
    return f


def _c_datas_misturadas(p):
    linhas = [["ID", "Data da compra"]]
    formatos = ["01/02/2025", "2025-02-03", 45692, "05-fev-2025", "6 de fevereiro de 2025", "07/02/2025 14:30"]
    for i in range(30):
        linhas.append([i + 1, formatos[i % len(formatos)]])
    return _xlsx(p, "datas.xlsx", linhas)


def _v_datas_misturadas(r):
    df = _tabela(r)
    f = _aprovado(r)
    if not f:
        col = df.get("data_da_compra")
        if col is None or not pd.api.types.is_datetime64_any_dtype(col):
            f.append("coluna de data com formatos misturados deveria virar data")
        else:
            esperado = [pd.Timestamp(x) for x in ("2025-02-01", "2025-02-03", "2025-02-04", "2025-02-05", "2025-02-06",
                                                    "2025-02-07 14:30")]
            obtido = list(col.iloc[:6])
            if obtido != esperado:
                f.append(f"datas esperadas {[str(x) for x in esperado]}, vieram {[str(x) for x in obtido]}")
    return f


def _c_totais(p):
    linhas = [["Região", "Loja", "Vendas"]]
    for reg in ("Sul", "Norte"):
        linhas += [[reg, f"{reg}-{i}", 100] for i in range(5)]
        linhas.append([f"Subtotal {reg}", None, 500])
    linhas.append(["TOTAL GERAL", None, 1000])
    linhas.append(["Total Express", "Loja TE", 80])  # loja cujo nome começa com "Total": NÃO é total
    return _xlsx(p, "totais.xlsx", linhas)


def _v_totais(r):
    df = _tabela(r)
    f = _aprovado(r)
    if not f:
        if len(df) != 11:
            f.append(f"esperava 11 linhas (10 lojas + 'Total Express'), vieram {len(df)}")
        if df["vendas"].sum() != 1080:
            f.append(f"soma de vendas deveria ser 1080 (sem contar subtotais), deu {df['vendas'].sum()}")
    return f


def _c_csv_latin1(p):
    caminho = p / "sistema_antigo.csv"
    texto = "Código;Município;Valor\n" + "".join(f"{i:05d};São João;1.234,{i:02d}\n" for i in range(1, 31))
    caminho.write_bytes(texto.encode("cp1252"))
    return caminho


def _v_csv_latin1(r):
    df = _tabela(r)
    f = _aprovado(r)
    if not f:
        if df["municipio"].iloc[0] != "São João":
            f.append(f"acentos do CSV antigo (Windows-1252) deveriam ser preservados, veio {df['municipio'].iloc[0]!r}")
        if df["codigo"].iloc[0] != "00001":
            f.append("código com zero à esquerda deveria continuar texto (00001)")
    return f


def _c_nps_pergunta(p):
    linhas = [["Respondente", "Região",
               "Em uma escala de 0 a 10, o quanto você recomendaria nossa empresa a um amigo ou colega?",
               "De 1 a 5, qual a sua satisfação geral com o atendimento?"]]
    linhas += [[i, ["Sul", "Norte", "Centro"][i % 3], i % 11, i % 5 + 1] for i in range(1, 121)]
    return _xlsx(p, "pesquisa_texto_longo.xlsx", linhas)


def _v_nps_pergunta(r):
    f = _aprovado(r)
    if f:
        return f
    ind = {i["indicador"] for i in (r.analise or {}).get("_indicadores") or [] if i.get("segmento") == "Total"}
    if "NPS" not in ind:
        f.append("a pergunta 'o quanto você recomendaria' deveria ser reconhecida como NPS")
    if not any(x.startswith("CSAT") for x in ind):
        f.append("a pergunta de satisfação de 1 a 5 deveria ser reconhecida como CSAT")
    segs = {i.get("segmento_coluna") for i in (r.analise or {}).get("_indicadores") or []}
    if "regiao" not in segs:
        f.append("a análise deveria segmentar sozinha por região")
    return f


def _c_escala_texto(p):
    escala = ["Muito satisfeito", "Satisfeito", "Nem satisfeito nem insatisfeito", "Insatisfeito", "Muito insatisfeito"]
    linhas = [["ID", "Canal", "Como você avalia o atendimento?"]]
    linhas += [[i, ["App", "Loja"][i % 2], escala[i % 5]] for i in range(1, 101)]
    return _xlsx(p, "escala_texto.xlsx", linhas)


def _v_escala_texto(r):
    f = _aprovado(r)
    if f:
        return f
    total = [i for i in (r.analise or {}).get("_indicadores") or [] if i.get("segmento") == "Total"]
    csat = [i for i in total if str(i["indicador"]).startswith("CSAT")]
    if not csat:
        f.append("escala escrita (Muito satisfeito… Muito insatisfeito) deveria virar CSAT top-2-box")
    elif abs(float(csat[0]["valor"]) - 40.0) > 0.1:
        f.append(f"CSAT top-2-box esperado 40%, veio {csat[0]['valor']}")
    return f


def _c_cabecalho_sujo(p):
    linhas = [["  ID  ", "Nome\ndo cliente", None, "Valor (R$)", "% Desconto", "Nota", "Nota"]]
    linhas += [[i, f"Cliente {i}", None, f"{i},50", f"{i % 3}%", i % 10, (i + 1) % 10] for i in range(1, 31)]
    return _xlsx(p, "cabecalho_sujo.xlsx", linhas)


def _v_cabecalho_sujo(r):
    df = _tabela(r)
    f = _aprovado(r)
    if not f:
        for c in ("id", "nome_do_cliente", "valor_r", "desconto", "nota", "nota_2"):
            if c not in df.columns:
                f.append(f"coluna {c} deveria existir; vieram {list(df.columns)}")
    return f


def _c_erro_formula(p):
    linhas = [["ID", "Receita", "Margem"]] + [[i, 100, 0.2] for i in range(1, 21)]
    linhas[7][2] = "#DIV/0!"
    return _xlsx(p, "erro_formula.xlsx", linhas)


def _v_erro_formula(r):
    if r.status != "reprovado":
        return ["planilha com #DIV/0! deveria ser barrada (o número certo não existe)"]
    if not any("linha 8" in " ".join(o.exemplos) for o in r.relatorio.erros):
        return ["o erro deveria apontar a linha exata do Excel (linha 8)"]
    return []


def _c_vazia(p):
    return _xlsx(p, "so_cabecalho.xlsx", [["ID", "Nota"]])


def _v_vazia(r):
    if r.status != "reprovado":
        return ["planilha só com cabeçalho deveria ser barrada com explicação"]
    return []


def _c_corrompido(p):
    caminho = p / "corrompido.xlsx"
    caminho.write_bytes(b"isto nao e um excel" * 50)
    return caminho


def _v_corrompido(r):
    if r.status != "reprovado":
        return ["arquivo corrompido deveria ser barrado (e o robô não pode cair)"]
    if not any("não pôde ser aberto" in o.mensagem or "corromp" in o.mensagem for o in r.relatorio.erros):
        return ["a mensagem deveria explicar que o arquivo está corrompido"]
    return []


def _c_ids_e_booleanos(p):
    linhas = [["Matrícula", "CEP", "Ativo", "Telefone"]]
    linhas += [[f"{i:06d}", f"0{i:04d}-000", "Sim" if i % 2 else "Não", f"(11) 9{i:04d}-0000"] for i in range(1, 31)]
    return _xlsx(p, "ids.xlsx", linhas)


def _v_ids_e_booleanos(r):
    df = _tabela(r)
    f = _aprovado(r)
    if not f:
        if df["matricula"].iloc[0] != "000001":
            f.append("matrícula com zeros à esquerda deveria ser preservada")
        if df["cep"].iloc[0] != "00001-000":
            f.append("CEP deveria ser preservado como texto")
    return f


def _c_regra_flexivel(p):
    linhas = [["Cod. Resposta", "NOTA (0-10)", "Data Resp."]] + [[i, i % 11, "01/03/2026"] for i in range(1, 31)]
    return _xlsx(p, "nps_campanha_marco.xlsx", linhas)


def _v_regra_flexivel(r):
    df = _tabela(r)
    f = _aprovado(r)
    if not f and not {"id_resposta", "nota_nps", "data_resposta"} <= set(df.columns):
        f.append(f"os sinônimos da regra deveriam renomear as colunas; vieram {list(df.columns)}")
    return f


CASOS = [
    Caso("cabecalho_baixo", "Cabeçalho na 7ª linha", "Relatórios exportados vêm com título, data e filtros antes da tabela.",
         ["leitura"], _c_cabecalho_baixo, _v_cabecalho_baixo),
    Caso("cabecalho_duas_linhas", "Cabeçalho em duas linhas (grupos mesclados)",
         "Pesquisas agrupam colunas: 'Atendimento > Nota' e 'Produto > Nota'.", ["leitura", "limpeza"],
         _c_cabecalho_duas_linhas, _v_cabecalho_duas_linhas),
    Caso("mesclado_vertical", "Células mescladas na coluna", "A região aparece uma vez e vale para várias lojas.",
         ["leitura"], _c_mesclado_vertical, _v_mesclado_vertical),
    Caso("varias_abas", "Uma aba por mês + aba de instruções + aba oculta",
         "Planilhas mensais acumulam abas; instruções e abas ocultas não podem entrar.", ["leitura", "supervisor"],
         _c_varias_abas, _v_varias_abas),
    Caso("numeros_misturados", "Moeda, %, decimal e negativo entre parênteses", "Financeiro: R$ 1.234,56 / 12,5% / (150,00).",
         ["limpeza", "excel_formulas"], _c_numeros_misturados, _v_numeros_misturados),
    Caso("numeros_us", "CSV de sistema americano", "Exportação com 1,234.50 (ponto decimal).", ["leitura", "limpeza"],
         _c_numeros_us, _v_numeros_us),
    Caso("datas_misturadas", "Datas em 6 formatos (inclui número do Excel e mês por extenso)",
         "Bases montadas por várias pessoas.", ["limpeza"], _c_datas_misturadas, _v_datas_misturadas),
    Caso("totais", "Subtotais no meio e loja chamada 'Total Express'",
         "Relatório com subtotais; nem tudo que começa com 'Total' é total.", ["limpeza", "supervisor"],
         _c_totais, _v_totais),
    Caso("csv_latin1", "CSV antigo (Windows-1252, ponto e vírgula)", "Sistemas legados exportam sem UTF-8.",
         ["leitura"], _c_csv_latin1, _v_csv_latin1),
    Caso("nps_pergunta", "NPS e CSAT com o texto da pergunta no cabeçalho",
         "Ferramentas de pesquisa exportam a pergunta inteira como nome da coluna.",
         ["analise_indicadores", "analise_perfil", "excel_pesquisa"], _c_nps_pergunta, _v_nps_pergunta),
    Caso("escala_texto", "Escala de satisfação escrita por extenso", "'Muito satisfeito… Muito insatisfeito' em vez de 1 a 5.",
         ["analise_indicadores", "analise_estatistica"], _c_escala_texto, _v_escala_texto),
    Caso("cabecalho_sujo", "Cabeçalho com quebra de linha, símbolos, coluna sem nome e nomes repetidos",
         "Planilhas feitas à mão.", ["limpeza"], _c_cabecalho_sujo, _v_cabecalho_sujo),
    Caso("erro_formula", "Erro de fórmula (#DIV/0!)", "O número certo não existe: tem de barrar e apontar a linha.",
         ["leitura", "validacao"], _c_erro_formula, _v_erro_formula),
    Caso("so_cabecalho", "Planilha só com cabeçalho", "Exportação que veio vazia.", ["validacao"], _c_vazia, _v_vazia),
    Caso("corrompido", "Arquivo corrompido", "Download interrompido ou arquivo renomeado.", ["leitura", "dev_pleno"],
         _c_corrompido, _v_corrompido),
    Caso("ids_e_booleanos", "Matrícula e CEP com zeros, Sim/Não, telefone", "Cadastros de RH e CRM.", ["limpeza"],
         _c_ids_e_booleanos, _v_ids_e_booleanos),
    Caso("regra_flexivel", "Mesma pesquisa com cabeçalhos diferentes",
         "Cada campanha muda o nome da coluna; a regra aceita sinônimos.", ["validacao", "coord_engenharia"],
         _c_regra_flexivel, _v_regra_flexivel,
         regra={"padrao": "nps_campanha*", "sinonimos": {"id_resposta": ["cod_resposta", "codigo_da_resposta"],
                                                         "nota_nps": ["nota_0_10", "nota"],
                                                         "data_resposta": ["data_resp"]},
                "colunas_obrigatorias": ["id_resposta", "nota_nps"], "tipos": {"data_resposta": "data"}}),
]
POR_ID = {c.id: c for c in CASOS}


# =============================================================================== execução
def rodar(ids: list[str] | None = None) -> dict:
    """Roda os casos no pipeline real, isolado numa pasta temporária. Nunca toca as pastas de produção."""
    from robo.config import carregar_config
    from robo.pipeline import processar_arquivo
    resultados = []
    inicio = time.perf_counter()
    for caso in [POR_ID[i] for i in ids] if ids else CASOS:
        with tempfile.TemporaryDirectory(prefix="robo_treino_") as d:
            base = Path(d)
            entrada = base / "OneDrive" / "Bases-Limpeza"
            entrada.mkdir(parents=True)
            regras = [caso.regra] if caso.regra else []
            cfg = carregar_config(base / "nao_existe.yaml", sobrescrever={
                "pastas": {"entrada": str(entrada), "logs": str(base / "logs"), "agentes": str(base / "agentes"),
                           "habilidades": str(base / "habilidades")},
                "regras": regras, "banco": {"habilitado": False},
                "contratacao_automatica": {"habilitada": False}})
            t0 = time.perf_counter()
            try:
                arquivo = caso.gerar(entrada)
                r = processar_arquivo(arquivo, cfg)
                falhas = caso.verificar(r)
            except Exception as e:  # o robô não pode cair; se cair, é falha grave do treinamento
                falhas = [f"o robô quebrou: {type(e).__name__}: {e}"]
            resultados.append({"id": caso.id, "titulo": caso.titulo, "problema": caso.problema,
                               "responsaveis": caso.responsaveis, "ok": not falhas, "falhas": falhas,
                               "segundos": round(time.perf_counter() - t0, 2)})
    return {"data_hora": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "total": len(resultados),
            "aprovados": sum(r["ok"] for r in resultados), "segundos": round(time.perf_counter() - inicio, 1),
            "casos": resultados}


def registrar(cfg, resultado: dict) -> None:
    pasta = Path(cfg.logs) / "rh"
    pasta.mkdir(parents=True, exist_ok=True)
    with open(pasta / "treinamentos.jsonl", "a", encoding="utf-8") as f:
        f.write(json.dumps(resultado, ensure_ascii=False) + "\n")


def historico(cfg, limite: int = 30) -> list[dict]:
    arq = Path(cfg.logs) / "rh" / "treinamentos.jsonl"
    if not arq.exists():
        return []
    linhas = arq.read_text(encoding="utf-8").splitlines()[-limite:]
    saida = []
    for l in linhas:
        try:
            saida.append(json.loads(l))
        except json.JSONDecodeError:
            continue
    return saida
