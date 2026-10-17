"""Leitura de planilhas (.xlsx, .xlsm, .xls, .csv, .txt) com detecção automática de
codificação, separador e linha de cabeçalho.

Importante: o índice de cada DataFrame devolvido é o NÚMERO DA LINHA no arquivo original.
Assim, quando um erro aparece, o relatório diz exatamente "linha 57" — igual ao que você vê no Excel.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

EXTENSOES_EXCEL = {".xlsx", ".xlsm", ".xls"}
EXTENSOES_TEXTO = {".csv", ".txt"}
EXTENSOES_SUPORTADAS = EXTENSOES_EXCEL | EXTENSOES_TEXTO

CODIFICACOES = ("utf-8-sig", "cp1252", "latin-1")
LIMITE_MESCLAS_MB = 30  # até este tamanho o Excel é aberto completo para enxergar células mescladas
SEPARADORES = ";,\t|"


class ErroLeitura(Exception):
    pass


def _vazio(v) -> bool:
    if v is None:
        return True
    if isinstance(v, str):
        return not v.strip()
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def _decodificar(conteudo: bytes) -> tuple[str, str]:
    for cod in CODIFICACOES:
        try:
            return conteudo.decode(cod), cod
        except UnicodeDecodeError:
            continue
    raise ErroLeitura("Não consegui identificar a codificação do arquivo de texto.")


def _detectar_separador(texto: str) -> str:
    amostra = "\n".join(texto.splitlines()[:50])
    try:
        return csv.Sniffer().sniff(amostra, delimiters=SEPARADORES).delimiter
    except csv.Error:
        contagens = {s: amostra.count(s) for s in SEPARADORES}
        melhor = max(contagens, key=contagens.get)
        return melhor if contagens[melhor] else ";"


def _ler_openpyxl(caminho: Path) -> dict[str, pd.DataFrame]:
    """Leitura direta pelo openpyxl. Motivo: o pandas.read_excel transforma células com erro de
    fórmula (#REF!, #DIV/0!, #N/D) em vazio sem avisar. Aqui elas chegam como texto e o robô acusa.
    data_only=True lê o RESULTADO das fórmulas (o que você vê na tela), não a fórmula em si.
    Abas ocultas são ignoradas.

    Células mescladas (arquivos até 30 MB):
      - mescla VERTICAL numa coluna (ex.: "Sul" valendo para 3 lojas): o valor é repetido em todas as linhas;
      - mescla HORIZONTAL (ex.: grupo "Atendimento" sobre "Nota" e "Comentário"): vira rótulo de grupo,
        usado para diferenciar colunas com o mesmo nome (atendimento_nota × produto_nota).
    """
    completo = caminho.stat().st_size <= LIMITE_MESCLAS_MB * 1024 * 1024
    wb = load_workbook(caminho, read_only=not completo, data_only=True)
    try:
        abas = {}
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            df = pd.DataFrame(list(ws.iter_rows(values_only=True)), dtype=object)
            if completo and ws.merged_cells.ranges and not df.empty:
                horizontais = []
                for rg in ws.merged_cells.ranges:
                    r1, c1, r2, c2 = rg.min_row - 1, rg.min_col - 1, rg.max_row - 1, rg.max_col - 1
                    if r1 >= len(df) or c1 >= len(df.columns):
                        continue
                    valor = df.iat[r1, c1]
                    if r2 > r1 and c1 == c2:
                        df.iloc[r1:r2 + 1, c1] = valor
                    elif r1 == r2 and c2 > c1:
                        horizontais.append((r1 + 1, c1, min(c2, len(df.columns) - 1), valor))
                df.attrs["mesclas_horizontais"] = horizontais
            abas[ws.title] = df
        return abas
    finally:
        wb.close()


def _ler_brutos(caminho: Path) -> dict[str, pd.DataFrame]:
    ext = caminho.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            return _ler_openpyxl(caminho)
        except Exception as e:  # arquivo corrompido, protegido por senha etc.
            raise ErroLeitura(f"O Excel não pôde ser aberto ({type(e).__name__}: {e}). "
                              "Verifique se o arquivo não está corrompido ou protegido por senha.") from e
    if ext == ".xls":
        try:
            # dtype=object: lê o valor exatamente como está na célula, sem o pandas "adivinhar" tipos
            return pd.read_excel(caminho, sheet_name=None, header=None, dtype=object, engine="xlrd")
        except ImportError as e:
            raise ErroLeitura(f"Falta uma biblioteca para ler {ext}: {e}. Rode instalar.bat novamente.") from e
        except Exception as e:  # arquivo corrompido, protegido por senha etc.
            raise ErroLeitura(f"O Excel não pôde ser aberto ({type(e).__name__}: {e}). "
                              "Verifique se o arquivo não está corrompido ou protegido por senha.") from e
    if ext in EXTENSOES_TEXTO:
        texto, _ = _decodificar(caminho.read_bytes())
        if not texto.strip():
            return {}
        sep = _detectar_separador(texto)
        try:
            df = pd.read_csv(io.StringIO(texto), sep=sep, header=None, dtype=str,
                             keep_default_na=False, engine="python", skip_blank_lines=False)
        except Exception as e:
            raise ErroLeitura(f"Não consegui ler o texto como tabela ({e}).") from e
        return {caminho.stem: df}
    raise ErroLeitura(f"Formato '{ext}' não suportado. Use: {', '.join(sorted(EXTENSOES_SUPORTADAS))}.")


def _localizar_cabecalho(df: pd.DataFrame, max_linhas: int) -> int:
    """Posição da linha de cabeçalho: ignora títulos/linhas soltas acima da tabela.

    Critério: primeira linha (entre as N primeiras) que tenha pelo menos 60% do maior número
    de células preenchidas da janela e cujas células sejam majoritariamente texto.
    """
    janela = df.head(max_linhas)
    preenchidas = janela.notna().sum(axis=1)
    maximo = int(preenchidas.max()) if len(preenchidas) else 0
    for pos, (_, linha) in enumerate(janela.iterrows()):
        valores = linha.dropna()
        if len(valores) == 0 or len(valores) < max(1, 0.6 * maximo):
            continue
        textos = sum(isinstance(v, str) for v in valores)
        if textos / len(valores) >= 0.5:
            return pos
    return 0


def _montar_tabela(bruto: pd.DataFrame, max_linhas: int) -> pd.DataFrame | None:
    df = bruto.copy()
    # número da linha como o usuário vê no Excel (1 = primeira linha)
    df.index = range(1, len(df) + 1)
    df = df.map(lambda v: pd.NA if _vazio(v) else v)
    df = df.dropna(how="all").dropna(axis=1, how="all")
    if df.empty:
        return None
    pos = _localizar_cabecalho(df, max_linhas)
    linha_cab = int(df.index[pos])
    cabecalho = df.iloc[pos].tolist()
    corpo = df.iloc[pos + 1:].copy()
    nomes = []
    for i, nome in enumerate(cabecalho, 1):
        nomes.append(f"coluna_{i}" if _vazio(nome) else " ".join(str(nome).split()))
    nomes = _desempatar(nomes, list(df.columns), _grupos_acima(bruto, linha_cab))
    corpo.columns = nomes
    corpo.attrs["linha_cabecalho"] = linha_cab
    return corpo


def _grupos_acima(bruto: pd.DataFrame, linha_cab: int) -> dict[int, str]:
    """Rótulo de grupo de cada coluna, vindo da linha logo acima do cabeçalho (mesclas horizontais)."""
    grupos = {}
    for linha, c1, c2, valor in bruto.attrs.get("mesclas_horizontais") or []:
        if linha == linha_cab - 1 and not _vazio(valor):
            for c in range(c1, c2 + 1):
                grupos[c] = str(valor).strip()
    return grupos


def _desempatar(nomes: list[str], posicoes: list[int], grupos: dict[int, str]) -> list[str]:
    """Nomes repetidos ("Nota", "Nota") ganham o grupo de cima ("Atendimento Nota") ou um número ("Nota (2)")."""
    contagem: dict[str, int] = {}
    for n in nomes:
        contagem[n.lower()] = contagem.get(n.lower(), 0) + 1
    saida, vistos = [], {}
    for n, pos in zip(nomes, posicoes):
        if contagem[n.lower()] > 1 and grupos.get(pos):
            n = f"{grupos[pos]} {n}"
        chave = n.lower()
        vistos[chave] = vistos.get(chave, 0) + 1
        saida.append(n if vistos[chave] == 1 else f"{n} ({vistos[chave]})")
    return saida


def ler_arquivo(caminho: str | Path, abas: str = "primeira", linhas_busca_cabecalho: int = 20) -> dict[str, pd.DataFrame]:
    """Devolve {nome_da_aba: DataFrame} já com o cabeçalho certo e sem linhas/colunas 100% vazias."""
    caminho = Path(caminho)
    if not caminho.exists():
        raise ErroLeitura(f"Arquivo não encontrado: {caminho}")
    brutos = _ler_brutos(caminho)
    tabelas: dict[str, pd.DataFrame] = {}
    for nome_aba, bruto in brutos.items():
        tabela = _montar_tabela(bruto, linhas_busca_cabecalho)
        if tabela is None:
            continue
        tabelas[str(nome_aba)] = tabela
    if abas == "primeira" and tabelas:
        return _empilhar_abas_iguais(tabelas)
    return tabelas


def _empilhar_abas_iguais(tabelas: dict[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
    """Modo padrão: lê a primeira aba e EMPILHA as outras que têm exatamente o mesmo cabeçalho
    (ex.: uma aba por mês). Abas com estrutura diferente (instruções, resumos) não entram, mas ficam
    listadas para o relatório avisar — nada é ignorado em silêncio.
    Nas abas empilhadas, a "linha" vira "Aba!linha" (ex.: Fev!12), para o erro continuar apontando o lugar exato.
    """
    from robo.limpeza import normalizar_nome_coluna
    nomes = list(tabelas)
    primeira = tabelas[nomes[0]]
    assinatura = lambda df: [normalizar_nome_coluna(c) for c in df.columns]
    iguais = [n for n in nomes[1:] if assinatura(tabelas[n]) == assinatura(primeira)]
    ignoradas = [n for n in nomes[1:] if n not in iguais]
    if iguais:
        partes = []
        for n in [nomes[0]] + iguais:
            df = tabelas[n].copy()
            df.index = [f"{n}!{i}" for i in df.index]
            df["_aba_origem"] = n
            partes.append(df)
        resultado = pd.concat(partes)
        resultado.attrs["linha_cabecalho"] = primeira.attrs.get("linha_cabecalho")
        resultado.attrs["abas_empilhadas"] = [nomes[0]] + iguais
    else:
        resultado = primeira
    resultado.attrs["abas_ignoradas"] = ignoradas
    return {nomes[0]: resultado}
