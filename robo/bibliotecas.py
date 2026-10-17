"""Guilda Python: um especialista para cada biblioteca que o robô usa de verdade.

Cada especialista fiscaliza os COMPORTAMENTOS da sua biblioteca dos quais o robô depende. Se uma atualização mudar
um deles (ex.: o pandas passar a manter grupos vazios), o especialista pega na revisão antes de virar número errado.
Também captura avisos de depreciação (FutureWarning/DeprecationWarning) emitidos pelo código do robô.

Só há especialista para biblioteca em uso (requisitos.txt + sqlite3 da biblioteca padrão). Contratar especialista de
scikit-learn, statsmodels ou matplotlib, que o robô não usa, seria cargo sem trabalho real.
"""

from __future__ import annotations

import io
import sqlite3
import tempfile
import warnings
from pathlib import Path


def _pandas() -> list[tuple[str, bool]]:
    import pandas as pd
    df = pd.DataFrame({"g": ["a", None, "a", "b"], "v": [1.0, 2.0, 3.0, None]})
    agrup = df.groupby("g", dropna=True)["v"].mean()
    lido = pd.read_csv(io.StringIO("id;cep\n001;01310-100\n"), sep=";", dtype=str)
    return [
        ("groupby descarta grupo vazio (a calculadora conta com isso)", list(agrup.index) == ["a", "b"]),
        ("média ignora vazio", agrup["a"] == 2.0 and pd.isna(agrup["b"])),
        ("to_numeric com coerce vira NaN em texto", pd.isna(pd.to_numeric(pd.Series(["x"]), errors="coerce")[0])),
        ("dtype=str preserva zero à esquerda", lido["id"][0] == "001"),
        ("Int64 aceita vazio sem virar float", str(pd.array([1, None], dtype="Int64").dtype) == "Int64"),
        ("value_counts(normalize) soma 1", abs(pd.Series(list("aab")).value_counts(normalize=True).sum() - 1) < 1e-12),
    ]


def _numpy() -> list[tuple[str, bool]]:
    import numpy as np
    x = np.linalg.lstsq(np.array([[1.0, 0], [0, 2.0], [1, 1]]), np.array([1.0, 4, 3]), rcond=None)[0]
    a, b = np.random.default_rng(42).normal(size=3), np.random.default_rng(42).normal(size=3)
    return [
        ("mínimos quadrados resolve sistema conhecido (drivers dependem disso)", np.allclose(x, [1, 2])),
        ("gerador com semente é reproduzível (segmentação depende disso)", np.array_equal(a, b)),
        ("correlação de série idêntica é 1", np.isclose(np.corrcoef([1, 2, 3], [2, 4, 6])[0, 1], 1.0)),
    ]


def _openpyxl() -> list[tuple[str, bool]]:
    from openpyxl import Workbook, load_workbook
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "t.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"], ws["A2"], ws["B1"] = 10, "=A1*2", 1234.5
        ws["B1"].number_format = "#,##0.00"
        ws.merge_cells("C1:D1")
        wb.save(p)
        f = load_workbook(p)
        v = load_workbook(p, data_only=True)
        return [
            ("fórmula é gravada como fórmula (Excel recalcula)", f.active["A2"].value == "=A1*2"),
            ("data_only sem recálculo devolve vazio (não inventa valor)", v.active["A2"].value is None),
            ("formato de número persiste", f.active["B1"].number_format == "#,##0.00"),
            ("célula mesclada é detectada", "C1:D1" in [str(r) for r in f.active.merged_cells.ranges]),
        ]


def _xlrd() -> list[tuple[str, bool]]:
    import xlrd
    with tempfile.TemporaryDirectory() as d:
        from openpyxl import Workbook
        p = Path(d) / "t.xlsx"
        Workbook().save(p)
        try:
            xlrd.open_workbook(str(p))
            recusou = False
        except Exception:
            recusou = True
    maior = int(xlrd.__VERSION__.split(".")[0]) if hasattr(xlrd, "__VERSION__") else int(xlrd.__version__.split(".")[0])
    return [("versão 2+ (só .xls antigo, sem macro)", maior >= 2),
            ("recusa .xlsx (o robô usa openpyxl para .xlsx)", recusou)]


def _yaml() -> list[tuple[str, bool]]:
    import yaml
    try:
        yaml.safe_load("!!python/object/apply:os.system ['echo x']")
        bloqueou = False
    except yaml.YAMLError:
        bloqueou = True
    volta = yaml.safe_load(yaml.safe_dump({"região": "São Paulo"}, allow_unicode=True))
    return [("safe_load recusa código embutido no config (segurança)", bloqueou),
            ("acento sobrevive à ida e volta", volta == {"região": "São Paulo"})]


def _psycopg() -> list[tuple[str, bool]]:
    import psycopg
    from psycopg import sql
    q = sql.SQL("SELECT {} FROM {}").format(sql.Identifier('col"x'), sql.Identifier("s", "t"))
    return [("versão 3+", int(psycopg.__version__.split(".")[0]) >= 3),
            ("composição de SQL com identificador", isinstance(q, sql.Composed)),
            ("aspas dentro do nome são escapadas (contra injeção)", sql.Identifier('col"x').as_string(None) == '"col""x"'),
            ("valor vai como parâmetro, nunca colado no texto", sql.SQL("x = {}").format(sql.Placeholder()).as_string(None) == "x = %s")]


def _sqlite() -> list[tuple[str, bool]]:
    with tempfile.TemporaryDirectory() as d:
        c = sqlite3.connect(Path(d) / "t.db")
        c.execute("CREATE TABLE t (x INTEGER)")
        c.commit()
        try:
            with c:
                c.execute("INSERT INTO t VALUES (1)")
                raise RuntimeError
        except RuntimeError:
            pass
        sobrou = c.execute("SELECT COUNT(*) FROM t").fetchone()[0]
        modo = c.execute("PRAGMA journal_mode=WAL").fetchone()[0]
        c.close()
    return [("transação desfaz tudo em erro (memória da empresa depende disso)", sobrou == 0),
            ("modo WAL disponível", str(modo).lower() == "wal")]


BIBLIOTECAS = {
    "pandas": ("esp_pandas", _pandas), "numpy": ("esp_numpy", _numpy), "openpyxl": ("esp_openpyxl", _openpyxl),
    "xlrd": ("esp_xlrd", _xlrd), "PyYAML": ("esp_yaml", _yaml), "psycopg": ("esp_psycopg", _psycopg),
    "sqlite3": ("esp_sqlite", _sqlite),
}


def versao(nome: str) -> str:
    mod = {"PyYAML": "yaml", "sqlite3": "sqlite3"}.get(nome, nome)
    import importlib
    m = importlib.import_module(mod)
    return getattr(m, "__version__", None) or getattr(m, "sqlite_version", None) or "?"


def revisar(nome: str) -> dict:
    """Roda os guardas da biblioteca e captura avisos de depreciação."""
    resp, funcao = BIBLIOTECAS[nome]
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        try:
            checks = funcao()
        except Exception as e:
            checks = [(f"revisão não rodou ({type(e).__name__}: {e})", False)]
    falhas = [n for n, ok in checks if not ok]
    deprec = sorted({f"{w.category.__name__}: {str(w.message)[:160]}" for w in avisos
                     if issubclass(w.category, (DeprecationWarning, FutureWarning))})
    return {"biblioteca": nome, "especialista": resp, "versao": versao(nome), "checks": checks, "falhas": falhas,
            "depreciacoes": deprec, "ok": not falhas}


def capturar_depreciacoes(funcao) -> tuple[object, list[str]]:
    """Roda código do robô e devolve (resultado, avisos de depreciação que ele emitiu)."""
    with warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        r = funcao()
    return r, sorted({f"{w.category.__name__}: {str(w.message)[:200]} ({Path(w.filename).name}:{w.lineno})"
                      for w in avisos if issubclass(w.category, (DeprecationWarning, FutureWarning))
                      and "robo" in str(w.filename)})


def especialista_do_aviso(texto: str) -> str:
    t = texto.lower()
    for nome, (resp, _) in BIBLIOTECAS.items():
        if nome.lower() in t or (nome == "PyYAML" and "yaml" in t):
            return resp
    return "esp_pandas" if "groupby" in t or "dataframe" in t or "series" in t else "head_python"


def provas() -> dict:
    """Prova de cada especialista: a revisão da própria biblioteca passa inteira na versão instalada."""
    r = {}
    for nome, (resp, funcao) in BIBLIOTECAS.items():
        try:
            r[resp] = [(f"[{nome}] {n}", ok) for n, ok in funcao()]
        except Exception as e:
            r[resp] = [(f"[{nome}] revisão", False)]
    r["head_python"] = [("Sabe quem responde por cada biblioteca", especialista_do_aviso("FutureWarning in pandas groupby") == "esp_pandas"),
                        ("Aviso desconhecido sobe para o head", especialista_do_aviso("algo estranho") == "head_python")]
    return r
