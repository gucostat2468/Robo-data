"""RH do Robô de Planilhas: cargos, níveis, permissões e PROCESSO SELETIVO TÉCNICO.

Ninguém trabalha sem passar nas provas:
- cada habilidade tem casos de prova com resposta conhecida (executam o código de verdade);
- cada agente nativo tem provas da sua função;
- um agente só é "contratado" (trabalha) se TODAS as provas das suas habilidades passarem
  e se o nível dele permitir o risco de cada habilidade.

Níveis (do menor para o maior) e o que cada um pode fazer:
  estagiário   trabalha em MODO OBSERVAÇÃO: roda as habilidades numa cópia e só relata o que faria
  júnior       só habilidades de risco BAIXO (validar, padronizar formato)
  pleno        risco BAIXO e MÉDIO (de-para, conversões, colunas calculadas…)
  sênior       todos os riscos, inclusive ALTO (remover linhas/colunas, anonimizar)
  especialista todos os riscos; referência técnica da área

Importante (transparência): os agentes são programas determinísticos. O "conhecimento" deles é a
metodologia codificada nas regras e comprovada pelas provas — não é uma IA que opina.
"""

from __future__ import annotations

import io
import tempfile
from functools import lru_cache
from pathlib import Path

import pandas as pd

NIVEIS = ["estagiario", "junior", "pleno", "senior", "especialista"]
NOME_NIVEL = {"estagiario": "Estagiário", "junior": "Júnior", "pleno": "Pleno", "senior": "Sênior", "especialista": "Especialista",
              "coordenador": "Coordenador(a)", "gerente": "Gerente"}
RISCO_MAXIMO = {"estagiario": "alto", "junior": "baixo", "pleno": "medio", "senior": "alto", "especialista": "alto",
                "coordenador": "alto", "gerente": "alto"}
# cargos de gestão: não executam habilidades nos dados; respondem por pessoas, metas e liberação
CARGOS_GESTAO = ["coordenador", "gerente"]
ORDEM_RISCO = {"baixo": 0, "medio": 1, "alto": 2}
NOME_RISCO = {"baixo": "baixo", "medio": "médio", "alto": "alto"}
AREAS = ["Engenharia de Dados", "Qualidade de Dados", "Governança e Privacidade", "Análise de Dados",
         "Estatística", "Customer Insights e Pesquisa", "Storytelling e Negócio"]

# ----------------------------------------------------------------------------- risco das habilidades
RISCO_BASE = {
    "padronizar_texto": "baixo", "normalizar_documento": "baixo", "validar_email": "baixo", "normalizar_telefone": "baixo",
    "normalizar_cep": "baixo", "validar_regex": "baixo", "validar_datas": "baixo", "completude_minima": "baixo",
    "classificar_nps": "baixo", "classificar_csat": "baixo", "extrair_data": "baixo", "juntar_colunas": "baixo",
    "checar_duplicados": "baixo", "outliers": "baixo",
    "mapear_valores": "medio", "converter_tipo": "medio", "arredondar": "medio", "classificar_faixas": "medio",
    "coluna_calculada": "medio", "dividir_coluna": "medio", "categorizar_texto": "medio", "renomear_colunas": "medio",
    "preencher_vazios": "medio",
    "filtrar_linhas": "alto", "selecionar_colunas": "alto", "anonimizar": "alto",
}


def risco_de(habilidade_id: str, parametros: dict | None = None) -> str:
    """Risco da habilidade COM os parâmetros escolhidos (apagar linha é sempre alto)."""
    p = parametros or {}
    risco = RISCO_BASE.get(habilidade_id, "medio")
    if p.get("invalidos") == "remover" or p.get("acao") in ("remover", "manter_primeiro", "manter_ultimo"):
        risco = "alto"
    elif p.get("invalidos") == "vazio" and risco == "baixo":
        risco = "medio"
    if habilidade_id == "preencher_vazios" and p.get("metodo") in ("media", "mediana", "moda"):
        risco = "alto"  # imputação estatística altera o resultado de indicadores
    if habilidade_id == "filtrar_linhas" and p.get("acao") in ("aviso", "erro"):
        risco = "baixo"
    return risco


def nivel_permite(nivel: str, risco: str) -> bool:
    return ORDEM_RISCO[risco] <= ORDEM_RISCO[RISCO_MAXIMO.get(nivel, "medio")]


# ----------------------------------------------------------------------------- provas das habilidades
def _df(**colunas):
    df = pd.DataFrame(colunas)
    df.index = range(2, len(df) + 2)
    return df


PROVAS_HABILIDADES: dict[str, list] = {
    "padronizar_texto": [
        ("Nome próprio em título, com preposição minúscula", _df(n=["MARIA DA SILVA"]), {"colunas": "n", "modo": "titulo"},
         lambda d, r: d["n"].iloc[0] == "Maria da Silva"),
        ("Sigla em maiúsculo", _df(n=["sp"]), {"colunas": "n", "modo": "maiusculo"}, lambda d, r: d["n"].iloc[0] == "SP")],
    "mapear_valores": [
        ("De-para ignora acento e maiúsculas", _df(m=["A Distância", "online"]),
         {"coluna": "m", "mapa": {"a distancia": "EAD", "online": "EAD"}}, lambda d, r: list(d["m"]) == ["EAD", "EAD"]),
        ("Valor fora da lista oficial reprova quando configurado", _df(m=["Híbrido"]),
         {"coluna": "m", "mapa": {"ead": "EAD"}, "nao_mapeados": "erro"}, lambda d, r: r.tem_erros)],
    "converter_tipo": [
        ("Número brasileiro com moeda", _df(v=["R$ 1.234,56"]), {"colunas": "v", "tipo": "decimal"},
         lambda d, r: abs(d["v"].iloc[0] - 1234.56) < 1e-9),
        ("Texto que não é número reprova", _df(v=["dez"]), {"colunas": "v", "tipo": "inteiro"}, lambda d, r: r.tem_erros)],
    "arredondar": [("Duas casas", _df(v=[1.23456]), {"colunas": "v", "casas": 2}, lambda d, r: d["v"].iloc[0] == 1.23)],
    "normalizar_documento": [
        ("CPF válido formatado", _df(c=["52998224725"]), {"colunas": "c", "formato": "formatado"},
         lambda d, r: d["c"].iloc[0] == "529.982.247-25" and not r.tem_erros),
        ("Recupera zero perdido pelo Excel", _df(c=[1234567890.0]), {"colunas": "c"}, lambda d, r: d["c"].iloc[0] == "01234567890"),
        ("Dígito verificador errado reprova", _df(c=["52998224724"]), {"colunas": "c"}, lambda d, r: r.tem_erros),
        ("CNPJ válido", _df(c=["11.222.333/0001-81"]), {"colunas": "c"}, lambda d, r: d["c"].iloc[0] == "11222333000181")],
    "validar_email": [
        ("Minúsculas e sem espaços", _df(e=[" Ana@X.COM "]), {"colunas": "e"}, lambda d, r: d["e"].iloc[0] == "ana@x.com"),
        ("Sem arroba é inválido", _df(e=["ana.x.com"]), {"colunas": "e"}, lambda d, r: r.tem_erros)],
    "normalizar_telefone": [
        ("Celular com +55", _df(t=["+55 (11) 98765-4321"]), {"colunas": "t"}, lambda d, r: d["t"].iloc[0] == "(11) 98765-4321"),
        ("Número curto é inválido", _df(t=["123"]), {"colunas": "t"}, lambda d, r: r.tem_erros)],
    "normalizar_cep": [("Zero à esquerda recuperado", _df(c=[1310100]), {"colunas": "c"}, lambda d, r: d["c"].iloc[0] == "01310-100")],
    "validar_regex": [("Formato de matrícula", _df(m=["AB123456", "X1"]), {"coluna": "m", "padrao": r"^[A-Z]{2}\d{6}$"},
                       lambda d, r: r.tem_erros and r.erros[0].quantidade == 1)],
    "validar_datas": [("Data futura reprova", _df(d=pd.to_datetime(["2020-01-01", "2099-01-01"])), {"coluna": "d"},
                       lambda d, r: r.tem_erros and r.erros[0].quantidade == 1)],
    "checar_duplicados": [
        ("Chave repetida reprova", _df(i=["1", "1"]), {"colunas": "i"}, lambda d, r: r.tem_erros),
        ("Mantém a resposta mais recente", _df(i=["1", "1"], n=[5, 9]), {"colunas": "i", "acao": "manter_ultimo"},
         lambda d, r: list(d["n"]) == [9])],
    "completude_minima": [("Abaixo do mínimo reprova", _df(c=[1, None, None, None]), {"colunas": "c", "minimo_pct": 90},
                           lambda d, r: r.tem_erros)],
    "outliers": [("Sinaliza valor extremo pelo IQR", _df(v=[10.0, 11, 12, 13, 9, 10, 1000]), {"coluna": "v"},
                  lambda d, r: d["v_outlier"].iloc[-1] == "Sim" and d["v_outlier"].iloc[0] == "Não")],
    "filtrar_linhas": [("Remove registros de teste", _df(s=["ok", "teste"]), {"condicao": "s != 'teste'"}, lambda d, r: len(d) == 1)],
    "preencher_vazios": [("Valor anterior (células mescladas)", _df(g=["A", None, "B"]), {"colunas": "g", "metodo": "anterior"},
                          lambda d, r: list(d["g"]) == ["A", "A", "B"])],
    "coluna_calculada": [("Quantidade × preço", _df(q=[2, 3], p=[1.5, 2.0]), {"nova_coluna": "t", "expressao": "q * p"},
                          lambda d, r: list(d["t"]) == [3.0, 6.0])],
    "classificar_faixas": [("Limite inicia a faixa seguinte", _df(i=[17, 18, 30]),
                            {"coluna": "i", "nova_coluna": "f", "limites": "18, 25", "rotulos": "<18, 18-24, 25+"},
                            lambda d, r: list(d["f"]) == ["<18", "18-24", "25+"])],
    "extrair_data": [("Ano e mês", _df(d=pd.to_datetime(["2026-10-17"])), {"coluna": "d", "partes": "ano, ano_mes"},
                      lambda d, r: d["d_ano"].iloc[0] == 2026 and d["d_ano_mes"].iloc[0] == "2026-10")],
    "dividir_coluna": [("Cidade - UF", _df(c=["São Paulo - SP"]), {"coluna": "c", "separador": " - ", "novas_colunas": "cidade, uf"},
                        lambda d, r: d["cidade"].iloc[0] == "São Paulo" and d["uf"].iloc[0] == "SP")],
    "juntar_colunas": [("Ignora vazios", _df(a=["X", "Y"], b=["1", None]), {"colunas": "a, b", "separador": "/", "nova_coluna": "c"},
                        lambda d, r: list(d["c"]) == ["X/1", "Y"])],
    "classificar_nps": [
        ("Metodologia Bain 9-10 / 7-8 / 0-6", _df(n=[10, 9, 8, 7, 6, 0]), {"coluna": "n"},
         lambda d, r: list(d["grupo_nps"]) == ["Promotor", "Promotor", "Neutro", "Neutro", "Detrator", "Detrator"]),
        ("Nota 11 reprova", _df(n=[10, 11]), {"coluna": "n"}, lambda d, r: r.tem_erros)],
    "classificar_csat": [("Top-2-box na escala 1-5", _df(n=[5, 4, 3]), {"coluna": "n"}, lambda d, r: list(d["satisfeito"]) == ["Sim", "Sim", "Não"])],
    "categorizar_texto": [("Temas por palavra-chave", _df(c=["Mensalidade cara", "Atendente demorou", "Adorei"]),
                           {"coluna": "c", "temas": {"Preço": "caro, mensalidade", "Atendimento": "atend"}},
                           lambda d, r: list(d["tema"]) == ["Preço", "Atendimento", "Outros"])],
    "selecionar_colunas": [("Exclui coluna interna", _df(a=[1], obs=["x"]), {"modo": "excluir", "colunas": "obs"},
                            lambda d, r: list(d.columns) == ["a"])],
    "renomear_colunas": [("Renomeia", _df(vlr=[1]), {"mapa": {"vlr": "valor"}}, lambda d, r: list(d.columns) == ["valor"])],
    "anonimizar": [("Hash estável e irreversível", _df(c=["52998224725", "52998224725"]), {"colunas": "c", "sal": "k"},
                    lambda d, r: d["c"].iloc[0] == d["c"].iloc[1] and d["c"].iloc[0] != "52998224725"),
                   ("Máscara de e-mail", _df(e=["ana@x.com"]), {"colunas": "e", "metodo": "mascarar"}, lambda d, r: d["e"].iloc[0] == "a***@x.com")],
}


def _rodar_caso(habilidade_id, descricao, df, parametros, verificar) -> dict:
    from robo import habilidades
    from robo.relatorio import Relatorio
    rel = Relatorio("prova.xlsx")
    try:
        saida = habilidades.executar_habilidade(df.copy(), {"habilidade": habilidade_id, "parametros": parametros}, rel, "prova", "prova.xlsx", "Prova")
        ok = bool(verificar(saida, rel))
        motivo = "" if ok else "; ".join(o.mensagem for o in rel.erros) or "resultado diferente do esperado"
    except Exception as e:
        ok, motivo = False, f"{type(e).__name__}: {e}"
    return {"prova": descricao, "aprovado": ok, "motivo": motivo}


@lru_cache(maxsize=None)
def _provas_habilidade_cache(habilidade_id: str) -> tuple:
    casos = PROVAS_HABILIDADES.get(habilidade_id, [])
    return tuple(tuple(_rodar_caso(habilidade_id, *c).items()) for c in casos)


def provas_da_habilidade(habilidade_id: str) -> list[dict]:
    """Executa (uma vez por sessão) os casos de prova da habilidade."""
    from robo import habilidades
    sk = habilidades.REGISTRO.get(habilidade_id)
    if sk is not None and sk.origem == "personalizada":
        provas = getattr(sk.funcao, "provas", None) or []
        return [_rodar_caso(habilidade_id, *c) for c in provas]
    return [dict(item) for item in _provas_habilidade_cache(habilidade_id)]


# ----------------------------------------------------------------------------- provas dos agentes nativos e analistas
def _prova_leitura():
    from openpyxl import Workbook

    from robo.leitura import ler_arquivo
    with tempfile.TemporaryDirectory() as d:
        wb = Workbook(); ws = wb.active
        ws.append(["Relatório X"]); ws.append([]); ws.append(["ID", "Valor"]); ws.append([1, "#REF!"])
        wb.save(Path(d) / "a.xlsx")
        df = ler_arquivo(Path(d) / "a.xlsx")["Sheet"]
        csv = Path(d) / "b.csv"; csv.write_bytes("Nome;Valor\nJoão;1,5\n".encode("cp1252"))
        dc = ler_arquivo(csv)["b"]
    return [("Encontra o cabeçalho abaixo do título", list(df.columns) == ["ID", "Valor"]),
            ("Preserva o número da linha do Excel", df.index[0] == 4),
            ("Não deixa erro de fórmula sumir", df.iloc[0, 1] == "#REF!"),
            ("Lê CSV com ';' em codificação Windows", dc.iloc[0, 0] == "João")]


def _prova_limpeza():
    from robo.config import PADROES
    from robo.limpeza import limpar_tabela
    from robo.relatorio import Relatorio
    rel = Relatorio("p")
    df = _df(**{"Grupo": ["A", "B", "B", "C", "Total"], "Valor (R$)": ["R$ 1.234,56", "10,00", "10,00", "N/A", "500"],
                "Matrícula": ["00123", "00456", "00456", "00789", None]})
    out = limpar_tabela(df, "dados", PADROES["limpeza"], {}, rel)
    return [("Nome de coluna padronizado", "valor_r" in out.columns),
            ("Moeda brasileira vira número", out["valor_r"].iloc[0] == 1234.56),
            ("Zeros à esquerda preservados", out["matricula"].iloc[0] == "00123"),
            ("Linha de Total removida (não soma em dobro no Power BI)", "Total" not in set(out["grupo"])),
            ("Duplicata exata removida", len(out) == 3),
            ("Marcador N/A vira vazio", out["valor_r"].isna().sum() == 1)]


def _prova_validacao():
    from robo.relatorio import Relatorio
    from robo.validacao import validar_tabela
    regra = {"colunas_obrigatorias": ["id", "nota"], "chave_unica": ["id"], "faixas": {"nota": {"min": 0, "max": 10}}}
    r1 = Relatorio("p"); validar_tabela(_df(id=["1", "1"], nota=[5, 11]), "t", regra, r1)
    r2 = Relatorio("p"); validar_tabela(_df(id=["1"]), "t", regra, r2)
    r3 = Relatorio("p"); validar_tabela(_df(id=["1", "2"], nota=[0, 10]), "t", regra, r3)
    msgs = " ".join(o.mensagem for o in r1.erros)
    return [("Barra chave repetida", "Chave única" in msgs), ("Barra valor fora da faixa", "fora da faixa" in msgs),
            ("Barra coluna obrigatória ausente", any("obrigatórias" in o.mensagem for o in r2.erros)),
            ("Aprova base correta", not r3.tem_erros)]


def _prova_modelagem():
    from robo.modelagem import modelar, tabela_calendario
    from robo.relatorio import Relatorio
    df = _df(curso=["Direito", "Medicina", "Direito"], nota=[1, 2, 3], data=pd.to_datetime(["2026-01-05"] * 3))
    t = modelar(df, "d", {"fato": "fato", "dimensoes": {"dim_curso": {"colunas": ["curso"], "chave": "id_curso"}}}, Relatorio("p"))
    fato, dim = t["fato"], t["dim_curso"]
    junto = fato.merge(dim, on="id_curso")
    return [("Cria fato e dimensão", set(t) == {"fato", "dim_curso"}),
            ("Dimensão sem repetição", len(dim) == 2),
            ("Fato se liga à dimensão sem perder linhas", list(junto.sort_values("nota")["curso"]) == ["Direito", "Medicina", "Direito"]),
            ("Calendário cobre o ano inteiro", len(tabela_calendario(pd.Series(pd.to_datetime(["2026-03-01"])))) == 365)]


def _prova_banco():
    import numpy as np

    from robo.banco import _valor_pg, tipo_pg
    return [("Inteiro vira bigint", tipo_pg(pd.Series([1, 2], dtype="Int64")) == "bigint"),
            ("Data sem hora vira date", tipo_pg(pd.Series(pd.to_datetime(["2026-01-01"]))) == "date"),
            ("Vazio vira NULL", _valor_pg(np.nan, "double precision") is None),
            ("Número em coluna texto vira texto sem '.0'", _valor_pg(12.0, "text") == "12")]


def _prova_publicacao():
    from robo.escrita import gravar_tabelas
    with tempfile.TemporaryDirectory() as d:
        df = _df(a=[1, 2], data=pd.to_datetime(["2026-01-01", "2026-01-02"]))
        arquivos, _ = gravar_tabelas({"dados": df}, Path(d), "x", {"formato": "ambos", "csv_separador": ";", "csv_decimal": ","})
        lido = pd.read_excel(Path(d) / "x.xlsx")
        sem_temporario = not any(p.name.startswith("~robo_tmp") for p in Path(d).iterdir())
    return [("Grava Excel e CSV", len(arquivos) == 2), ("Releitura idêntica", list(lido["a"]) == [1, 2]),
            ("Gravação atômica (sem arquivo temporário sobrando)", sem_temporario)]


def _prova_recepcao():
    from robo.monitor import eh_candidato, pode_abrir
    with tempfile.TemporaryDirectory() as d:
        bloqueio = Path(d) / "~$base.xlsx"; bloqueio.write_text("x")
        normal = Path(d) / "base.xlsx"; normal.write_text("x")
        ini = Path(d) / "desktop.ini"; ini.write_text("x")
        return [("Ignora arquivo de bloqueio do Excel (~$)", not eh_candidato(bloqueio)),
                ("Ignora arquivos do sistema", not eh_candidato(ini)),
                ("Aceita planilha normal", eh_candidato(normal)),
                ("Confere se o arquivo pode ser aberto", pode_abrir(normal))]


def _base_nps(medicina_ruim=True, n=400, semente=5):
    import numpy as np
    rng = np.random.default_rng(semente)
    cursos = rng.choice(["Direito", "Psicologia", "Medicina"], n)
    notas = rng.choice(range(11), n, p=[.02, .02, .02, .03, .03, .05, .08, .1, .15, .2, .3])
    if medicina_ruim:
        notas = np.where((cursos == "Medicina") & (rng.random(n) < .5), rng.choice(range(0, 6), n), notas)
    return pd.DataFrame({"nota_nps": pd.array(notas, dtype="Int64"), "curso": cursos})


def _prova_analistas():
    import math

    from robo.analise import Analise, calcular_indicador, holm
    r = calcular_indicador("nps", pd.Series([10, 10, 9, 8, 7, 6, 0, 10]), 1.96)
    p, d = .5, .25
    margem_ok = abs(r["margem"] - 100 * 1.96 * math.sqrt((p + d - (p - d) ** 2) / 8)) < 1e-9
    a = Analise(_base_nps(), "t", "x.xlsx", "x", {}, [])
    a.fazer_perfil(); a.fazer_indicadores(); a.fazer_estatistica()
    sig = a.destaques[a.destaques["significativo"]]
    b = Analise(_base_nps(medicina_ruim=False, semente=9), "t", "x.xlsx", "x", {}, [])
    b.fazer_indicadores(); b.fazer_estatistica()
    foto = Analise(_base_nps(), "t", "x.xlsx", "x", {}, []).tirar_foto()
    m = Analise(_base_nps().drop(columns=["curso"]), "t", "y.xlsx", "x", {}, [])
    m.fazer_monitoramento([foto])
    a.fazer_monitoramento([]); a.fazer_resumo()
    return {
        "analise_perfil": [("Reconhece variável categórica", a.papeis["curso"] == "categoria"),
                           ("Descreve média e quartis", "p75" in a.perfil.columns)],
        "analise_indicadores": [("NPS = %promotores − %detratores", abs(r["valor"] - 25) < 1e-9),
                                ("Margem de erro de 95% correta", margem_ok),
                                ("Reconhece a coluna de NPS", a.indicadores_cols == {"nps": "nota_nps"})],
        "analise_estatistica": [("Correção de Holm correta", holm([0.001, 0.04, 0.03], 0.05) == [True, False, False]),
                                ("Encontra o efeito real (Medicina)", "Medicina" in set(sig["segmento"])),
                                ("Não inventa diferença onde não há", not b.destaques["significativo"].any())],
        "analise_monitoramento": [("Percebe coluna que sumiu", "coluna sumiu" in set(m.alertas["tipo"])),
                                  ("Compara indicador entre cargas", "indicador" in set(m.alertas["tipo"]))],
        "analise_relator": [("Resumo cita o indicador com margem", any("NPS:" in l and "margem" in l for l in a.resumo)),
                            ("Resumo separa diferença real de acaso", any("Diferenças reais" in l for l in a.resumo))],
    }


class _ConexaoFalsa:
    """Simula o PostgreSQL para a prova da DBA (contagens controladas)."""
    def __init__(self, contagens):
        self.contagens = contagens

    def execute(self, comando, params=None):
        texto = comando.as_string(None) if hasattr(comando, "as_string") else str(comando)
        valor = 0 if "LEFT JOIN" in texto else self.contagens.pop(0)
        return type("R", (), {"fetchone": lambda _s: (valor,)})()


def _prova_dba():
    from robo.banco import Banco, ErroCargaDados
    b = Banco({}, Path("."))
    b.esperado_tratado = {("tratado", "fato"): 10}
    b.chaves_estrangeiras = [("tratado", "fato", "id_curso", "dim_curso")]
    aprovou = bool(b.reconciliar(_ConexaoFalsa([5, 10]), "lote", {"base": 5}))
    try:
        b.reconciliar(_ConexaoFalsa([5, 9]), "lote", {"base": 5})
        barrou = False
    except ErroCargaDados:
        barrou = True
    return [("Aprova carga que bate com o esperado", aprovou), ("Desfaz carga com linha perdida no caminho", barrou)]


def _prova_supervisor():
    from types import SimpleNamespace

    from robo.config import PADROES
    from robo.producao import conferencia_supervisor
    from robo.relatorio import Relatorio
    cfg = SimpleNamespace(saida={"nome_aba": "dados", "colunas_auditoria": False})
    bruto = _df(a=[1, 2, 3])
    ok = Relatorio("p"); ok.info("Linhas duplicadas removidas", tabela="dados", quantidade=1)
    t_ok = SimpleNamespace(brutas={"S": bruto}, analisaveis={"dados": bruto.iloc[1:]}, tabelas={"dados": bruto.iloc[1:]}, modelo=None)
    ruim = Relatorio("p")
    t_ruim = SimpleNamespace(brutas={"S": bruto}, analisaveis={"dados": bruto.iloc[1:]}, tabelas={"dados": bruto.iloc[1:]}, modelo=None)
    conferencia_supervisor(t_ok, ok, cfg, "x")
    conferencia_supervisor(t_ruim, ruim, cfg, "x")
    return [("Aceita remoção com motivo registrado", not ok.tem_erros),
            ("Barra linha que sumiu sem explicação", ruim.tem_erros)]


def _prova_coordenadores():
    from types import SimpleNamespace

    from robo.producao import checklist_analise, checklist_engenharia
    from robo.relatorio import Relatorio
    equipe = SimpleNamespace(problemas={})
    liberado = checklist_engenharia(Relatorio("p"), [{"ok": True, "bloqueia": True}], equipe, False, None, [])["liberado"]
    negado = checklist_engenharia(Relatorio("p"), [{"ok": False, "bloqueia": True}], equipe, False, None, [])["liberado"]
    a = checklist_analise({"resumo": ["x"], "_indicadores": [{"n": 40, "margem": 5}],
                           "_destaques": [{"significativo": True, "n_grupo": 12}]}, 30)
    return {"coord_engenharia": [("Libera com checklist verde", liberado), ("Não libera com item vermelho", not negado)],
            "coord_analise": [("Aponta destaque com base pequena", not a["aprovado"]),
                              ("Lista a ressalva", any("base menor" in x for x in a["ressalvas"]))]}


def _prova_excel():
    from robo.excel import Entrega, _criterio
    df = _df(curso=["A", "A", "B*", "=C"], nota_nps=pd.array([10, 6, 9, 3], dtype="Int64"),
             loja=["L1", "L2", "L1", "L2"], valor_venda=[10.0, 20.0, 30.0, 40.0], quantidade=[1, 2, 3, 4])
    papeis = {"curso": "categoria", "nota_nps": "numero", "loja": "categoria", "valor_venda": "numero", "quantidade": "numero"}
    with tempfile.TemporaryDirectory() as d:
        e = Entrega(df, "p", "p.xlsx", papeis, {"nps": "nota_nps"}, {})
        info = e.gerar(Path(d) / "p.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(Path(d) / "p.xlsx")
        dados_igual = wb["Dados"]["A5"].data_type == "s"
        formulas = [c.value for c in wb["Indicadores"]["B"] if isinstance(c.value, str) and c.value.startswith("=")]
        trade = [c.value for c in wb["Trade"]["B"] if isinstance(c.value, str) and c.value.startswith("=SUMIFS")]
    esperado = {x[0]: x[2] for x in e.conferencias}
    return {
        "excel_supervisor": [("Monta a conferência de todos os números", info["conferencias"] >= 15),
                             ("Resumo com status da conferência", "Resumo" in info["abas"])],
        "excel_pesquisa": [("Escapa curingas do CONT.SES (* ? ~)", _criterio("B*") == '"B~*"'),
                           ("Base por segmento correta", esperado.get("[curso] base de A") == 2)],
        "excel_formulas": [("NPS em fórmula", any("B8-B9" in f for f in formulas)),
                           ("NPS esperado correto", abs(esperado["Indicador: NPS"] - 0.0) < 1e-9)],
        "excel_trade": [("Vendas por loja em SOMASES", len(trade) == 2),
                        ("Soma por loja correta", esperado.get("[loja] valor_venda de L2") == 60.0)],
        "excel_dinamicas": [("Texto com '=' fica como texto (sem injeção de fórmula)", dados_igual),
                            ("Abas organizadas", info["abas"][0] == "Resumo")],
    }


PROVAS_NATIVAS = {"supervisor": _prova_supervisor, "dba": _prova_dba,"recepcao": _prova_recepcao, "leitura": _prova_leitura, "limpeza": _prova_limpeza,
                  "validacao": _prova_validacao, "modelagem": _prova_modelagem, "banco": _prova_banco,
                  "publicacao": _prova_publicacao}


@lru_cache(maxsize=None)
def _provas_nativas_cache() -> dict:
    resultado = {}
    for agente, funcao in PROVAS_NATIVAS.items():
        try:
            resultado[agente] = [{"prova": n, "aprovado": bool(ok), "motivo": "" if ok else "resultado diferente do esperado"}
                                 for n, ok in funcao()]
        except Exception as e:
            resultado[agente] = [{"prova": "Execução das provas", "aprovado": False, "motivo": f"{type(e).__name__}: {e}"}]
    for grupo in (_prova_excel,):
        try:
            for agente, casos in grupo().items():
                resultado[agente] = [{"prova": n, "aprovado": bool(ok), "motivo": "" if ok else "resultado diferente do esperado"}
                                     for n, ok in casos]
        except Exception as e:
            for a in ("excel_supervisor", "excel_pesquisa", "excel_formulas", "excel_trade", "excel_dinamicas"):
                resultado[a] = [{"prova": "Execução das provas", "aprovado": False, "motivo": f"{type(e).__name__}: {e}"}]
    try:
        for agente, casos in _prova_coordenadores().items():
            resultado[agente] = [{"prova": n, "aprovado": bool(ok), "motivo": "" if ok else "resultado diferente do esperado"}
                                 for n, ok in casos]
    except Exception as e:
        for a in ("coord_engenharia", "coord_analise"):
            resultado[a] = [{"prova": "Execução das provas", "aprovado": False, "motivo": f"{type(e).__name__}: {e}"}]
    try:
        for agente, casos in _prova_analistas().items():
            resultado[agente] = [{"prova": n, "aprovado": bool(ok), "motivo": "" if ok else "resultado diferente do esperado"}
                                 for n, ok in casos]
    except Exception as e:
        for a in ("analise_perfil", "analise_indicadores", "analise_estatistica", "analise_monitoramento", "analise_relator"):
            resultado[a] = [{"prova": "Execução das provas", "aprovado": False, "motivo": f"{type(e).__name__}: {e}"}]
    # quadro de pessoal (gerentes, plenos, juniores, estagiários) e equipe do canal
    from robo import canal, quadro
    from robo import ti
    for grupo, ids in ((quadro.provas_quadro, [p["id"] for p in quadro.QUADRO]), (ti.provas_ti, list(ti.POR_ID)),
                       (canal.provas_canal, ["canal_nina", "canal_iara", "canal_helena", "ger_operacoes"])):
        try:
            for agente, casos in grupo().items():
                destino = agente.removesuffix("_canal")
                resultado.setdefault(destino, [])
                resultado[destino] = resultado[destino] + [
                    {"prova": n, "aprovado": bool(ok), "motivo": "" if ok else "resultado diferente do esperado"} for n, ok in casos]
        except Exception as e:
            for a in ids:
                resultado[a] = [{"prova": "Execução das provas", "aprovado": False, "motivo": f"{type(e).__name__}: {e}"}]
    return resultado


def processo_seletivo(agente: dict) -> dict:
    """Resultado do processo seletivo técnico de um agente (nativo ou personalizado)."""
    provas = list(_provas_nativas_cache().get(agente["id"], []))
    pendencias = []
    nivel = agente.get("nivel") or "pleno"
    for cfg in agente.get("habilidades") or []:
        if cfg.get("ativo", True) is False:
            continue
        sid = cfg.get("habilidade", "")
        casos = provas_da_habilidade(sid)
        provas += [{**c, "prova": f"[{sid}] {c['prova']}"} for c in casos]
        if not casos:
            pendencias.append(f"habilidade '{sid}' não tem prova técnica cadastrada")
        risco = risco_de(sid, cfg.get("parametros"))
        if not nivel_permite(nivel, risco):
            pendencias.append(f"nível {NOME_NIVEL.get(nivel, nivel)} não pode usar '{sid}' (risco {NOME_RISCO[risco]}); "
                              f"promova o agente ou passe a tarefa a um sênior")
    acertos = sum(p["aprovado"] for p in provas)
    reprovadas = [p for p in provas if not p["aprovado"]]
    bloqueios = [f"reprovado na prova: {p['prova']} ({p['motivo']})" for p in reprovadas] + \
                [x for x in pendencias if "não pode usar" in x]
    return {"total": len(provas), "acertos": acertos, "aprovado": not bloqueios, "bloqueios": bloqueios,
            "ressalvas": [x for x in pendencias if "não pode usar" not in x], "provas": provas}


# ----------------------------------------------------------------------------- fichas profissionais
FICHAS = {
    "recepcao": {"cargo": "Analista Sênior de Ingestão de Dados", "nivel": "senior", "area": "Engenharia de Dados",
                 "metodologias": ["Ingestão por verificação de estabilidade (arquivo só entra quando o OneDrive termina de sincronizar)",
                                  "Idempotência: um arquivo nunca é processado duas vezes ao mesmo tempo",
                                  "Detecção de arquivo em uso (Excel aberto) antes de mover"],
                 "responsabilidades": ["Vigiar Bases-Limpeza", "Organizar a fila de chegada", "Adiar sem perder o que não está pronto"],
                 "criterios": ["Zero leitura de arquivo incompleto", "Arquivo em uso nunca é movido"],
                 "limites": ["Não abre nem altera o conteúdo das planilhas"]},
    "leitura": {"cargo": "Engenheiro de Dados Sênior — Extração", "nivel": "senior", "area": "Engenharia de Dados",
                "metodologias": ["Detecção de codificação (UTF-8, Windows-1252, Latin-1) e separador", "Cabeçalho por densidade de preenchimento",
                                 "Leitura do valor calculado das fórmulas, com captura de #REF!, #DIV/0!, #N/D"],
                "responsabilidades": ["Ler xlsx, xlsm, xls, csv e txt", "Preservar o número da linha original para rastreabilidade"],
                "criterios": ["Todo erro aponta a linha exata do Excel"], "limites": ["Não interpreta tipos nem corrige valores"]},
    "limpeza": {"cargo": "Especialista em Qualidade de Dados", "nivel": "especialista", "area": "Qualidade de Dados",
                "metodologias": ["Dimensões de qualidade DAMA-DMBOK: completude, validade, consistência e unicidade",
                                 "Padronização de nomenclatura (snake_case, sem acento)", "Conversão numérica e de datas no padrão brasileiro",
                                 "Preservação de identificadores com zero à esquerda"],
                "responsabilidades": ["Padronizar colunas e textos", "Converter tipos com limiar de confiança de 95%",
                                      "Remover linhas vazias, de total e duplicatas exatas"],
                "criterios": ["Toda alteração fica registrada no relatório", "Nenhuma conversão abaixo do limiar de confiança"],
                "limites": ["Não inventa valores: imputação só com habilidade explícita e nível adequado"]},
    "validacao": {"cargo": "Auditor Sênior de Dados", "nivel": "senior", "area": "Governança e Privacidade",
                  "metodologias": ["Contratos de dados: colunas obrigatórias, chaves únicas, domínios e faixas",
                                   "Princípio tudo-ou-nada (fail fast)", "Rastreabilidade linha a linha"],
                  "responsabilidades": ["Aprovar ou barrar cada arquivo", "Explicar o que corrigir e onde"],
                  "criterios": ["Nenhum arquivo com erro chega a Bases-Tratadas ou ao banco"], "limites": ["Não corrige dados: só aprova ou barra"]},
    "modelagem": {"cargo": "Arquiteta de Dados Sênior — Modelagem Dimensional", "nivel": "senior", "area": "Engenharia de Dados",
                  "metodologias": ["Modelagem dimensional de Kimball (esquema estrela)", "Chaves substitutas estáveis entre cargas",
                                   "Dimensão calendário completa", "Boas práticas de Power BI (relação 1:*, filtro único)"],
                  "responsabilidades": ["Separar fato e dimensões", "Gerar dicionário de dados"],
                  "criterios": ["A fato se liga às dimensões sem perder nem duplicar linhas"],
                  "limites": ["Não cria medidas de negócio — elas ficam no DAX ou no SQL analítico"]},
    "banco": {"cargo": "Engenheiro de Dados Sênior — PostgreSQL e ELT", "nivel": "senior", "area": "Engenharia de Dados",
              "metodologias": ["Arquitetura medalhão (bruto, tratado, analítico)", "Transações ACID e cargas idempotentes",
                               "Carga em lote (COPY)", "Evolução de schema com alargamento seguro de tipos"],
              "responsabilidades": ["Carregar as camadas numa única transação", "Registrar lotes e ocorrências"],
              "criterios": ["Ou entra inteiro, ou nada entra", "Reprocessar não duplica"],
              "limites": ["Nunca apaga tabelas (DROP) nem estreita tipos"]},
    "publicacao": {"cargo": "Analista Sênior de Governança e Publicação", "nivel": "senior", "area": "Governança e Privacidade",
                   "metodologias": ["Gravação atômica (sem arquivo pela metade)", "Retenção e versionamento", "Trilha de auditoria (origem e data)"],
                   "responsabilidades": ["Publicar em Bases-Tratadas", "Arquivar originais e reprovados"],
                   "criterios": ["Power BI nunca enxerga arquivo incompleto"], "limites": ["Não publica nada com erro"]},
    "supervisor": {"cargo": "Supervisor Sênior de Planilhas e Qualidade de Produção", "nivel": "senior", "area": "Qualidade de Dados",
                   "metodologias": ["Princípio dos quatro olhos (quem faz não é quem aprova)", "Conciliação de linhas origem × destino",
                                    "Integridade referencial fato ↔ dimensão", "Checklist de pré-produção"],
                   "responsabilidades": ["Rondar o 1º andar conferindo o trabalho de todos", "Barrar o que não fecha antes de publicar"],
                   "criterios": ["Toda linha removida tem motivo registrado", "Nenhuma chave órfã"],
                   "limites": ["Não executa etapas: só confere"]},
    "dba": {"cargo": "DBA Sênior PostgreSQL", "nivel": "senior", "area": "Engenharia de Dados",
            "metodologias": ["Reconciliação pós-carga dentro da transação (antes do commit)", "Integridade referencial em SQL",
                             "Transações ACID, rollback total em divergência", "Monitoramento de tamanho e saúde das tabelas"],
            "responsabilidades": ["Acompanhar tudo o que acontece no banco", "Desfazer a carga que não bate"],
            "criterios": ["Contagem no banco = contagem esperada, por lote"], "limites": ["Nunca corrige dado dentro do banco à mão"]},
    "coord_engenharia": {"cargo": "Coordenadora de Engenharia de Dados", "nivel": "coordenador", "area": "Engenharia de Dados",
                         "metodologias": ["Gate de produção (definition of done)", "Gestão de qualidade por checklist",
                                          "Boas práticas de engenharia: idempotência, atomicidade, rastreabilidade, testes"],
                         "responsabilidades": ["Assinar a liberação para produção", "Garantir que a equipe está aprovada no RH"],
                         "criterios": ["Só libera com 100% do checklist verde"], "limites": ["Não pula etapas por pressa"]},
    "coord_analise": {"cargo": "Coordenador de Análise e Ciência de Dados", "nivel": "coordenador", "area": "Análise de Dados",
                      "metodologias": ["Revisão por pares de análises", "Boas práticas de inferência (base mínima, margem, comparações múltiplas)",
                                       "Reprodutibilidade: toda análise refeita a partir dos dados publicados"],
                      "responsabilidades": ["Revisar as análises do 2º andar", "Apontar ressalvas antes da apresentação"],
                      "criterios": ["Nenhum achado comunicado sem base e margem"], "limites": ["Não altera números: aponta ressalvas"]},
    "excel_supervisor": {"cargo": "Gerente da Sala Excel · ex-Microsoft", "nivel": "gerente", "area": "Storytelling e Negócio",
                         "metodologias": ["Excel de ponta a ponta: fórmulas, matrizes dinâmicas, Power Query, Power Pivot, tabelas dinâmicas",
                                          "Auditoria de planilhas (rastreabilidade célula a célula)", "Conferência independente de cada número"],
                         "responsabilidades": ["Coordenar a Sala Excel", "Responder dúvidas de Excel no canal", "Assinar a entrega",
                                               "Coordenar as equipes de limpeza: barrar o que o Excel corromperia "
                                               "(ID de 16+ dígitos como número, textos acima de 32.767 caracteres, "
                                               "abas acima de 1.048.576 linhas, datas antes de 1900)"],
                         "criterios": ["Entrega só vale com 100% das conferências OK"],
                         "limites": ["Não cola valores: tudo em fórmula auditável"]},
    "canal_helena": {"cargo": "Coordenadora de RH e Desenvolvimento da Equipe", "nivel": "coordenador",
                     "area": "Governança e Privacidade",
                     "metodologias": ["Seleção por competência comprovada (provas com resposta conhecida)",
                                      "Trilha de carreira estagiário → júnior → pleno → sênior → especialista"],
                     "responsabilidades": ["Conduz o processo seletivo técnico de toda contratação",
                                           "Conta na sprint quem está aprovado, em observação ou em treinamento",
                                           "Avalia propostas de contratação automática antes de valerem"],
                     "criterios": ["Ninguém trabalha sem passar nas provas", "Nível compatível com o risco da tarefa"],
                     "limites": ["Não contrata por indicação: só por prova", "Não promove sem evidência de trabalho correto"]},
    "canal_iara": {"cargo": "Especialista Sênior em Inovação e Evolução da Plataforma", "nivel": "senior",
                   "area": "Storytelling e Negócio",
                   "metodologias": ["Melhoria contínua (PDCA)", "Priorização por impacto × esforço"],
                   "responsabilidades": ["Propõe melhorias com o dado que as motivou e o primeiro passo",
                                         "Mantém o backlog de ideias da sprint"],
                   "criterios": ["Toda ideia tem motivo e primeiro passo", "No máximo 3 ideias por sprint (foco)"],
                   "limites": ["Não implementa sozinha: propõe e o dono decide", "Não propõe o que não tem dado para sustentar"]},
    "canal_nina": {"cargo": "Assistente Júnior do Canal", "nivel": "junior", "area": "Storytelling e Negócio",
                   "metodologias": ["Roteamento de intenções por palavras-chave (offline, sem IA externa)"],
                   "responsabilidades": ["Lê cada mensagem do dono e chama quem sabe responder",
                                         "Abre o canal com o status do dia"],
                   "criterios": ["Pergunta técnica vai ao especialista certo", "Sem certeza, diz que não entendeu"],
                   "limites": ["Não inventa resposta", "Não decide nada: encaminha"]},
    "excel_pesquisa": {"cargo": "Especialista em Pesquisa de Mercado — Tabulação", "nivel": "especialista",
                       "area": "Customer Insights e Pesquisa",
                       "metodologias": ["Plano tabular e tabulação cruzada", "Base mínima e sinalização de base pequena",
                                        "NPS por segmento (Bain)"],
                       "responsabilidades": ["Frequências e percentuais por segmento"], "criterios": ["Total do plano = base preenchida"],
                       "limites": ["Não pondera sem plano de ponderação definido"]},
    "excel_formulas": {"cargo": "Especialista em Fórmulas e Power Query", "nivel": "especialista", "area": "Engenharia de Dados",
                       "metodologias": ["CONT.SES, SOMASES, SE, SEERRO, RAIZ", "Fórmulas robustas a divisão por zero",
                                        "Margem de erro em fórmula"],
                       "responsabilidades": ["Indicadores em fórmula"], "criterios": ["Nenhum #DIV/0! ou #VALOR! na entrega"],
                       "limites": ["Não usa macros (VBA) em entregas"]},
    "excel_trade": {"cargo": "Especialista em Trade Marketing", "nivel": "especialista", "area": "Customer Insights e Pesquisa",
                    "metodologias": ["Sell-in e sell-out", "Share por loja, canal, região e SKU", "Preço médio e mix",
                                     "Curva ABC de produtos"],
                    "responsabilidades": ["Montar a visão comercial da base"], "criterios": ["Soma dos shares = 100%"],
                    "limites": ["Não mistura sell-in com sell-out na mesma conta"]},
    "excel_dinamicas": {"cargo": "Especialista em Tabelas Dinâmicas e Dashboards Excel", "nivel": "especialista",
                        "area": "Storytelling e Negócio",
                        "metodologias": ["Tabelas dinâmicas e segmentação", "Formatação executiva", "Segurança: sem injeção de fórmula"],
                        "responsabilidades": ["Organizar a entrega e o Resumo"], "criterios": ["Quem abre entende em 30 segundos"],
                        "limites": ["Não esconde dados: tudo rastreável na aba Dados"]},
    "analise_perfil": {"cargo": "Cientista de Dados Sênior — Análise Exploratória", "nivel": "senior", "area": "Análise de Dados",
                       "metodologias": ["Análise exploratória (Tukey)", "Estatística descritiva robusta (mediana, quartis)",
                                        "Classificação do papel de cada variável"],
                       "responsabilidades": ["Raio-X de todas as colunas"], "criterios": ["Coluna com menos de 80% de preenchimento é sinalizada"],
                       "limites": ["Descreve; não tira conclusões causais"]},
    "analise_indicadores": {"cargo": "Especialista em Customer Insights — Métricas de Experiência", "nivel": "especialista",
                            "area": "Customer Insights e Pesquisa",
                            "metodologias": ["NPS (Reichheld/Bain)", "CSAT top-2-box", "CES", "Erro amostral com 95% de confiança", "Base mínima de 30"],
                            "responsabilidades": ["Calcular os indicadores no geral e por segmento"],
                            "criterios": ["Todo indicador sai com n e margem de erro"], "limites": ["Não compara escalas diferentes"]},
    "analise_estatistica": {"cargo": "Estatístico Sênior", "nivel": "senior", "area": "Estatística",
                            "metodologias": ["Testes z para proporções e médias (Welch)", "Correção de Holm-Bonferroni para comparações múltiplas",
                                             "Correlação de Pearson"],
                            "responsabilidades": ["Separar diferença real de acaso", "Apontar relações entre variáveis"],
                            "criterios": ["Grupo abaixo da base mínima nunca vira destaque"],
                            "limites": ["Correlação não é causa: nunca afirma causalidade"]},
    "analise_monitoramento": {"cargo": "Engenheira Sênior de Observabilidade de Dados", "nivel": "senior", "area": "Qualidade de Dados",
                              "metodologias": ["Detecção de drift: volume, schema, nulos e distribuição", "Contratos entre ondas",
                                               "Comparação de indicadores com significância"],
                              "responsabilidades": ["Avisar o que mudou em relação às cargas anteriores"],
                              "criterios": ["Variação dentro da margem nunca é chamada de mudança"], "limites": ["Não bloqueia publicação: alerta"]},
    "analise_relator": {"cargo": "Consultor Sênior de Storytelling com Dados", "nivel": "senior", "area": "Storytelling e Negócio",
                        "metodologias": ["Princípio da Pirâmide (Minto)", "Linguagem executiva", "Transparência sobre incerteza"],
                        "responsabilidades": ["Escrever o resumo executivo de cada base"],
                        "criterios": ["Só comunica como achado o que é estatisticamente confiável"],
                        "limites": ["Não inventa contexto que não está nos dados"]},
}


def _poder_de_bloqueio(p: dict, esc: dict) -> str:
    if p["nivel"] == "estagiario" or not esc["bloqueia"] or p["funcao"] is None:
        return "Não bloqueia: relata para quem decide"
    if p["fase"] in ("engenharia", "publicacao"):
        return "Pode bloquear a publicação (o arquivo vai para Bases-Erros com o motivo)"
    return "Achado vira ressalva registrada na entrega e no canal (o dado já está publicado)"


def _nome_pessoa(pid: str) -> str:
    try:
        from robo.canal import pessoas
        from robo.quadro import POR_ID
        from robo.ti import POR_ID as TI
        if pid in TI:
            return f"{TI[pid]['nome']} ({TI[pid]['cargo'].split(' — ')[0]})"
        if pid in POR_ID:
            return f"{POR_ID[pid]['nome']} ({POR_ID[pid]['cargo'].split(' — ')[0]})"
        p = pessoas().get(pid)
        return f"{p['nome']} ({p['papel']})" if p else pid
    except Exception:
        return pid


def ficha(agente: dict) -> dict:
    """Ficha profissional + resultado do processo seletivo."""
    base = dict(FICHAS.get(agente["id"], {}))
    from robo import quadro, ti
    if not base and agente["id"] in ti.POR_ID:
        p = ti.POR_ID[agente["id"]]
        esc = quadro.ESCOPO_NIVEL[p["nivel"]]
        base = {"cargo": p["cargo"], "nivel": p["nivel"], "area": p["area"], "metodologias": p["metodologias"],
                "responsabilidades": p["escopo"], "criterios": p["criterios"],
                "limites": (p.get("nao_faz") or [quadro._frase(esc["nao_faz"])]) + [
                    f"Decide: {esc['decide']}", "Não bloqueia planilhas: alerta no canal #ti e entra na sprint"],
                "gestor": "Você (dono)" if ti.GESTORES[p["id"]] == "dono" else _nome_pessoa(ti.GESTORES[p["id"]])}
    if not base and agente["id"] in quadro.POR_ID:
        p = quadro.POR_ID[agente["id"]]
        esc = quadro.ESCOPO_NIVEL[p["nivel"]]
        base = {"cargo": p["cargo"], "nivel": p["nivel"], "area": p["area"], "metodologias": p["metodologias"],
                "responsabilidades": p["escopo"], "criterios": p["criterios"],
                "limites": p["nao_faz"] + [f"Decide: {esc['decide']}", _poder_de_bloqueio(p, esc)]}
    if not base:
        from robo import habilidades
        cats = sorted({habilidades.REGISTRO[s["habilidade"]].categoria for s in agente.get("habilidades") or [] if s.get("habilidade") in habilidades.REGISTRO})
        base = {"cargo": agente.get("cargo") or agente.get("papel") or "Agente", "nivel": agente.get("nivel") or "pleno",
                "area": agente.get("area") or "Qualidade de Dados",
                "metodologias": [f"Conhecimento de {c} (ver 'Conhecimento' de cada habilidade)" for c in cats],
                "responsabilidades": [agente.get("descricao")] if agente.get("descricao") else [],
                "criterios": ["Todas as provas técnicas das suas habilidades aprovadas"],
                "limites": [f"Nível {NOME_NIVEL.get(agente.get('nivel') or 'pleno')}: risco máximo "
                            f"{NOME_RISCO[RISCO_MAXIMO.get(agente.get('nivel') or 'pleno', 'medio')]}"]
                + (["Estagiário: trabalha em observação, não altera os dados publicados"] if agente.get("nivel") == "estagiario" else []),
                "mentor": agente.get("mentor")}
    gestor = base.get("gestor") or quadro.GESTORES_EXISTENTES.get(agente["id"]) or (quadro.POR_ID.get(agente["id"]) or {}).get("gestor")
    if gestor and not base.get("gestor"):
        base["gestor"] = "Você (dono)" if gestor == "dono" else _nome_pessoa(gestor)
    base["nivel_id"] = base.get("nivel")
    base["nivel"] = NOME_NIVEL.get(base.get("nivel"), base.get("nivel"))
    base["provas"] = processo_seletivo({**agente, "nivel": base["nivel_id"]})
    return base
