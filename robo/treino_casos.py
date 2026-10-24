"""Catálogo de casos de treino: uma COMPETÊNCIA por ferramenta de contrato que dá para corrigir com resposta conhecida.

Cada competência tem 5 níveis e quatro partes, todas código de verdade:
  gerar(nivel, semente, pasta, cfg)  → o caso: dados sintéticos (arquivo ou parâmetros) + GABARITO calculado pelo gerador,
                                       que conhece a verdade porque criou os dados limpos antes de sujá-los;
  executar(cfg, sandbox, pid, caso)  → o funcionário usa a ferramenta REAL do contrato dele (ferramentas.executar);
  corrigir(caso, resposta)           → (acertou, nota 0–1, tipo_erro, detalhe) — comparação objetiva com o gabarito;
  conferir(caso, cfg)                → conferência INDEPENDENTE do revisor da Sala de Casos: refaz o gabarito por outro
                                       caminho (laço em Python puro em vez de pandas, validadores próprios) e confirma
                                       que a sujeira do nível foi mesmo plantada. Caso com problema não vai para ninguém.

Os mesmos casos servem a todos que estão no mesmo nível: dá para comparar pessoas de forma justa e, no futuro, usar
os registros em Machine Learning (semente + versão do gerador reproduzem o caso exatamente).
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from typing import Callable

import numpy as np
import pandas as pd

from robo.limpeza import normalizar_nome_coluna, sem_acento

NIVEIS = {1: "básico", 2: "intermediário", 3: "avançado", 4: "difícil", 5: "desafio"}


@dataclass
class Competencia:
    id: str
    nome: str
    ferramenta: str
    designer: str
    versao: str
    niveis: dict
    gerar: Callable
    executar: Callable
    corrigir: Callable
    conferir: Callable
    requer_banco: bool = False
    usa_conhecimento_real: bool = False
    descricao: str = ""
    extra: dict = field(default_factory=dict)


# =============================================================================== utilidades
def _rng(semente: int, nivel: int, sal: str) -> np.random.Generator:
    h = int(hashlib.sha256(f"{sal}|{semente}|{nivel}".encode()).hexdigest()[:12], 16)
    return np.random.default_rng(h)


def conteudo_hash(caminho: Path | None, extra: dict) -> str:
    """Hash do CONTEÚDO (valores das células), não dos bytes: o .xlsx guarda a hora em que foi salvo."""
    h = hashlib.sha256(json.dumps(extra, sort_keys=True, ensure_ascii=False, default=str).encode())
    if caminho and Path(caminho).exists():
        p = Path(caminho)
        if p.suffix.lower() in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            wb = load_workbook(p, read_only=True, data_only=False)
            for ws in wb.worksheets:
                h.update(ws.title.encode())
                for linha in ws.iter_rows(values_only=True):
                    h.update(repr(linha).encode())
            wb.close()
        else:
            h.update(p.read_bytes())
    return h.hexdigest()


def _xlsx(caminho: Path, cabecalho: list, linhas: list, titulo: str | None = None) -> Path:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    if titulo:
        ws.append([titulo])
        ws.append([])
    ws.append(cabecalho)
    for l in linhas:
        ws.append(l)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho


def _usar(ferramenta: str, cfg, pid: str, **params) -> dict:
    from robo import ferramentas
    return ferramentas.executar(ferramenta, cfg, quem=pid, **params)


def _fechado(a: float, b: float, tol: float = 1e-6) -> bool:
    return a is not None and b is not None and not (isinstance(a, float) and math.isnan(a)) and \
        math.isclose(float(a), float(b), rel_tol=tol, abs_tol=tol)


def _excecao(resposta: dict):
    if resposta.get("excecao"):
        return (False, 0.0, "excecao_no_codigo", f"a ferramenta quebrou: {resposta['texto'][:300]}")
    return None


# =============================================================================== 1. limpeza (diagnosticar_planilha)
def _g_limpeza(nivel, semente, pasta, cfg):
    from robo import laboratorio
    from openpyxl import load_workbook
    area = laboratorio.AREAS[semente % len(laboratorio.AREAS)]
    caso = laboratorio.aplicar_nivel(laboratorio.gerar(area, semente), nivel, semente)
    arq = pasta / f"treino_{area}_{semente}.xlsx"
    laboratorio.escrever_sujo(caso, arq)
    rng = _rng(semente, nivel, "limpeza")
    defeito = nivel >= 2 and rng.random() < 0.3  # no nível 5 o desafio é a sujeira de sistema antigo, não só barrar
    celulas = []
    if defeito:  # fórmula quebrada no meio dos dados: a planilha TEM que ser barrada
        wb = load_workbook(arq)
        ws = wb.active
        cab = next(i for i, r in enumerate(ws.iter_rows(values_only=True), 1) if r and r[0] and str(r[0]).strip().upper() == "ID")
        alvo_col = ws.max_column if ws.cell(cab, ws.max_column).value else ws.max_column - 1
        for k in rng.choice(np.arange(cab + 2, cab + 2 + min(40, caso["n"] - 2)), 2, replace=False):
            ws.cell(int(k), alvo_col, "#REF!")
            celulas.append(int(k))
        wb.save(arq)
    return {"arquivo": str(arq), "parametros": {"area": area, "nivel_sujeira": nivel, "defeito": defeito},
            "gabarito": {"aprovaria": not defeito, "linhas": caso["n"], "celulas_ref": celulas},
            "resumo": f"planilha de {area} com sujeira nível {nivel}" + (" e fórmula quebrada (#REF!)" if defeito else "")}


def _x_limpeza(cfg, sb, pid, caso):
    return _usar("diagnosticar_planilha", sb, pid, arquivo=caso["arquivo"])


def _c_limpeza(caso, r):
    if (e := _excecao(r)):
        return e
    g, d = caso["gabarito"], r.get("dados") or {}
    if g["aprovaria"] and not d.get("aprovaria"):
        return (False, 0.0, "barrou_planilha_boa", "barrou uma planilha que deveria passar: " + "; ".join((d.get("erros") or [])[:2]))
    if not g["aprovaria"]:
        achou = any("fórmula" in str(x).lower() or "#ref" in str(x).lower() for x in d.get("erros") or [])
        return (True, 1.0, "", "barrou a fórmula quebrada") if not d.get("aprovaria") and achou else \
            (False, 0.0, "aprovou_planilha_com_defeito", "deixou passar célula #REF!")
    if d.get("linhas_saida") != g["linhas"]:
        return (False, 0.5, "linhas_erradas", f"{d.get('linhas_saida')} linha(s) na saída; o certo é {g['linhas']} "
                                               "(Total, subtotal, rodapé ou duplicata mal tratados)")
    return (True, 1.0, "", f"{g['linhas']} linhas certas, sem falso alarme")


def _v_limpeza(caso, cfg):
    """Revisor: relê o arquivo cru (openpyxl, sem o pipeline) e confirma a verdade e a sujeira plantada."""
    from openpyxl import load_workbook
    problemas = []
    wb = load_workbook(caso["arquivo"], read_only=True)
    linhas = list(wb.active.iter_rows(values_only=True))
    wb.close()
    ids = set()
    for l in linhas:
        if l and isinstance(l[0], int) and 10_000 <= l[0] < 20_000:
            ids.add(l[0])
    if len(ids) != caso["gabarito"]["linhas"]:
        problemas.append(f"o arquivo tem {len(ids)} IDs distintos; o gabarito diz {caso['gabarito']['linhas']}")
    texto = " ".join(str(c) for l in linhas for c in (l or ()) if isinstance(c, str))
    if "Total" not in texto:
        problemas.append("linha de Total não foi plantada")
    nivel = caso["parametros"]["nivel_sujeira"]
    if nivel >= 2 and "Subtotal" not in texto:
        problemas.append("nível ≥ 2 sem subtotal plantado")
    if nivel >= 3 and "Gerado em" not in texto:
        problemas.append("nível ≥ 3 sem rodapé de sistema")
    refs = sum(1 for l in linhas for c in (l or ()) if c == "#REF!")
    if caso["parametros"]["defeito"] != bool(refs):
        problemas.append(f"defeito plantado={caso['parametros']['defeito']} mas o arquivo tem {refs} #REF!")
    if caso["gabarito"]["aprovaria"] != (refs == 0):
        problemas.append(f"gabarito diz aprovaria={caso['gabarito']['aprovaria']}, mas o arquivo tem {refs} célula(s) #REF!")
    linhas_ref = sorted(i for i, l in enumerate(linhas, 1) if any(c == "#REF!" for c in (l or ())))
    if linhas_ref != sorted(caso["gabarito"]["celulas_ref"]):
        problemas.append(f"#REF! no arquivo nas linhas {linhas_ref}; o gabarito diz {sorted(caso['gabarito']['celulas_ref'])}")
    return problemas


# =============================================================================== 2. cálculo (calcular)
_SEGUNDA = {"varejo": "Canal de Venda", "educacao_nps": "Modalidade", "financeiro": "Conta Contábil"}


def _g_calculo(nivel, semente, pasta, cfg):
    from robo import laboratorio
    area = laboratorio.AREAS[(semente // 7) % len(laboratorio.AREAS)]
    caso = laboratorio.aplicar_nivel(laboratorio.gerar(area, semente), nivel, semente)
    arq = pasta / f"treino_calc_{area}_{semente}.xlsx"
    laboratorio.escrever_sujo(caso, arq)
    verdade = laboratorio.verdade(caso)
    seg, met = normalizar_nome_coluna(caso["segmento"]), normalizar_nome_coluna(caso["metrica"])
    op = caso["op"]
    nome_op = {"nps": "nps", "media": "média", "soma": "soma", "mediana": "mediana"}
    rng = _rng(semente, nivel, "calculo")
    por, filtros = [seg], []
    if nivel == 2:
        valor = str(rng.choice(sorted(verdade[seg].dropna().unique())))
        por, filtros = [], [(seg, "=", valor)]
    elif nivel == 3 and op != "nps":
        op = "mediana"
    elif nivel == 5 and area in _SEGUNDA:
        por = [seg, normalizar_nome_coluna(_SEGUNDA[area])]
    pergunta = (f"{nome_op[op]} de {met}" if op != "nps" else "nps") + (f" por {' e '.join(por)}" if por else "") + \
        "".join(f" onde {c} = {v}" for c, s, v in filtros)
    dados = verdade
    for c, s, v in filtros:
        dados = dados[dados[c].astype(str) == v]
    from robo import conferencia
    esperado = conferencia.valores_pandas(op, dados, met, por)
    return {"arquivo": str(arq), "parametros": {"area": area, "pergunta": pergunta, "op": op, "por": por, "filtros": filtros,
                                                "metrica": met, "nivel_sujeira": nivel},
            "gabarito": {"valores": {k: (None if isinstance(v, float) and math.isnan(v) else v) for k, v in esperado.items()}},
            "resumo": f"“{pergunta}” numa planilha de {area} com sujeira nível {nivel}"}


def _x_calculo(cfg, sb, pid, caso):
    return _usar("calcular", sb, pid, pergunta=caso["parametros"]["pergunta"], arquivo=caso["arquivo"])


def _c_calculo(caso, r):
    if (e := _excecao(r)):
        return e
    calc = (r.get("dados") or {}).get("calculo")
    if not calc:
        return (False, 0.0, "nao_entendeu_o_pedido", r.get("texto", "")[:200])
    p = caso["parametros"]
    if calc["op"] != p["op"]:
        return (False, 0.0, "operacao_errada", f"fez {calc['op']}; o pedido era {p['op']}")
    g, obt = caso["gabarito"]["valores"], calc.get("valores") or {}
    erradas = [k for k in g if g[k] is not None and not _fechado(obt.get(k), g[k], 1e-6)]
    sobrando = [k for k in obt if k not in g]
    if erradas or sobrando:
        nota = 1 - len(erradas) / max(1, len(g))
        det = "; ".join(f"{k}: deu {obt.get(k)}, certo {g[k]}" for k in erradas[:3]) + \
            (f"; grupos que não existem: {sobrando[:3]}" if sobrando else "")
        return (False, round(max(0.0, nota), 2), "numero_errado", det)
    return (True, 1.0, "", f"{len(g)} valor(es) certos")


def _v_calculo(caso, cfg):
    """Revisor: refaz a conta num laço em Python puro (outro caminho que não o pandas) a partir da verdade."""
    from robo import laboratorio
    p = caso["parametros"]
    semente = caso["semente"]
    base = laboratorio.aplicar_nivel(laboratorio.gerar(p["area"], semente), p["nivel_sujeira"], semente)
    rot = {normalizar_nome_coluna(k): list(v) for k, v in base["rotulos"].items()}
    grupos: dict = {}
    for i in range(base["n"]):
        if any(str(rot[c][i]) != v for c, s, v in p["filtros"]):
            continue
        x = rot[p["metrica"]][i]
        if x is None or (isinstance(x, float) and math.isnan(x)):
            continue
        chave = " / ".join(str(rot[c][i]) for c in p["por"]) if p["por"] else "(total)"
        grupos.setdefault(chave, []).append(float(x))
    problemas = []
    for chave, xs in grupos.items():
        if p["op"] == "soma":
            v = math.fsum(xs)
        elif p["op"] == "media":
            v = math.fsum(xs) / len(xs)
        elif p["op"] == "mediana":
            s = sorted(xs)
            v = s[len(s) // 2] if len(s) % 2 else (s[len(s) // 2 - 1] + s[len(s) // 2]) / 2
        else:  # nps: só notas inteiras de 0 a 10
            ok = [x for x in xs if float(x).is_integer() and 0 <= x <= 10]
            v = 100 * (sum(x >= 9 for x in ok) - sum(x <= 6 for x in ok)) / len(ok) if ok else float("nan")
        g = caso["gabarito"]["valores"].get(chave)
        if g is None or not _fechado(v, g, 1e-9):
            problemas.append(f"{chave}: gabarito {g}, conta independente {v}")
    if set(grupos) != set(caso["gabarito"]["valores"]):
        problemas.append("grupos do gabarito diferentes dos da conta independente")
    return problemas


# =============================================================================== 3. indicadores (perfil_indicadores)
def _g_indicadores(nivel, semente, pasta, cfg):
    from robo import laboratorio
    caso = laboratorio.aplicar_nivel(laboratorio.gerar("educacao_nps", semente), nivel, semente)
    rng = _rng(semente, nivel, "indicadores")
    nota = caso["rotulos"]["Nota NPS"] = np.array(caso["rotulos"]["Nota NPS"], dtype=object)
    codigos = []
    if nivel >= 3:  # "99 = não sabe": não é nota, tem que ficar fora do NPS
        livres = [i for i in range(caso["n"]) if nota[i] is not None]
        for i in rng.choice(livres, max(3, caso["n"] // 25), replace=False):
            nota[int(i)] = 99
            codigos.append(int(i))
    arq = pasta / f"treino_nps_{semente}.xlsx"
    laboratorio.escrever_sujo(caso, arq)
    validas = [int(x) for x in nota if x is not None and 0 <= int(x) <= 10]
    nps = 100 * (sum(x >= 9 for x in validas) - sum(x <= 6 for x in validas)) / len(validas)
    return {"arquivo": str(arq), "parametros": {"nivel_sujeira": nivel, "codigos_99": codigos},
            "gabarito": {"nps": nps, "n": len(validas)},
            "resumo": f"pesquisa NPS com sujeira nível {nivel}" + (f" e {len(codigos)} código(s) 99" if codigos else "")}


def _x_indicadores(cfg, sb, pid, caso):
    return _usar("perfil_indicadores", sb, pid, arquivo=caso["arquivo"])


def _c_indicadores(caso, r):
    if (e := _excecao(r)):
        return e
    g = caso["gabarito"]
    linhas = [l for tab in (r.get("dados") or {}).values() if isinstance(tab, dict) for l in tab.get("indicadores") or []]
    geral = next((l for l in linhas if l.get("indicador") == "NPS" and l.get("segmento") == "Total"), None)
    if not geral:
        return (False, 0.0, "nao_achou_o_indicador", "não reconheceu a coluna de NPS")
    if int(geral.get("n") or 0) != g["n"]:
        return (False, 0.3, "base_errada", f"n={geral.get('n')}; o certo é {g['n']} (códigos 99 e vazios ficam fora)")
    if abs(float(geral["valor"]) - g["nps"]) > 0.051:
        return (False, 0.5, "numero_errado", f"NPS {geral['valor']}; o certo é {g['nps']:.1f}")
    return (True, 1.0, "", f"NPS {geral['valor']} com n={g['n']}")


def _v_indicadores(caso, cfg):
    from openpyxl import load_workbook
    wb = load_workbook(caso["arquivo"], read_only=True)
    linhas = list(wb.active.iter_rows(values_only=True))
    wb.close()
    cab_i = next(i for i, l in enumerate(linhas) if l and l[0] and str(l[0]).strip().upper() == "ID")
    cab = [str(c or "").strip().lower() for c in linhas[cab_i]]
    j = cab.index("nota nps")
    vistos, notas = set(), []
    for l in linhas[cab_i + 1:]:
        if not l or not isinstance(l[0], int) or l[0] in vistos:
            continue
        vistos.add(l[0])
        v = l[j]
        try:
            x = float(str(v).replace(",", "")) if v not in (None, "-", "N/A", "") else None
        except ValueError:
            x = None
        if x is not None and x.is_integer() and 0 <= x <= 10:
            notas.append(x)
    problemas = []
    if len(notas) != caso["gabarito"]["n"]:
        problemas.append(f"releitura achou {len(notas)} notas válidas; gabarito diz {caso['gabarito']['n']}")
    elif notas:
        nps = 100 * (sum(x >= 9 for x in notas) - sum(x <= 6 for x in notas)) / len(notas)
        if abs(nps - caso["gabarito"]["nps"]) > 1e-6:
            problemas.append(f"NPS relido {nps}; gabarito {caso['gabarito']['nps']}")
    return problemas


# =============================================================================== 4. LGPD (dados_pessoais)
def _cpf(rng, valido=True) -> str:
    n = [int(x) for x in rng.integers(0, 10, 9)]
    for k in (10, 11):
        d = (sum(v * w for v, w in zip(n, range(k, 1, -1))) * 10) % 11
        n.append(0 if d == 10 else d)
    if not valido:
        n[-1] = (n[-1] + 1 + int(rng.integers(0, 8))) % 10
    return "".join(map(str, n))


def cpf_valido(c: str) -> bool:
    c = re.sub(r"\D", "", str(c))
    if len(c) != 11 or len(set(c)) == 1:
        return False
    n = [int(x) for x in c]
    for k in (9, 10):
        d = (sum(v * w for v, w in zip(n[:k], range(k + 1, 1, -1))) * 10) % 11
        if (0 if d == 10 else d) != n[k]:
            return False
    return True


NOMES = ["ana", "bruno", "carla", "diego", "elisa", "fabio", "gabi", "hugo", "iris", "joao", "katia", "luis"]
SOBRENOMES = ["silva", "souza", "lima", "costa", "rocha", "alves", "melo", "dias"]


def _g_lgpd(nivel, semente, pasta, cfg):
    rng = _rng(semente, nivel, "lgpd")
    n = 120
    cpfs = [_cpf(rng) for _ in range(n)]
    emails = [f"{rng.choice(NOMES)}.{rng.choice(SOBRENOMES)}{int(rng.integers(1, 99))}@{rng.choice(['gmail.com', 'empresa.com.br', 'outlook.com'])}"
              for _ in range(n)]
    tels = [f"({int(rng.integers(11, 99))}) 9{int(rng.integers(1000, 9999))}-{int(rng.integers(1000, 9999))}" for _ in range(n)]
    nomes_col = {1: ("email", "cpf", "telefone"), 2: ("contato", "documento", "fone_cliente")}.get(nivel, ("campo_a", "campo_b", "campo_c"))
    if nivel >= 3:  # formatos crus: CPF só com dígitos, telefone com +55
        cpfs_txt = cpfs
        tels = [f"+55 {t}" for t in tels]
    else:
        cpfs_txt = [f"{c[:3]}.{c[3:6]}.{c[6:9]}-{c[9:]}" for c in cpfs]
    colunas = {"id": list(range(1, n + 1)), "valor_compra": [round(float(x), 2) for x in rng.lognormal(5, .5, n)],
               "codigo_produto": [f"PRD-{int(x):05d}" for x in rng.integers(0, 99999, n)],
               nomes_col[0]: emails, nomes_col[1]: cpfs_txt, nomes_col[2]: tels}
    pessoais = list(nomes_col)
    iscas = []
    if nivel >= 4:  # isca: número de pedido COM a cara de CPF, mas dígito verificador inválido — não é dado pessoal
        colunas["numero_pedido"] = [(lambda c: f"{c[:3]}.{c[3:6]}.{c[6:9]}-{c[9:]}")(_cpf(rng, valido=False)) for _ in range(n)]
        iscas.append("numero_pedido")
    if nivel >= 5:  # e-mail escondido no meio do texto livre: também é dado pessoal
        colunas["observacao"] = [f"cliente pediu retorno em {e}" if rng.random() < 0.8 else "sem observação" for e in emails]
        pessoais.append("observacao")
    cab = list(colunas)
    linhas = [[colunas[c][i] for c in cab] for i in range(n)]
    arq = _xlsx(pasta / f"treino_lgpd_{semente}.xlsx", cab, linhas)
    return {"arquivo": str(arq), "parametros": {"colunas_pessoais": pessoais, "iscas": iscas},
            "gabarito": {"pessoais": sorted(normalizar_nome_coluna(c) for c in pessoais)},
            "resumo": f"cadastro com {len(pessoais)} coluna(s) de dado pessoal" + (f" e {len(iscas)} isca(s)" if iscas else "")}


def _x_lgpd(cfg, sb, pid, caso):
    return _usar("dados_pessoais", sb, pid, arquivo=caso["arquivo"])


def _c_lgpd(caso, r):
    if (e := _excecao(r)):
        return e
    achados = (r.get("dados") or {}).get("achados") or []
    obtidas = sorted({m.group(1) for a in achados for m in [re.search(r"\.([^\].]+)\] parece", a)] if m})
    certas = caso["gabarito"]["pessoais"]
    faltou, sobrou = sorted(set(certas) - set(obtidas)), sorted(set(obtidas) - set(certas))
    if faltou or sobrou:
        tipo = "deixou_dado_pessoal_passar" if faltou else "falso_alarme"
        nota = len(set(certas) & set(obtidas)) / max(1, len(set(certas) | set(obtidas)))
        return (False, round(nota, 2), tipo, (f"não apontou {faltou}" if faltou else "") + ("; " if faltou and sobrou else "")
                + (f"apontou sem ser dado pessoal: {sobrou}" if sobrou else ""))
    return (True, 1.0, "", f"apontou exatamente {certas}")


def _v_lgpd(caso, cfg):
    from openpyxl import load_workbook
    wb = load_workbook(caso["arquivo"], read_only=True)
    linhas = list(wb.active.iter_rows(values_only=True))
    wb.close()
    cab = [str(c) for c in linhas[0]]
    col = {c: [l[i] for l in linhas[1:]] for i, c in enumerate(cab)}
    problemas = []
    for c in caso["parametros"]["colunas_pessoais"]:
        vals = [str(v) for v in col.get(c, [])]
        tem = sum(bool(re.search(r"[\w.]+@[\w.]+\.\w{2,}", v)) or cpf_valido(v) or
                  len(re.sub(r"\D", "", v)) in (10, 11, 12, 13) and "(" in v for v in vals)
        if tem < 0.6 * len(vals):
            problemas.append(f"coluna {c} deveria ter dado pessoal em ≥60% das linhas (tem {tem}/{len(vals)})")
    for c in caso["parametros"]["iscas"]:
        if any(cpf_valido(v) for v in col.get(c, [])):
            problemas.append(f"isca {c} tem CPF válido (deveria ser só parecido)")
    # gabarito refeito pelo CONTEÚDO de cada coluna (e-mail, CPF com dígito válido, telefone com DDD) — outro caminho
    def pessoal(v) -> bool:
        v = str(v)
        return bool(re.search(r"[\w.]+@[\w.]+\.\w{2,}", v)) or cpf_valido(v) or \
            (len(re.sub(r"\D", "", v)) in (10, 11, 12, 13) and "(" in v)
    refeito = sorted(normalizar_nome_coluna(c) for c, vals in col.items()
                     if vals and sum(pessoal(v) for v in vals) >= 0.6 * len(vals))
    if refeito != sorted(caso["gabarito"]["pessoais"]):
        problemas.append(f"pelo conteúdo, as colunas pessoais são {refeito}; o gabarito diz {sorted(caso['gabarito']['pessoais'])}")
    return problemas


# =============================================================================== 5. anomalias
def _tukey(xs: list[float]) -> tuple[float, float]:
    s = sorted(xs)

    def q(p):  # quantil linear (mesmo método padrão do pandas), calculado à mão
        pos = (len(s) - 1) * p
        lo = int(math.floor(pos))
        return s[lo] + (s[min(lo + 1, len(s) - 1)] - s[lo]) * (pos - lo)
    q1, q3 = q(.25), q(.75)
    return q1 - 3 * (q3 - q1), q3 + 3 * (q3 - q1)


def _g_anomalias(nivel, semente, pasta, cfg):
    rng = _rng(semente, nivel, "anomalias")
    n = 240
    ncols = 1 if nivel == 1 else 2 if nivel <= 3 else 3
    colunas = {"id": list(range(1, n + 1))}
    nomes = ["tempo_atendimento", "valor_pedido", "dias_entrega"][:ncols]
    for c in nomes:
        base = rng.lognormal(3, 0.35, n) if nivel >= 3 else rng.normal(100, 10, n)
        k = 1 if nivel == 1 else int(rng.integers(2, 4))
        idx = rng.choice(n, k, replace=False)
        for j, i in enumerate(idx):
            sinal = -1 if (nivel >= 5 and j % 2) else 1
            base[i] = float(np.median(base)) + sinal * 40 * float(np.std(base))
        colunas[c] = [round(float(x), 2) for x in base]
    verdade = {}
    for c in nomes:
        lo, hi = _tukey(colunas[c])
        verdade[c] = sum(1 for x in colunas[c] if x < lo or x > hi)
    escrito = {k: list(v) for k, v in colunas.items()}
    if nivel >= 4:  # ruído de texto: vazios escritos como "N/A" e números em formato brasileiro
        for c in nomes:
            for i in rng.choice(n, 6, replace=False):
                escrito[c][int(i)] = "N/A"
                colunas[c][int(i)] = None
            escrito[c] = [f"{v:.2f}".replace(".", ",") if isinstance(v, float) else v for v in escrito[c]]
        for c in nomes:  # verdade recalculada sem os vazios
            xs = [x for x in colunas[c] if x is not None]
            lo, hi = _tukey(xs)
            verdade[c] = sum(1 for x in xs if x < lo or x > hi)
    cab = list(escrito)
    arq = _xlsx(pasta / f"treino_anomalias_{semente}.xlsx", cab, [[escrito[c][i] for c in cab] for i in range(n)])
    return {"arquivo": str(arq), "parametros": {"colunas": nomes},
            "gabarito": {"por_coluna": {c: v for c, v in verdade.items() if v}},
            "resumo": f"{ncols} coluna(s) numérica(s) com {sum(verdade.values())} valor(es) fora da faixa de Tukey (3×IQR)"}


def _x_anomalias(cfg, sb, pid, caso):
    return _usar("anomalias", sb, pid, arquivo=caso["arquivo"])


def _c_anomalias(caso, r):
    if (e := _excecao(r)):
        return e
    achados = ((r.get("dados") or {}).get("anomalias") or {}).get("achados") or []
    obt = {a["coluna"]: a["quantidade"] for a in achados if a["coluna"] in caso["parametros"]["colunas"]}
    g = caso["gabarito"]["por_coluna"]
    if obt != g:
        return (False, 0.0, "anomalia_nao_achada" if sum(obt.values()) < sum(g.values()) else "falso_alarme",
                f"achou {obt}; o certo é {g}")
    return (True, 1.0, "", f"achou {sum(g.values())} anomalia(s) certas")


def _v_anomalias(caso, cfg):
    """Revisor: relê a planilha crua (N/A e vírgula decimal tratados à mão) e reconta pela regra de Tukey (3×IQR)."""
    from openpyxl import load_workbook
    if not caso["gabarito"]["por_coluna"]:
        return ["caso sem nenhuma anomalia: não treina nada"]
    wb = load_workbook(caso["arquivo"], read_only=True)
    linhas = list(wb.active.iter_rows(values_only=True))
    wb.close()
    cab = [str(c) for c in linhas[0]]
    refeito = {}
    for c in caso["parametros"]["colunas"]:
        xs = []
        for l in linhas[1:]:
            v = l[cab.index(c)]
            if v is None or str(v).strip().upper() == "N/A":
                continue
            xs.append(float(str(v).replace(",", ".")) if isinstance(v, str) else float(v))
        lo, hi = _tukey(xs)
        n = sum(1 for x in xs if x < lo or x > hi)
        if n:
            refeito[c] = n
    if refeito != caso["gabarito"]["por_coluna"]:
        return [f"recontando no arquivo: {refeito}; o gabarito diz {caso['gabarito']['por_coluna']}"]
    return []


# =============================================================================== 6. segmentação
def _g_segmentacao(nivel, semente, pasta, cfg):
    rng = _rng(semente, nivel, "segmentacao")
    k = {1: 2, 2: 3, 3: 3, 4: 4, 5: 4}[nivel]
    dims = ["renda_mensal", "idade", "gasto_mensal", "visitas_mes"][: (2 if nivel <= 2 else 3)]
    sep = {1: 9.0, 2: 8.0, 3: 7.0, 4: 7.0, 5: 6.5}[nivel]
    # centros equidistantes (polígono regular / tetraedro): nenhum par de grupos mais perto que outro — gabarito sem
    # ambiguidade (antes, dois centros sorteados perto um do outro faziam '3 grupos' parecer 2)
    if len(dims) == 2:
        ang = rng.uniform(0, 2 * np.pi) + np.arange(k) * 2 * np.pi / k
        centros = np.c_[np.cos(ang), np.sin(ang)] * sep / (2 * np.sin(np.pi / k))
    else:
        base = np.array([[1, 1, 1], [1, -1, -1], [-1, 1, -1], [-1, -1, 1]], dtype=float)[:k]
        rot, _ = np.linalg.qr(rng.normal(0, 1, (3, 3)))
        centros = base @ rot * sep / (2 * np.sqrt(2))
    n_por = 90
    X = np.vstack([rng.normal(centros[g], 1.0, (n_por, len(dims))) for g in range(k)])
    rotulos = np.repeat(np.arange(k), n_por)
    escalas = {"renda_mensal": (3000, 400), "idade": (40, 3), "gasto_mensal": (800, 90), "visitas_mes": (6, 0.8)}
    colunas = {"id": list(range(1, len(X) + 1))}
    for j, d in enumerate(dims):
        m, s = escalas[d]
        colunas[d] = [round(float(m + s * x), 2) for x in X[:, j]]
    if nivel >= 3:  # coluna sem estrutura nenhuma (ruído): não pode inventar grupo a mais por causa dela
        colunas["nota_aleatoria"] = [round(float(x), 2) for x in rng.normal(5, 0.05, len(X))]
    cab = list(colunas)
    arq = _xlsx(pasta / f"treino_seg_{semente}.xlsx", cab, [[colunas[c][i] for c in cab] for i in range(len(X))])
    return {"arquivo": str(arq), "parametros": {"dims": dims, "rotulos": rotulos.tolist()},
            "gabarito": {"k": k}, "resumo": f"{len(X)} clientes em {k} perfis plantados ({', '.join(dims)})"}


def _x_segmentacao(cfg, sb, pid, caso):
    return _usar("segmentacao", sb, pid, arquivo=caso["arquivo"])


def _c_segmentacao(caso, r):
    if (e := _excecao(r)):
        return e
    sg = (r.get("dados") or {}).get("segmentos") or {}
    if not sg.get("ok"):
        return (False, 0.0, "nao_segmentou", r.get("texto", "")[:200])
    if sg["k"] != caso["gabarito"]["k"]:
        return (False, 0.0, "grupos_errados", f"achou {sg['k']} grupo(s); foram plantados {caso['gabarito']['k']}")
    return (True, 1.0, "", f"{sg['k']} grupos, silhueta {sg.get('silhueta', 0):.2f}")


def _v_segmentacao(caso, cfg):
    """Revisor: silhueta dos rótulos VERDADEIROS (conta própria) — grupo mal definido não serve de gabarito."""
    from openpyxl import load_workbook
    wb = load_workbook(caso["arquivo"], read_only=True)
    linhas = list(wb.active.iter_rows(values_only=True))
    wb.close()
    cab = list(linhas[0])
    idx = [cab.index(d) for d in caso["parametros"]["dims"]]
    X = np.array([[float(l[i]) for i in idx] for l in linhas[1:]])
    X = (X - X.mean(0)) / X.std(0)
    rot = np.array(caso["parametros"]["rotulos"])
    D = np.sqrt(((X[:, None, :] - X[None]) ** 2).sum(-1))
    s = []
    for i in range(len(X)):
        a = D[i, rot == rot[i]].sum() / max(1, (rot == rot[i]).sum() - 1)
        b = min(D[i, rot == j].mean() for j in set(rot.tolist()) if j != rot[i])
        s.append((b - a) / max(a, b))
    sil = float(np.mean(s))
    problemas = [] if sil >= 0.5 else [f"grupos plantados mal separados (silhueta {sil:.2f} < 0,50)"]
    if len(set(rot.tolist())) != caso["gabarito"]["k"]:
        problemas.append(f"foram plantados {len(set(rot.tolist()))} grupo(s); o gabarito diz {caso['gabarito']['k']}")
    return problemas


# =============================================================================== 7. drivers
ATRIBUTOS = ["nota_atendimento", "nota_preco", "nota_entrega", "nota_app", "nota_variedade"]


def _g_drivers(nivel, semente, pasta, cfg):
    rng = _rng(semente, nivel, "drivers")
    n = 320
    attrs = ATRIBUTOS[: (3 if nivel <= 2 else 5)]
    driver = str(rng.choice(attrs))
    Z = {a: rng.normal(0, 1, n) for a in attrs}
    if nivel >= 4:  # distrator parecido com o driver (correlação ~0,6): só olhar um par engana
        outro = next(a for a in attrs if a != driver)
        Z[outro] = 0.6 * Z[driver] + 0.8 * rng.normal(0, 1, n)
    ruido = {1: 0.4, 2: 0.6, 3: 0.8, 4: 0.9, 5: 1.0}[nivel]
    latente = 0.9 * Z[driver] + ruido * rng.normal(0, 1, n)
    colunas = {"id": list(range(1, n + 1))}
    for a in attrs:
        colunas[a] = [round(float(np.clip(7 + 1.5 * x, 0, 10)), 1) for x in Z[a]]
    colunas["nota_nps"] = [int(np.clip(round(7 + 2 * x), 0, 10)) for x in latente]
    if nivel >= 5:
        for i in rng.choice(n, 15, replace=False):
            colunas[attrs[int(i) % len(attrs)]][int(i)] = None
    cab = list(colunas)
    arq = _xlsx(pasta / f"treino_drivers_{semente}.xlsx", cab, [[colunas[c][i] for c in cab] for i in range(n)])
    return {"arquivo": str(arq), "parametros": {"atributos": attrs}, "gabarito": {"driver": driver},
            "resumo": f"pesquisa com {len(attrs)} atributos; o que move o NPS foi plantado"}


def _x_drivers(cfg, sb, pid, caso):
    return _usar("drivers", sb, pid, arquivo=caso["arquivo"])


def _c_drivers(caso, r):
    if (e := _excecao(r)):
        return e
    dr = (r.get("dados") or {}).get("drivers") or {}
    if not dr.get("ok"):
        return (False, 0.0, "nao_rodou", r.get("texto", "")[:200])
    topo = dr["drivers"][0]["atributo"]
    return (True, 1.0, "", f"1º driver: {topo}") if topo == caso["gabarito"]["driver"] else \
        (False, 0.0, "driver_errado", f"apontou {topo}; o plantado é {caso['gabarito']['driver']}")


def _v_drivers(caso, cfg):
    from openpyxl import load_workbook
    wb = load_workbook(caso["arquivo"], read_only=True)
    linhas = list(wb.active.iter_rows(values_only=True))
    wb.close()
    cab = list(linhas[0])
    df = pd.DataFrame(linhas[1:], columns=cab).apply(pd.to_numeric, errors="coerce").dropna()
    cor = {a: abs(float(np.corrcoef(df[a], df["nota_nps"])[0, 1])) for a in caso["parametros"]["atributos"]}
    ordem = sorted(cor, key=lambda a: -cor[a])
    if ordem[0] != caso["gabarito"]["driver"]:
        return [f"na conta independente o mais correlacionado é {ordem[0]}, não o plantado"]
    if cor[ordem[0]] - cor[ordem[1]] < 0.05:
        return ["driver plantado sem folga sobre o 2º (diferença de correlação < 0,05): gabarito ambíguo"]
    return []


# =============================================================================== 8. conhecimento (todos)
def _palavras(t: str) -> list[str]:
    return [w for w in re.findall(r"[\wÀ-ÿ]+", t) if len(w) >= 3]


def _g_conhecimento(nivel, semente, pasta, cfg):
    from robo import licoes
    with licoes._tb(cfg).tx() as c:
        # lições geradas pelo próprio treino são quase iguais entre si (mesmo erro, nível 4 e 5): gabarito ambíguo
        cand = [dict(r) for r in c.execute("SELECT id, titulo, pergunta, origem_tipo, origem_ref FROM licoes "
                                           "WHERE arquivada=0 AND length(pergunta) >= 12 AND origem_tipo <> 'treino' ORDER BY id")]
    if not cand:
        raise ValueError("a empresa ainda não tem lições com pergunta: nada para perguntar")
    rng = _rng(semente, nivel, "conhecimento")
    l = cand[int(rng.integers(0, len(cand)))]
    p = l["pergunta"]
    ws = _palavras(p)
    if nivel == 2:
        p = " ".join(sem_acento(w).lower() for w in rng.permutation(ws))
    elif nivel >= 3:
        # fica com as palavras MAIS raras no caderno (as que identificam o assunto), embaralhadas: a pergunta continua
        # tendo resposta; o difícil é achar com menos palavras (e, nos níveis 4 e 5, com erro de digitação)
        freq = {}
        for x in cand:
            for w in set(sem_acento(x["pergunta"] + " " + x["titulo"]).lower().split()):
                freq[w] = freq.get(w, 0) + 1
        k = max(2, int(len(ws) * (0.6 if nivel == 3 else 0.5)))
        raras = sorted(dict.fromkeys(sem_acento(w).lower() for w in ws), key=lambda w: (freq.get(w, 0), w))[:k]
        p = " ".join(rng.permutation(raras))
        sem_erro = p
        if nivel >= 4:  # erro de digitação numa palavra
            w = list(p)
            pos = int(rng.integers(1, max(2, len(w) - 1)))
            if w[pos].isalpha() and w[pos - 1].isalpha():
                w[pos], w[pos - 1] = w[pos - 1], w[pos]
            p = "".join(w)
    return {"arquivo": None, "parametros": {"pergunta": p, "pergunta_original": l["pergunta"], "licao": l["id"],
                                            "pergunta_sem_erro": locals().get("sem_erro", p),
                                            "titulo": l["titulo"], "origem": f"{l['origem_tipo']}:{l['origem_ref']}"},
            "gabarito": {"licao": l["id"], "titulo": l["titulo"]},
            "resumo": f"achar a lição “{l['titulo'][:60]}” pela pergunta “{p[:60]}”"}


def _x_conhecimento(cfg, sb, pid, caso):
    return _usar("conhecimento", cfg, pid, pergunta=caso["parametros"]["pergunta"], registrar_uso=False)


def _c_conhecimento(caso, r):
    if (e := _excecao(r)):
        return e
    if not r.get("ok"):
        return (False, 0.0, "nao_encontrou", "não achou nada no conhecimento")
    fonte = str((r.get("dados") or {}).get("fonte") or "")
    if fonte.split(".v")[0] == caso["gabarito"]["licao"].split(".v")[0]:
        return (True, 1.0, "", f"achou a lição certa ({fonte})")
    origem = caso["parametros"]["origem"]
    if origem.startswith("manual:") and "#" in origem:  # a mesma seção do manual, respondida direto do .md
        arquivo, secao = origem.split(":", 1)[1].split("#", 1)
        if Path(fonte).name == arquivo and r.get("texto", "").startswith(f"{secao}:"):
            return (True, 1.0, "", f"achou a seção certa do manual ({arquivo} › {secao})")
    return (False, 0.0, "achou_outra_coisa", f"respondeu com {fonte or r.get('evidencia', '')[:80]}")


def _v_conhecimento(caso, cfg):
    """Revisor: a pergunta ORIGINAL tem que achar a lição em 1º lugar no índice — senão o gabarito não é justo."""
    from robo import licoes
    alvo = caso["gabarito"]["licao"].split(".v")[0]
    problemas = []
    for chave, rotulo in (("pergunta_original", "a pergunta original"), ("pergunta_sem_erro", "a pergunta do nível (sem o erro de digitação)")):
        achadas = [l["id"].split(".v")[0] for l in (licoes.buscar(cfg, caso["parametros"].get(chave) or "") or [])[:1]]
        if alvo not in achadas:
            problemas.append(f"{rotulo} não acha esta lição em 1º lugar no índice: caso injusto")
    with licoes._tb(cfg).tx() as c:  # o título do gabarito é o da lição no caderno (não outro)
        r = c.execute("SELECT titulo FROM licoes WHERE id=?", (caso["gabarito"]["licao"],)).fetchone()
    if not r:
        problemas.append(f"a lição {caso['gabarito']['licao']} do gabarito não existe no caderno")
    elif r["titulo"] != caso["gabarito"]["titulo"]:
        problemas.append(f"o título no caderno é “{r['titulo']}”; o gabarito diz “{caso['gabarito']['titulo']}”")
    return problemas


# =============================================================================== 9. histórico (todos)
def _g_historico(nivel, semente, pasta, cfg):
    from robo import historico
    rng = _rng(semente, nivel, "historico")
    logs = pasta / "logs"
    logs.mkdir(parents=True, exist_ok=True)
    n = {1: 15, 2: 200, 3: 200, 4: 3000, 5: 400}[nivel]
    status = ["sucesso", "reprovado", "adiado"]
    t0 = datetime(2026, 1, 1, 8, 0)
    registros = []
    for i in range(n):
        nome = f"base_{int(rng.integers(1, 50)):02d}.xlsx" if nivel < 5 else f"base;{i}'\"x.xlsx"
        registros.append({"data_hora": (t0 + timedelta(minutes=7 * i)).strftime("%Y-%m-%d %H:%M:%S"), "arquivo": nome,
                          "status": str(rng.choice(status, p=[0.7, 0.25, 0.05])), "linhas_entrada": int(rng.integers(10, 900)),
                          "linhas_saida": 0, "erros": 0, "avisos": 0, "duracao_seg": 1.0, "saidas": "", "id": f"exec_{i}"})
    pendentes = registros[-3:] if nivel in (3, 5) else []
    historico._gravar_linhas(logs, registros[: len(registros) - len(pendentes)])
    if pendentes:  # o CSV estava aberto no Excel: as últimas ficaram na fila de pendentes
        (logs / "historico_pendente.jsonl").write_text("\n".join(json.dumps(p, ensure_ascii=False) for p in pendentes) + "\n",
                                                        encoding="utf-8")
    ultimos = [[r["arquivo"], r["status"]] for r in reversed(registros)][:10]
    return {"arquivo": None, "parametros": {"logs": str(logs), "n": n, "pendentes": len(pendentes)},
            "gabarito": {"ultimos": ultimos}, "resumo": f"histórico com {n} execução(ões)" +
            (f", {len(pendentes)} esperando o CSV ser fechado" if pendentes else "")}


def _x_historico(cfg, sb, pid, caso):
    from robo.config import Config
    import copy
    dados = copy.deepcopy(sb.dados)
    dados["pastas"]["logs"] = caso["parametros"]["logs"]
    return _usar("historico", Config(dados, sb.base), pid, limite=10)


def _c_historico(caso, r):
    if (e := _excecao(r)):
        return e
    obt = [[x.get("arquivo"), x.get("status")] for x in (r.get("dados") or {}).get("execucoes") or []]
    g = caso["gabarito"]["ultimos"]
    if obt != g:
        dif = next((i for i, (a, b) in enumerate(zip(obt, g)) if a != b), min(len(obt), len(g)))
        return (False, round(dif / len(g), 2), "historico_errado", f"a partir da posição {dif + 1}: {obt[dif:dif + 1]} ≠ {g[dif:dif + 1]}")
    return (True, 1.0, "", "as 10 últimas execuções certas e na ordem")


def _v_historico(caso, cfg):
    """Revisor: lê o CSV e a fila de pendentes direto (sem robo.historico) e refaz as 10 últimas execuções."""
    p = Path(caso["parametros"]["logs"]) / "historico_processamentos.csv"
    with open(p, encoding="utf-8-sig", newline="") as f:
        linhas = list(csv.DictReader(f, delimiter=";"))
    esperado = caso["parametros"]["n"] - caso["parametros"]["pendentes"]
    problemas = [] if len(linhas) == esperado else [f"CSV gravado com {len(linhas)} linha(s); deveria ter {esperado}"]
    pend = Path(caso["parametros"]["logs"]) / "historico_pendente.jsonl"
    if pend.exists():
        linhas += [json.loads(l) for l in pend.read_text(encoding="utf-8").splitlines() if l.strip()]
    linhas.sort(key=lambda r: r["data_hora"])
    refeito = [[r["arquivo"], r["status"]] for r in reversed(linhas)][:10]
    if refeito != caso["gabarito"]["ultimos"]:
        problemas.append("as 10 últimas execuções lidas do CSV não batem com o gabarito")
    return problemas


# =============================================================================== 10. ronda de TI
_PLANTIOS_TI = {
    "sem_senha_painel": ("sec_senior", "sem senha"),
    "senha_banco_no_config": ("sec_senior", "senha do PostgreSQL gravada"),
    "erros_no_log": ("dev_junior", "erro(s) de programa no log"),
    "codigo_novo": ("sec_senior", "código de habilidade novo/alterado"),
}


def _g_ronda_ti(nivel, semente, pasta, cfg):
    rng = _rng(semente, nivel, "ronda")
    ordem = list(_PLANTIOS_TI)
    if nivel <= 4:
        plantados = ordem[:nivel]
    else:  # desafio: só alguns problemas, sorteados — achar o que NÃO existe também conta
        plantados = [p for p in ordem if rng.random() < 0.5] or [ordem[int(rng.integers(0, 4))]]
    return {"arquivo": None, "parametros": {"plantados": plantados, "pasta": str(pasta)},
            "gabarito": {"problemas": {p: (p in plantados) for p in ordem}},
            "resumo": f"computador com {len(plantados)} problema(s) plantado(s): {', '.join(plantados)}"}


def _x_ronda_ti(cfg, sb, pid, caso):
    import copy
    from robo.config import Config
    from robo import ti
    p = caso["parametros"]
    pasta = Path(sb.dados["pastas"]["logs"]).parent / "ti_sandbox"  # computador de teste NOVO a cada tentativa
    dados = copy.deepcopy(sb.dados)
    dados["pastas"] = {k: str(pasta / k) for k in ("entrada", "logs", "agentes", "habilidades")}
    dados.setdefault("painel", {})["senha_hash"] = "" if "sem_senha_painel" in p["plantados"] else "pbkdf2$treino"
    dados["banco"] = {**(dados.get("banco") or {}), "habilitado": False,
                      "senha": "senha123" if "senha_banco_no_config" in p["plantados"] else ""}
    c = Config(dados, sb.base)
    c.criar_pastas()
    Path(c.habilidades).mkdir(parents=True, exist_ok=True)
    ti.ronda(c, anunciar=False)  # 1ª ronda: fotografa o estado (o "antes")
    if "erros_no_log" in p["plantados"]:
        agora = datetime.now().strftime("%d/%m/%Y %H:%M")
        limite = int(ti._c(c)["erros_log_alerta"])
        (Path(c.logs) / "robo.log").write_text("".join(f"{agora}:00 | ERROR | robo | falha {i}\n" for i in range(limite + 2)),
                                               encoding="utf-8")
    if "codigo_novo" in p["plantados"]:
        (Path(c.habilidades) / "habilidade_nova.py").write_text("# código que apareceu sem aviso\n", encoding="utf-8")
    return _usar("ronda_ti", c, pid)


def _c_ronda_ti(caso, r):
    if (e := _excecao(r)):
        return e
    itens = {i["id"]: " ".join(i.get("achados") or []) for i in (r.get("dados") or {}).get("itens") or []}
    erros = []
    for plantio, esperado in caso["gabarito"]["problemas"].items():
        quem, trecho = _PLANTIOS_TI[plantio]
        achou = trecho in itens.get(quem, "")
        if achou != esperado:
            erros.append(f"{plantio}: {'não detectou' if esperado else 'alarme falso'}")
    if erros:
        return (False, round(1 - len(erros) / 4, 2), "problema_nao_detectado" if any("não detectou" in e for e in erros)
                else "falso_alarme", "; ".join(erros))
    return (True, 1.0, "", f"detectou exatamente {sum(caso['gabarito']['problemas'].values())} problema(s)")


def _v_ronda_ti(caso, cfg):
    plantados = caso["parametros"]["plantados"]
    if not all(p in _PLANTIOS_TI for p in plantados):
        return ["plantio desconhecido"]
    if not plantados:
        return ["nenhum problema plantado"]
    refeito = {p: (p in plantados) for p in _PLANTIOS_TI}
    if refeito != caso["gabarito"]["problemas"]:
        return [f"pelos plantios, o gabarito seria {refeito}; está {caso['gabarito']['problemas']}"]
    return []


# =============================================================================== 11. resumo executivo
_RE_NUM = re.compile(r"\d+(?:[.,]\d+)*")


def _g_resumo(nivel, semente, pasta, cfg):
    rng = _rng(semente, nivel, "resumo")
    k = {1: 2, 2: 3, 3: 4, 4: 5, 5: 8}[nivel]
    temas = ["NPS da pesquisa de alunos", "Vendas por região", "Clima dos colaboradores", "Contas a receber",
             "Campanhas digitais", "Entregas", "Atendimento", "Retenção"]
    partes = []
    for i in range(k):
        v = round(float(rng.uniform(-50, 95)), 1)
        n = int(rng.integers(40, 3000))
        falhou = nivel >= 3 and i == k - 1
        partes.append({"titulo": temas[i], "natureza": "limitacao" if falhou else "calculo",
                       "evidencia": f"base_{i}.xlsx", "texto": ("Não foi feito: faltou a coluna de data." if falhou else
                                                              f"Resultado {str(v).replace('.', ',')} com base de {n} respostas.")})
    if nivel >= 4:
        partes.append({"titulo": "Etapa sem resultado", "natureza": "limitacao", "texto": ""})  # sem texto: não entra
    return {"arquivo": None, "parametros": {"partes": partes},
            "gabarito": {"titulos": [p["titulo"] for p in partes if p["texto"]],
                         "numeros": sorted({x for p in partes for x in _RE_NUM.findall(p["texto"])})},
            "resumo": f"resumir {sum(1 for p in partes if p['texto'])} resultado(s) sem inventar número"}


def _x_resumo(cfg, sb, pid, caso):
    return _usar("resumo_executivo", sb, pid, titulo="Resumo de treino", partes=caso["parametros"]["partes"], ref="treino")


def _c_resumo(caso, r):
    if (e := _excecao(r)):
        return e
    arts = r.get("artefatos") or []
    if not r.get("ok") or not arts or not Path(arts[0]).exists():
        return (False, 0.0, "nao_gerou", r.get("texto", "")[:200])
    md = Path(arts[0]).read_text(encoding="utf-8")
    corpo = "\n".join(l for l in md.splitlines() if not l.startswith("Gerado em"))
    faltam = [t for t in caso["gabarito"]["titulos"] if f"## {t}" not in md]
    permitidos = set(caso["gabarito"]["numeros"]) | {re.sub(r"\D", "", t) for t in caso["gabarito"]["titulos"]}
    inventados = sorted({x for x in _RE_NUM.findall(corpo) if x not in permitidos and not re.fullmatch(r"\d", x)} -
                        {x for x in _RE_NUM.findall(" ".join(p.get("evidencia", "") for p in caso["parametros"]["partes"]))})
    if faltam or inventados:
        return (False, 0.0, "numero_inventado" if inventados else "parte_omitida",
                (f"faltou: {faltam}" if faltam else "") + (f" número(s) que não estavam nas etapas: {inventados[:5]}" if inventados else ""))
    return (True, 1.0, "", f"{len(caso['gabarito']['titulos'])} seção(ões), nenhum número inventado")


def _v_resumo(caso, cfg):
    """Revisor: refaz títulos e números das partes palavra por palavra (sem a expressão regular do gerador)."""
    partes = caso["parametros"]["partes"]
    titulos = [p["titulo"] for p in partes if p.get("texto")]
    numeros = set()
    for p in partes:
        for palavra in str(p.get("texto") or "").split():
            palavra = palavra.strip("-+().:;!?%")
            if palavra and palavra[0].isdigit():
                numeros.add(palavra)
    problemas = [] if titulos else ["caso sem nenhuma parte com resultado"]
    if titulos != caso["gabarito"]["titulos"]:
        problemas.append(f"títulos com resultado: {titulos}; o gabarito diz {caso['gabarito']['titulos']}")
    if sorted(numeros) != caso["gabarito"]["numeros"]:
        problemas.append(f"números nas partes: {sorted(numeros)}; o gabarito diz {caso['gabarito']['numeros']}")
    return problemas


# =============================================================================== 12. desempenho (RH)
def _situacao_independente(a: int, e: int, minima: float) -> str:
    from robo.memoria import MINIMO_EVENTOS
    n = a + e
    if n < MINIMO_EVENTOS:
        return "sem_dados" if not n else ("atencao" if e else "ok")
    t = 100 * a / n
    if t < minima:
        return "insuficiente"
    return "atencao" if t < max(90, minima + 5) else "ok"


def _g_desempenho(nivel, semente, pasta, cfg):
    from robo import organizacao
    from robo.lideranca import EXIGENCIA
    from robo.memoria import TAXA_MINIMA
    rng = _rng(semente, nivel, "desempenho")
    gente = [p["id"] for p in organizacao.lista(None, com_contratados=False)]
    k = {1: 4, 2: 12, 3: 12, 4: 15, 5: 18}[nivel]
    escolhidos = list(rng.choice([g for g in gente if g != "esp_pandas"], k, replace=False))
    if nivel >= 5:
        escolhidos[0] = "esp_pandas"  # guilda Python: exigência de 90% (sem repetir a pessoa na lista)
    agora = datetime.now()
    eventos, gab = [], {}
    for pid in escolhidos:
        a, e = int(rng.integers(0, 15)), int(rng.integers(0, 6))
        if nivel >= 5 and pid == "esp_pandas":
            a, e = 18, 2  # exatamente 90%
        for t, q in (("acerto", a), ("erro", e)):
            for _ in range(q):
                eventos.append((pid, (agora - timedelta(hours=float(rng.uniform(1, 150)))).isoformat(timespec="seconds"), t))
        if nivel >= 3:  # fora da janela de 7 dias: não conta
            eventos += [(pid, (agora - timedelta(days=float(rng.uniform(8, 20)))).isoformat(timespec="seconds"), "erro")] * 2
        if nivel >= 4:  # outros tipos não entram na taxa
            eventos += [(pid, (agora - timedelta(hours=5)).isoformat(timespec="seconds"), t) for t in ("treinamento", "colaboracao", "aprendizado")]
        minima = EXIGENCIA.get(pid, TAXA_MINIMA)
        gab[pid] = {"acertos": a, "erros": e, "taxa": round(100 * a / (a + e), 1) if a + e else None,
                    "situacao": _situacao_independente(a, e, minima)}
    logs = pasta / "logs"
    from robo.trabalho import Trabalho
    with Trabalho(arquivo=logs / "trabalho.db").tx() as c:
        c.executemany("INSERT INTO experiencias (funcionario, data_hora, tipo, resumo, origem) VALUES (?,?,?,'treino','treino')",
                      eventos)
    return {"arquivo": None, "parametros": {"logs": str(logs), "pessoas": escolhidos},
            "gabarito": {"por_pessoa": {p: g for p, g in gab.items() if g["acertos"] + g["erros"]}},
            "resumo": f"semana sintética de {k} pessoa(s)" + (" com eventos fora da janela" if nivel >= 3 else "")}


def _x_desempenho(cfg, sb, pid, caso):
    import copy
    from robo.config import Config
    dados = copy.deepcopy(sb.dados)
    dados["pastas"]["logs"] = caso["parametros"]["logs"]
    return _usar("desempenho", Config(dados, sb.base), pid)


def _c_desempenho(caso, r):
    if (e := _excecao(r)):
        return e
    obt = {x["id"]: x for x in (r.get("dados") or {}).get("desempenho") or []}
    erros = []
    for pid, g in caso["gabarito"]["por_pessoa"].items():
        o = obt.get(pid)
        if not o:
            erros.append(f"{pid} sumiu do relatório")
        elif (o["acertos"], o["erros"], o["taxa"], o["situacao"]) != (g["acertos"], g["erros"], g["taxa"], g["situacao"]):
            erros.append(f"{pid}: deu {o['acertos']}/{o['erros']} {o['taxa']}% {o['situacao']}; certo {g['acertos']}/"
                         f"{g['erros']} {g['taxa']}% {g['situacao']}")
    if erros:
        return (False, round(1 - len(erros) / max(1, len(caso["gabarito"]["por_pessoa"])), 2), "conta_de_pessoas_errada",
                "; ".join(erros[:3]))
    return (True, 1.0, "", f"{len(caso['gabarito']['por_pessoa'])} pessoa(s) com taxa e situação certas")


def _v_desempenho(caso, cfg):
    """Revisor: reconta direto no banco sintético, por pessoa, só acerto/erro dentro dos 7 dias."""
    import sqlite3
    c = sqlite3.connect(Path(caso["parametros"]["logs"]) / "trabalho.db")
    desde = (datetime.now() - timedelta(days=7)).isoformat(timespec="seconds")
    contas = {}
    for f, t, n in c.execute("SELECT funcionario, tipo, COUNT(*) FROM experiencias WHERE data_hora>=? AND tipo IN "
                             "('acerto','erro') GROUP BY 1, 2", (desde,)):
        contas.setdefault(f, {"acerto": 0, "erro": 0})[t] = n
    c.close()
    problemas = []
    if len(set(caso["parametros"]["pessoas"])) != len(caso["parametros"]["pessoas"]):
        problemas.append("pessoa repetida na semana sintética")
    from robo.lideranca import EXIGENCIA
    from robo.memoria import TAXA_MINIMA
    for pid, g in caso["gabarito"]["por_pessoa"].items():
        x = contas.get(pid, {"acerto": 0, "erro": 0})
        if (x["acerto"], x["erro"]) != (g["acertos"], g["erros"]):
            problemas.append(f"{pid}: banco tem {x['acerto']}/{x['erro']}, gabarito diz {g['acertos']}/{g['erros']}")
        n = x["acerto"] + x["erro"]
        taxa = round(100 * x["acerto"] / n, 1) if n else None
        situacao = _situacao_independente(x["acerto"], x["erro"], EXIGENCIA.get(pid, TAXA_MINIMA))
        if (g["taxa"], g["situacao"]) != (taxa, situacao):
            problemas.append(f"{pid}: pela política, {taxa}% {situacao}; o gabarito diz {g['taxa']}% {g['situacao']}")
    com_eventos = sorted(f for f, x in contas.items() if f in caso["parametros"]["pessoas"] and x["acerto"] + x["erro"])
    if com_eventos != sorted(caso["gabarito"]["por_pessoa"]):
        problemas.append("as pessoas do gabarito não são as que têm acerto/erro na semana")
    return problemas


# =============================================================================== 13. SQL (consulta_sql) — só com PostgreSQL
def _g_sql(nivel, semente, pasta, cfg):
    from robo import conferencia, laboratorio
    area = laboratorio.AREAS[(semente // 3) % len(laboratorio.AREAS)]
    caso = laboratorio.gerar(area, semente)
    df = laboratorio.verdade(caso)
    seg, met = normalizar_nome_coluna(caso["segmento"]), normalizar_nome_coluna(caso["metrica"])
    op = caso["op"] if caso["op"] != "nps" else "media"
    rng = _rng(semente, nivel, "sql")
    por, filtros = [seg], []
    if nivel == 2:
        por, filtros = [], [(seg, "=", str(rng.choice(sorted(df[seg].unique()))))]
    elif nivel == 3:
        op = "contagem"
    elif nivel == 4:  # valor com apóstrofo: o SQL precisa escapar direito
        df[seg] = df[seg].astype(str) + np.where(np.arange(len(df)) % 2, " d'Oeste", "")
        filtros = [(seg, "=", str(df[seg].iloc[1]))]
        por = []
    elif nivel == 5 and area in _SEGUNDA:
        por = [seg, normalizar_nome_coluna(_SEGUNDA[area])]
    dados = df
    for c, s, v in filtros:
        dados = dados[dados[c].astype(str) == v]
    esperado = conferencia.valores_pandas(op, dados, met, por)
    tabela = f"treino_{area}_{semente}_{nivel}"
    arq = pasta / f"{tabela}.csv"
    df.to_csv(arq, index=False)
    return {"arquivo": str(arq), "parametros": {"tabela": tabela, "op": op, "metrica": met, "por": por, "filtros": filtros},
            "gabarito": {"valores": esperado}, "resumo": f"{op} de {met} em SQL" + (f" por {', '.join(por)}" if por else "")
            + (" com filtro" if filtros else "")}


def _x_sql(cfg, sb, pid, caso):
    from robo import conferencia
    from robo.banco import Banco, tipo_pg
    p = caso["parametros"]
    df = pd.read_csv(caso["arquivo"])
    b = Banco(cfg.banco, cfg.base)
    b.criar_banco_se_preciso()
    tipos = {c: tipo_pg(df[c]) for c in df.columns}
    with b.conectar() as conn:
        with conn.transaction():
            conn.execute("CREATE SCHEMA IF NOT EXISTS treino")
            tipos = b._garantir_tabela(conn, "treino", p["tabela"], tipos)
            b._copiar(conn, "treino", p["tabela"], df, tipos)
    try:
        consulta = conferencia.montar_sql(p["op"], p["metrica"], p["por"], p["filtros"], f'"treino"."{p["tabela"]}"', None)
        return _usar("consulta_sql", cfg, pid, sql=consulta)
    finally:
        with b.conectar() as conn:  # a tabela sintética é refeita pela semente: não precisa ocupar o banco
            conn.execute(f'DROP TABLE IF EXISTS "treino"."{p["tabela"]}"')


def _c_sql(caso, r):
    from robo import conferencia
    if (e := _excecao(r)):
        return e
    if not r.get("ok"):
        return (False, 0.0, "sql_recusado", r.get("texto", "")[:200])
    d = r.get("dados") or {}
    cols, linhas = d.get("colunas") or [], d.get("linhas") or []
    i = cols.index("valor") if "valor" in cols else None
    if i is None:
        return (False, 0.0, "resultado_sem_valor", f"colunas: {cols}")
    obt = {}
    for l in linhas:
        chave = conferencia._chave(tuple(l[:i])) if i > 1 else (conferencia._chave(l[0]) if i == 1 else "(total)")
        obt[chave] = float(l[i]) if l[i] not in (None, "None") else None
    g = caso["gabarito"]["valores"]
    erradas = [k for k in g if not _fechado(obt.get(k), g[k], 1e-6)]
    if erradas or set(obt) != set(g):
        return (False, round(1 - len(erradas) / max(1, len(g)), 2), "sql_numero_errado",
                "; ".join(f"{k}: deu {obt.get(k)}, certo {g[k]}" for k in erradas[:3]) or f"grupos {sorted(obt)[:3]}")
    return (True, 1.0, "", f"{len(g)} valor(es) certos no PostgreSQL")


def _v_sql(caso, cfg):
    """Revisor: refaz a conta lendo o CSV com o módulo csv (texto puro, sem pandas), filtro e grupos à mão."""
    from robo.conferencia import _chave
    p = caso["parametros"]
    with open(caso["arquivo"], encoding="utf-8", newline="") as f:
        linhas = list(csv.DictReader(f))
    if not linhas:
        return ["tabela vazia"]
    grupos: dict = {}
    for l in linhas:
        if any(str(l.get(c)) != str(v) for c, _s, v in p["filtros"]):
            continue
        if any(l.get(c) in (None, "") for c in p["por"]):
            continue
        chave = _chave(tuple(l[c] for c in p["por"])) if p["por"] else "(total)"
        x = l.get(p["metrica"]) if p["metrica"] else None
        try:
            x = float(x) if x not in (None, "") else None
        except ValueError:
            x = None
        grupos.setdefault(chave, []).append(x)
    refeito = {}
    for chave, xs in grupos.items():
        num = [x for x in xs if x is not None and not math.isnan(x)]
        if p["op"] == "contagem":
            refeito[chave] = float(len(xs))
        elif not num:
            refeito[chave] = float("nan")
        elif p["op"] == "soma":
            refeito[chave] = math.fsum(num)
        elif p["op"] == "media":
            refeito[chave] = math.fsum(num) / len(num)
        elif p["op"] == "mediana":
            o = sorted(num)
            refeito[chave] = o[len(o) // 2] if len(o) % 2 else (o[len(o) // 2 - 1] + o[len(o) // 2]) / 2
        elif p["op"] in ("minimo", "maximo"):
            refeito[chave] = min(num) if p["op"] == "minimo" else max(num)
        else:
            return [f"operação {p['op']} sem conferência independente"]
    g = caso["gabarito"]["valores"]
    problemas = []
    if set(refeito) != set(g):
        problemas.append(f"grupos refeitos {sorted(refeito)[:4]}; o gabarito tem {sorted(g)[:4]}")
    for chave in set(refeito) & set(g):
        a, b = refeito[chave], g[chave]
        if not ((a != a and (b is None or b != b)) or _fechado(a, b, 1e-9)):
            problemas.append(f"{chave}: conta independente {a}, gabarito {b}")
    return problemas



# =============================================================================== 14. orquestração (SLA das rotinas — Airflow)
def _orq_pool():
    from robo.orquestracao import APENAS_NO_ROBO
    from robo.rotinas import ROTINAS
    return {n: float(v[2]) * 60 for n, v in ROTINAS.items() if n not in APENAS_NO_ROBO}  # nome → intervalo em minutos


def _g_orquestracao(nivel, semente, pasta, cfg):
    rng = _rng(semente, nivel, "orquestracao")
    pool = _orq_pool()
    nomes = [str(x) for x in rng.permutation(sorted(pool))]
    base = {n: float(rng.uniform(5, 40)) for n in nomes}
    hist = {n: [{"nome": n, "fim_min_atras": 35 + 30 * k, "segundos": round(base[n] * float(rng.uniform(0.8, 1.2)), 1),
                 "status": "fez"} for k in range(8)] for n in nomes}
    estado = {n: {"nome": n, "verif_min_atras": round(0.3 * pool[n], 2), "ultimo_status": "fez"} for n in nomes}
    if nivel <= 4:
        qtd = {1: (1, 0, 0), 2: (1, 1, 0), 3: (1, 1, 1), 4: (1, 1, 1)}[nivel]
    else:
        qtd = tuple(int(rng.integers(0, 3)) for _ in range(3))
        if sum(qtd) == 0:
            qtd = (0, 0, 1)
    fila = list(nomes)
    erro, atraso, lenta = [fila.pop() for _ in range(qtd[0])], [fila.pop() for _ in range(qtd[1])], [fila.pop() for _ in range(qtd[2])]
    plantios = {"erro": erro, "atraso": atraso, "lenta": lenta}
    for n in erro:
        hist[n].append({"nome": n, "fim_min_atras": 1, "segundos": 3.0, "status": "erro"})
    for n in atraso:
        estado[n]["verif_min_atras"] = round(3 * pool[n] + 10, 2)
    for n in lenta:
        hist[n].append({"nome": n, "fim_min_atras": 1, "segundos": round(base[n] * 6 + 120, 1), "status": "fez"})
    if nivel >= 4:  # iscas: parecem problema mas a regra diz que não são
        recuperada, lenta_antiga, pouco_hist = fila.pop(), fila.pop(), fila.pop()
        hist[recuperada] += [{"nome": recuperada, "fim_min_atras": 3, "segundos": 2.0, "status": "erro"},
                             {"nome": recuperada, "fim_min_atras": 1, "segundos": base[recuperada], "status": "fez"}]
        hist[lenta_antiga][3]["segundos"] = round(base[lenta_antiga] * 10 + 300, 1)
        hist[pouco_hist] = hist[pouco_hist][-3:] + [{"nome": pouco_hist, "fim_min_atras": 1,
                                                    "segundos": round(base[pouco_hist] * 8 + 200, 1), "status": "fez"}]
        plantios["iscas"] = {"erro_recuperado": recuperada, "lenta_antiga": lenta_antiga, "historico_curto": pouco_hist}
        if nivel == 5 and fila:
            quase = fila.pop()
            estado[quase]["verif_min_atras"] = round(1.2 * pool[quase], 2)  # perto do limite, mas dentro
            plantios["iscas"]["perto_do_limite"] = quase
    from robo.orquestracao import APENAS_NO_ROBO
    em_dia = [{"nome": n, "verif_min_atras": 0.1, "ultimo_status": "fez"} for n in sorted(APENAS_NO_ROBO)]  # o canal (robô) em dia
    cenario = {"sessao_min_atras": 60 * 24 * 40, "estado": list(estado.values()) + em_dia,
               "execucoes": [e for n in nomes for e in hist[n]]}
    gabarito = {"atrasadas": sorted(atraso), "com_erro": sorted(erro), "lentas": sorted(lenta)}
    return {"arquivo": None, "parametros": {"cenario": cenario, "plantios": plantios},
            "gabarito": gabarito,
            "resumo": f"{len(nomes)} rotinas; plantado: {len(erro)} com erro, {len(atraso)} atrasada(s), {len(lenta)} lenta(s)"
                      + (f"; iscas: {', '.join(plantios['iscas'])}" if "iscas" in plantios else "")}


def _x_orquestracao(cfg, sb, pid, caso):
    import copy
    from robo.config import Config
    from robo import orquestracao
    raiz = Path(sb.dados["pastas"]["logs"]).parent
    n = 1
    while (raiz / f"orq_sandbox_{n}").exists():  # computador de teste NOVO a cada tentativa
        n += 1
    pasta = raiz / f"orq_sandbox_{n}"
    dados = copy.deepcopy(sb.dados)
    dados["pastas"] = {k: str(pasta / k) for k in ("entrada", "logs", "agentes", "habilidades")}
    c = Config(dados, sb.base)
    c.criar_pastas()
    orquestracao.gravar_cenario(c, caso["parametros"]["cenario"])
    return _usar("orquestracao", c, pid)


def _c_orquestracao(caso, r):
    if (e := _excecao(r)):
        return e
    d = r.get("dados") or {}
    g = caso["gabarito"]
    total = len(_orq_pool()) * 3
    faltou, sobrou = [], []
    for chave in ("atrasadas", "com_erro", "lentas"):
        esperado, achado = set(g[chave]), set(d.get(chave) or [])
        faltou += [f"{chave}: {x}" for x in sorted(esperado - achado)]
        sobrou += [f"{chave}: {x}" for x in sorted(achado - esperado)]
    if faltou or sobrou:
        detalhe = "; ".join([f"não detectou {x}" for x in faltou] + [f"alarme falso {x}" for x in sobrou])
        return (False, round(1 - (len(faltou) + len(sobrou)) / total, 2),
                "problema_nao_detectado" if faltou else "falso_alarme", detalhe)
    n = sum(len(g[k]) for k in g)
    return (True, 1.0, "", f"apontou exatamente {n} problema(s) e nenhuma isca")


def _v_orquestracao(caso, cfg):
    """Revisor: refaz o gabarito direto do cenário, com laços simples, e confere que o nível tem o que promete."""
    pool = _orq_pool()
    cen = caso["parametros"]["cenario"]
    problemas = []
    faltam = sorted(set(pool) - {e["nome"] for e in cen["estado"]})
    if faltam:  # caso feito antes de rotinas novas existirem: a ferramenta acusaria atraso nelas (caso injusto)
        problemas.append(f"cenário desatualizado: faltam as rotinas {', '.join(faltam)}")
    atrasadas = sorted(e["nome"] for e in cen["estado"] if e["nome"] in pool and e["verif_min_atras"] > 1.5 * pool[e["nome"]])
    com_erro, lentas = [], []
    for nome in pool:
        execs = sorted((e for e in cen["execucoes"] if e["nome"] == nome), key=lambda e: e["fim_min_atras"])
        if not execs:
            continue
        if execs[0]["status"] == "erro":
            com_erro.append(nome)
            continue
        feitas = [e["segundos"] for e in execs if e["status"] == "fez"]
        anteriores = sorted(feitas[1:21])
        if len(anteriores) >= 5:
            meio = len(anteriores) // 2
            med = anteriores[meio] if len(anteriores) % 2 else (anteriores[meio - 1] + anteriores[meio]) / 2
            if feitas[0] > 3 * med and feitas[0] > med + 60:
                lentas.append(nome)
    refeito = {"atrasadas": atrasadas, "com_erro": sorted(com_erro), "lentas": sorted(lentas)}
    for chave, valor in refeito.items():
        if valor != caso["gabarito"][chave]:
            problemas.append(f"gabarito de {chave} não confere: gerador {caso['gabarito'][chave]}, revisor {valor}")
    nivel = int(caso.get("nivel") or 0)
    if nivel >= 4 and len((caso["parametros"]["plantios"].get("iscas") or {})) < 3:
        problemas.append("nível 4+ sem as iscas prometidas")
    if not any(caso["gabarito"].values()):
        problemas.append("caso sem nenhum problema plantado")
    return problemas



# =============================================================================== 15. revisão de casos (Sala de Casos)
def revisao_independente(caso: dict, cfg, sha: str | None = None, pasta_tmp: Path | None = None) -> dict:
    """O exame que a revisora faz em QUALQUER caso, sem gravar nada (função pura, usada pelo treino.revisar e pela
    ferramenta revisar_caso):
      1. integridade      — o conteúdo de agora (arquivo + gabarito) é o mesmo que foi gravado (sha da criação)?
      2. reprodutibilidade — gerado de novo pela semente, sai igual ao gravado?
      3. conferência      — o gabarito confere por outro caminho (conferir da competência)?"""
    import shutil
    import tempfile
    c = CATALOGO[caso["competencia"]]
    problemas = []
    atual = conteudo_hash(caso.get("arquivo"), caso["gabarito"])
    sha = sha or atual
    if atual != sha:
        problemas.append("o conteúdo do caso mudou depois de gravado (arquivo ou gabarito alterado)")
    tmp = Path(pasta_tmp) if pasta_tmp else Path(tempfile.mkdtemp(prefix="revisao_"))
    try:
        shutil.rmtree(tmp, ignore_errors=True)
        tmp.mkdir(parents=True)
        de_novo = c.gerar(caso["nivel"], caso["semente"], tmp, cfg)
        if conteudo_hash(de_novo.get("arquivo"), de_novo["gabarito"]) != sha:
            problemas.append("refeito pela mesma semente, o caso saiu diferente (não é reprodutível)")
    except Exception as e:
        problemas.append(f"não consegui refazer o caso: {type(e).__name__}: {e}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    try:
        problemas += list(c.conferir(caso, cfg) or [])
    except Exception as e:
        problemas.append(f"a conferência independente quebrou: {type(e).__name__}: {e}")
    return {"aprovado": not problemas, "problemas": problemas}


DEFEITOS = {"nenhum": "caso correto (tem que ser aprovado)", "gabarito_grosseiro": "gabarito grosseiramente errado",
            "gabarito_sutil": "gabarito com erro sutil", "conteudo": "arquivo do caso alterado depois de gravado"}


def _folhas_gabarito(o, cam=()):
    """Folhas alteráveis do gabarito, em ordem fixa (número, verdadeiro/falso, texto ou lista simples)."""
    if isinstance(o, (bool, int, float, str)):
        yield cam, o
    elif isinstance(o, dict):
        for k in sorted(o, key=str):
            yield from _folhas_gabarito(o[k], cam + (k,))
    elif isinstance(o, list):
        if o and all(not isinstance(x, (dict, list)) for x in o):
            yield cam, o
        else:
            for i, x in enumerate(o):
                yield from _folhas_gabarito(x, cam + (i,))


def _alterar(v, sutil: bool):
    if isinstance(v, bool):
        return not v
    if isinstance(v, int):
        return v + 1 if sutil else v * 2 + 3
    if isinstance(v, float):
        return v + max(abs(v) * 0.005, 0.01) if sutil else v * 2 + 3
    if isinstance(v, str):
        return v + " (alterado)"
    return v[:-1] if v else ["intruso"]


def _sabotar_gabarito(gabarito: dict, rng, sutil: bool) -> str:
    folhas = [(c, v) for c, v in _folhas_gabarito(gabarito) if c]
    if sutil and any(isinstance(v, (bool, int, float)) for _, v in folhas):
        folhas = [(c, v) for c, v in folhas if isinstance(v, (bool, int, float))]
    cam, v = folhas[int(rng.integers(0, len(folhas)))]
    alvo = gabarito
    for k in cam[:-1]:
        alvo = alvo[k]
    alvo[cam[-1]] = _alterar(v, sutil)
    return "/".join(map(str, cam))


def _sabotar_arquivo(arquivo: str, rng) -> str:
    """Altera UMA célula de dado do arquivo (depois de o caso ter sido gravado)."""
    arq = Path(arquivo)
    if arq.suffix.lower() == ".csv":
        with open(arq, encoding="utf-8", newline="") as f:
            linhas = list(csv.reader(f))
        i = 1 + int(rng.integers(0, len(linhas) - 1))
        j = int(rng.integers(0, len(linhas[i])))
        linhas[i][j] = (linhas[i][j] or "0") + "9"
        with open(arq, "w", encoding="utf-8", newline="") as f:
            csv.writer(f).writerows(linhas)
        return f"linha {i + 1}, coluna {j + 1}"
    from openpyxl import load_workbook
    wb = load_workbook(arq)
    ws = wb.active
    cheias = [(c.row, c.column) for linha in ws.iter_rows(min_row=2) for c in linha if c.value not in (None, "")]
    r, col = cheias[int(rng.integers(0, len(cheias)))]
    v = ws.cell(r, col).value
    ws.cell(r, col, v * 3 + 7 if isinstance(v, (int, float)) and not isinstance(v, bool) else f"{v} alterado")
    wb.save(arq)
    return f"célula linha {r}, coluna {col}"


def _revisaveis() -> list[str]:
    """Competências que dá para revisar num computador de teste: o caso não depende do caderno real da empresa."""
    return [c.id for c in CATALOGO.values() if c.ferramenta != "revisar_caso" and not c.usa_conhecimento_real]


def _g_revisao(nivel, semente, pasta, cfg):
    rng = _rng(semente, nivel, "revisao")
    if nivel == 1:
        defeito = "gabarito_grosseiro"
    elif nivel == 2:
        defeito = "gabarito_sutil"
    elif nivel == 3:
        defeito = "nenhum" if semente % 2 == 0 else "gabarito_sutil"
    elif nivel == 4:
        defeito = "conteudo"
    else:
        defeito = str(rng.choice(list(DEFEITOS), p=[0.3, 0.2, 0.25, 0.25]))
    candidatas = _revisaveis()
    ordem = [candidatas[int(i)] for i in rng.permutation(len(candidatas))]
    for comp in ordem:
        sub_nivel, sub_semente = int(rng.integers(1, 6)), int(rng.integers(1, 1_000_000))
        destino = Path(pasta) / "caso_revisado"
        import shutil
        shutil.rmtree(destino, ignore_errors=True)
        destino.mkdir(parents=True)
        sub = CATALOGO[comp].gerar(sub_nivel, sub_semente, destino, cfg)
        if defeito == "conteudo" and not sub.get("arquivo"):
            continue  # alteração de arquivo só em caso que tem arquivo
        break
    sub.update({"competencia": comp, "nivel": sub_nivel, "semente": sub_semente})
    onde = ""
    if defeito in ("gabarito_grosseiro", "gabarito_sutil"):
        onde = _sabotar_gabarito(sub["gabarito"], rng, defeito == "gabarito_sutil")  # o designer gravou a resposta errada
    sha = conteudo_hash(sub.get("arquivo"), sub["gabarito"])
    if defeito == "conteudo":
        onde = _sabotar_arquivo(sub["arquivo"], rng)  # alterado DEPOIS de gravado
    revisar = {k: sub[k] for k in ("competencia", "nivel", "semente", "parametros", "gabarito", "arquivo") if k in sub}
    revisar["sha"] = sha
    return {"arquivo": sub.get("arquivo"), "parametros": {"caso": revisar},
            "gabarito": {"deve_aprovar": defeito == "nenhum", "defeito": defeito, "onde": onde},
            "resumo": f"revisar um caso de {CATALOGO[comp].nome} (nível {sub_nivel})"}


def _x_revisao(cfg, sb, pid, caso):
    return _usar("revisar_caso", sb, pid, caso=caso["parametros"]["caso"])


def _c_revisao(caso, r):
    if (e := _excecao(r)):
        return e
    d = r.get("dados") or {}
    if "aprovado" not in d:
        return (False, 0.0, "nao_revisou", r.get("texto", "")[:200])
    g = caso["gabarito"]
    if bool(d["aprovado"]) == g["deve_aprovar"]:
        return (True, 1.0, "", "aprovou o caso correto" if g["deve_aprovar"] else
                f"reprovou o caso com defeito ({DEFEITOS[g['defeito']]}): " + "; ".join(d.get("problemas") or [])[:150])
    if g["deve_aprovar"]:
        return (False, 0.0, "reprovou_caso_correto", "reprovou um caso correto: " + "; ".join(d.get("problemas") or [])[:200])
    return (False, 0.0, "aprovou_caso_com_defeito", f"aprovou um caso com defeito: {DEFEITOS[g['defeito']]} ({g['onde']})")


def _v_revisao(caso, cfg):
    """Revisor do caso de revisão: refaz o caso revisado pela semente e confirma, comparando o JSON direto (sem hash),
    que o defeito prometido existe de verdade — ou que o caso correto está mesmo correto."""
    import shutil
    import tempfile
    sub = caso["parametros"]["caso"]
    g = caso["gabarito"]
    tmp = Path(tempfile.mkdtemp(prefix="revisao_v_"))
    try:
        refeito = CATALOGO[sub["competencia"]].gerar(sub["nivel"], sub["semente"], tmp, cfg)
        mesmo_gabarito = json.dumps(refeito["gabarito"], sort_keys=True, default=str) == \
            json.dumps(sub["gabarito"], sort_keys=True, default=str)
        mesmo_arquivo = True
        if sub.get("arquivo"):
            mesmo_arquivo = conteudo_hash(sub["arquivo"], {}) == conteudo_hash(refeito.get("arquivo"), {})
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    problemas = []
    if g["defeito"] not in DEFEITOS:
        problemas.append(f"defeito desconhecido: {g['defeito']}")
    if g["deve_aprovar"] != (g["defeito"] == "nenhum"):
        problemas.append("deve_aprovar não combina com o defeito")
    if g["defeito"] == "nenhum":
        if not (mesmo_gabarito and mesmo_arquivo):
            problemas.append("o caso dito correto não é igual ao refeito pela semente")
        elif CATALOGO[sub["competencia"]].conferir(sub, cfg):
            problemas.append("o caso dito correto é reprovado pela própria conferência da competência")
    elif g["defeito"].startswith("gabarito") and mesmo_gabarito:
        problemas.append("prometeu gabarito errado, mas o gabarito é o certo")
    elif g["defeito"] == "conteudo" and mesmo_arquivo:
        problemas.append("prometeu arquivo alterado, mas o arquivo é igual ao original")
    return problemas


# =============================================================================== catálogo
CATALOGO: dict[str, Competencia] = {c.id: c for c in [
    Competencia("limpeza", "Limpar e barrar planilha", "diagnosticar_planilha", "casos_engenharia", "limpeza-4",
                {1: "título, R$, data em texto, duplicata, Total", 2: "+ nulos '-'/'N/A', subtotal, espaços; às vezes #REF!",
                 3: "+ número americano, datas misturadas, coluna vazia, rodapé", 4: "+ SAP '1.234,56-', '−' de PDF, 'R$ 1,2 mil'",
                 5: "sistema antigo: cabeçalho em duas linhas e repetido, linhas vazias, (1.234,56), 12 345,67; às vezes #REF!"},
                _g_limpeza, _x_limpeza, _c_limpeza, _v_limpeza, descricao="O pipeline real em diagnóstico: contar as linhas certas e barrar defeito."),
    Competencia("calculo", "Calcular o pedido", "calcular", "casos_analytics", "calculo-3",
                {1: "uma conta por segmento", 2: "conta com filtro", 3: "mediana por segmento", 4: "conta na sujeira nível 4",
                 5: "conta por dois segmentos na sujeira de sistema antigo"}, _g_calculo, _x_calculo, _c_calculo, _v_calculo,
                descricao="Entender o pedido em português e acertar o número de cada grupo."),
    Competencia("indicadores", "NPS com a base certa", "perfil_indicadores", "casos_analytics", "indicadores-4",
                {1: "NPS limpo", 2: "+ vazios", 3: "+ código 99 (não sabe)", 4: "+ sujeira de sistema", 5: "tudo junto + sistema antigo"},
                _g_indicadores, _x_indicadores, _c_indicadores, _v_indicadores, descricao="NPS e base (n) certos."),
    Competencia("lgpd", "Achar dado pessoal", "dados_pessoais", "casos_governanca", "lgpd-2",
                {1: "colunas com nome óbvio", 2: "nomes neutros", 3: "formatos crus (+55, só dígitos)",
                 4: "+ isca com cara de CPF (dígito inválido)", 5: "+ e-mail escondido em texto livre"},
                _g_lgpd, _x_lgpd, _c_lgpd, _v_lgpd, descricao="Apontar exatamente as colunas com dado pessoal (nem mais, nem menos)."),
    Competencia("anomalias", "Achar anomalias", "anomalias", "casos_analytics", "anomalias-2",
                {1: "1 valor absurdo", 2: "vários em 2 colunas", 3: "distribuição assimétrica", 4: "+ 'N/A' e vírgula decimal",
                 5: "+ anomalias para baixo"}, _g_anomalias, _x_anomalias, _c_anomalias, _v_anomalias,
                descricao="Regra de Tukey (3×IQR): quantas fora da faixa em cada coluna."),
    Competencia("segmentacao", "Achar os perfis", "segmentacao", "casos_analytics", "segmentacao-3",
                {1: "2 perfis", 2: "3 perfis", 3: "3 perfis + coluna de ruído", 4: "4 perfis", 5: "4 perfis mais próximos"},
                _g_segmentacao, _x_segmentacao, _c_segmentacao, _v_segmentacao, descricao="Número certo de grupos plantados."),
    Competencia("drivers", "Achar o que move o NPS", "drivers", "casos_analytics", "drivers-1",
                {1: "3 atributos, pouco ruído", 2: "mais ruído", 3: "5 atributos", 4: "+ distrator correlacionado",
                 5: "+ vazios"}, _g_drivers, _x_drivers, _c_drivers, _v_drivers, descricao="O 1º driver é o plantado."),
    Competencia("conhecimento", "Achar o conhecimento certo", "conhecimento", "casos_governanca", "conhecimento-4",
                {1: "pergunta exata", 2: "palavras embaralhadas, sem acento", 3: "só 60% das palavras",
                 4: "+ erro de digitação", 5: "metade das palavras + erro"},
                _g_conhecimento, _x_conhecimento, _c_conhecimento, _v_conhecimento, usa_conhecimento_real=True,
                descricao="Recuperar a lição certa do caderno da empresa (não conta como 'aplicada')."),
    Competencia("historico", "Ler o histórico", "historico", "casos_engenharia", "historico-2",
                {1: "15 execuções", 2: "200 execuções", 3: "+ linhas esperando o CSV fechar", 4: "3.000 execuções",
                 5: "nomes com ';' e aspas + pendentes"}, _g_historico, _x_historico, _c_historico, _v_historico,
                descricao="As 10 últimas execuções, na ordem certa."),
    Competencia("ronda_ti", "Ronda de TI", "ronda_ti", "casos_ti", "ronda-2",
                {1: "painel sem senha", 2: "+ senha do banco no config", 3: "+ erros no log", 4: "+ código novo sem aviso",
                 5: "problemas sorteados (achar também o que NÃO existe)"}, _g_ronda_ti, _x_ronda_ti, _c_ronda_ti, _v_ronda_ti,
                descricao="Detectar exatamente os problemas plantados no computador de teste."),
    Competencia("orquestracao", "SLA das rotinas (Airflow)", "orquestracao", "casos_ti", "orquestracao-2",
                {1: "uma rotina com erro", 2: "+ uma atrasada", 3: "+ uma lenta", 4: "+ iscas: erro já resolvido, lentidão antiga, "
                 "histórico curto", 5: "quantidades sorteadas + rotina perto do limite (achar também o que NÃO é problema)"},
                _g_orquestracao, _x_orquestracao, _c_orquestracao, _v_orquestracao,
                descricao="Apontar exatamente as rotinas atrasadas, com erro e lentas pela regra de SLA — nenhuma isca."),
    Competencia("resumo_executivo", "Resumo sem inventar", "resumo_executivo", "casos_governanca", "resumo-2",
                {1: "2 resultados", 2: "3 resultados", 3: "+ uma etapa que falhou", 4: "+ etapa vazia", 5: "8 resultados"},
                _g_resumo, _x_resumo, _c_resumo, _v_resumo, descricao="Todas as partes, nenhum número que não veio das etapas."),
    Competencia("desempenho", "Conta de desempenho (RH)", "desempenho", "casos_governanca", "desempenho-3",
                {1: "4 pessoas", 2: "12 pessoas", 3: "+ eventos fora da janela de 7 dias", 4: "+ tipos que não contam",
                 5: "+ exigência de 90% no limite"}, _g_desempenho, _x_desempenho, _c_desempenho, _v_desempenho,
                descricao="Taxa e situação de cada pessoa pela regra da política de RH."),
    Competencia("sql", "Mesma conta em SQL", "consulta_sql", "casos_engenharia", "sql-2",
                {1: "conta por segmento", 2: "com filtro", 3: "contagem", 4: "filtro com apóstrofo (d'Oeste)",
                 5: "dois segmentos"}, _g_sql, _x_sql, _c_sql, _v_sql, requer_banco=True,
                descricao="O SQL gerado roda no PostgreSQL (somente leitura) e bate com a verdade."),
    Competencia("revisao_casos", "Revisar caso de treino", "revisar_caso", "casos_governanca", "revisao-1",
                {1: "gabarito grosseiramente errado", 2: "gabarito com erro sutil", 3: "caso correto ou erro sutil "
                 "(aprovar o certo também conta)", 4: "arquivo alterado depois de gravado",
                 5: "mistura sorteada, com casos corretos"}, _g_revisao, _x_revisao, _c_revisao, _v_revisao,
                descricao="Sala de Casos: integridade, reprodutibilidade e conferência — reprovar o caso com defeito e aprovar o certo."),
]}
