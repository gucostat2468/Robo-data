"""Gravação das tabelas tratadas em Excel (formatado) e/ou CSV.

A gravação é "atômica": primeiro escreve um arquivo temporário e só no final troca pelo nome
definitivo. Assim o Power BI/OneDrive nunca enxerga um arquivo pela metade.
"""

from __future__ import annotations

import os
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LIMITE_LINHAS_EXCEL = 1_048_575
COR_CABECALHO = "1F3864"


def _nome_aba(nome: str, usados: set[str]) -> str:
    base = "".join(c for c in nome if c not in '[]:*?/\\')[:31] or "dados"
    nome_final, i = base, 2
    while nome_final.lower() in usados:
        sufixo = f"_{i}"
        nome_final = base[: 31 - len(sufixo)] + sufixo
        i += 1
    usados.add(nome_final.lower())
    return nome_final


def _neutralizar_formulas(ws):
    """Segurança (injeção de fórmula): texto vindo da planilha de origem que começa com '=' é gravado como TEXTO.
    A saída tratada nunca contém fórmulas ativas — uma célula "=HYPERLINK(...)" da origem não vira link clicável."""
    for linha in ws.iter_rows(min_row=1):
        for cel in linha:
            if cel.data_type == "f":
                cel.data_type = "s"


def _formatar_aba(ws, df: pd.DataFrame):
    _neutralizar_formulas(ws)
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=COR_CABECALHO)
        cel.alignment = Alignment(vertical="center")
    amostra = df.head(500)
    for i, col in enumerate(df.columns, 1):
        serie = df[col]
        letra = get_column_letter(i)
        if pd.api.types.is_datetime64_any_dtype(serie):
            com_hora = not (serie.dropna() == serie.dropna().dt.normalize()).all()
            formato, largura = ("dd/mm/yyyy hh:mm", 17) if com_hora else ("dd/mm/yyyy", 12)
            for (cel,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cel.number_format = formato
        elif pd.api.types.is_float_dtype(serie):
            for (cel,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cel.number_format = "#,##0.00"
            largura = 12
        else:
            textos = amostra[col].dropna().astype(str)
            largura = int(textos.str.len().quantile(0.9)) + 2 if len(textos) else 10
        ws.column_dimensions[letra].width = max(10, min(60, max(largura, len(str(col)) + 3)))


def _gravar_xlsx(tabelas: dict[str, pd.DataFrame], destino: Path):
    temporario = destino.with_name(f"~robo_tmp_{destino.name}")
    usados: set[str] = set()
    try:
        with pd.ExcelWriter(temporario, engine="openpyxl") as writer:
            for nome, df in tabelas.items():
                aba = _nome_aba(nome, usados)
                df.to_excel(writer, sheet_name=aba, index=False)
                _formatar_aba(writer.sheets[aba], df)
        os.replace(temporario, destino)
    finally:
        if temporario.exists():
            temporario.unlink(missing_ok=True)


def _gravar_csv(df: pd.DataFrame, destino: Path, sep: str, decimal: str):
    temporario = destino.with_name(f"~robo_tmp_{destino.name}")
    try:
        df.to_csv(temporario, sep=sep, decimal=decimal, index=False, encoding="utf-8-sig", date_format="%Y-%m-%d %H:%M:%S")
        os.replace(temporario, destino)
    finally:
        if temporario.exists():
            temporario.unlink(missing_ok=True)


def gravar_tabelas(tabelas: dict[str, pd.DataFrame], pasta: Path, nome_base: str, cfg_saida: dict) -> tuple[list[Path], list[str]]:
    """Devolve (arquivos gravados, avisos)."""
    pasta.mkdir(parents=True, exist_ok=True)
    formato = str(cfg_saida.get("formato", "xlsx")).lower()
    avisos = []
    if formato in ("xlsx", "ambos") and any(len(df) > LIMITE_LINHAS_EXCEL for df in tabelas.values()):
        avisos.append("Tabela maior que o limite do Excel (1.048.576 linhas): gravada em CSV")
        formato = "csv"

    gravados: list[Path] = []
    if formato in ("xlsx", "ambos"):
        destino = pasta / f"{nome_base}.xlsx"
        _gravar_xlsx(tabelas, destino)
        gravados.append(destino)
    if formato in ("csv", "ambos"):
        sep, dec = cfg_saida.get("csv_separador", ";"), cfg_saida.get("csv_decimal", ",")
        for nome, df in tabelas.items():
            destino = pasta / (f"{nome_base}.csv" if len(tabelas) == 1 else f"{nome_base}__{nome}.csv")
            _gravar_csv(df, destino, sep, dec)
            gravados.append(destino)
    if not gravados:
        raise ValueError(f"Formato de saída inválido: '{formato}' (use xlsx, csv ou ambos)")
    return gravados, avisos
