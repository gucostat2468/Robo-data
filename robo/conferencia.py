"""Confiabilidade dos resultados: todo cálculo vem com "como fazer à mão" e, quando possível, conferido no banco.

Duas garantias, sem IA e sem internet:
  1. Modo manual — para cada cálculo, o passo a passo para você refazer sozinho: a fórmula do Excel (CONT.SES,
     SOMASES, MÉDIASES…) e a consulta SQL equivalente. Assim nenhum número depende de confiar no robô.
  2. Dupla checagem — com o PostgreSQL ligado, a mesma conta é refeita DENTRO do banco, por um caminho diferente
     (SQL em vez de pandas), só com as linhas do mesmo arquivo. Os dois resultados são comparados grupo a grupo:
       conferido      → os dois caminhos dão o mesmo número;
       divergente     → não batem: o resultado é marcado como NÃO confiável e vira handover para a Dora (DBA);
       nao_aplicavel  → não dá para conferir no banco (e o motivo é dito: banco desligado, arquivo anexado,
                        coluna que só existe na dimensão, filtro de texto…). Nunca é dito "conferido" sem ter sido.
"""

from __future__ import annotations

import datetime as dt
import math

import pandas as pd

OPS_SQL = {
    "contagem": "COUNT(*)",
    "soma": "SUM({c})", "media": "AVG({c})", "minimo": "MIN({c})", "maximo": "MAX({c})",
    "mediana": "percentile_cont(0.5) WITHIN GROUP (ORDER BY {c})", "desvio": "stddev_samp({c})",
    "distintos": "COUNT(DISTINCT {c})",
    "nps": "100.0 * (COUNT(*) FILTER (WHERE {c} >= 9) - COUNT(*) FILTER (WHERE {c} <= 6)) / NULLIF(COUNT({c}), 0)",
}
EXCEL = {  # fórmula do Excel em português (o mesmo vale em inglês com os nomes entre parênteses)
    "contagem": ("CONT.SES (COUNTIFS)", "=CONT.SES({filtros_excel})"),
    "soma": ("SOMASES (SUMIFS)", "=SOMASES(Tabela[{c}]{sep}{filtros_excel})"),
    "media": ("MÉDIASES (AVERAGEIFS)", "=MÉDIASES(Tabela[{c}]{sep}{filtros_excel})"),
    "minimo": ("MÍNIMOSES (MINIFS)", "=MÍNIMOSES(Tabela[{c}]{sep}{filtros_excel})"),
    "maximo": ("MÁXIMOSES (MAXIFS)", "=MÁXIMOSES(Tabela[{c}]{sep}{filtros_excel})"),
}


def _ident(nome: str) -> str:
    return '"' + str(nome).replace('"', '""') + '"'


def _literal(v) -> str:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return repr(float(v)) if not float(v).is_integer() else str(int(v))
    return "'" + str(v).replace("'", "''") + "'"


def _escala_csat(serie: pd.Series) -> int:
    v = pd.to_numeric(serie, errors="coerce").dropna()
    return 5 if not len(v) or v.max() <= 5 else 7 if v.max() <= 7 else 10


# --------------------------------------------------------------------------- valores no pandas (por grupo)
def _chave(k) -> str:
    if isinstance(k, tuple):
        return " / ".join(_chave(x) for x in k)
    if isinstance(k, (pd.Timestamp, dt.datetime)):
        k = pd.Timestamp(k)
        return k.strftime("%Y-%m-%d") if k == k.normalize() else k.isoformat()
    if isinstance(k, dt.date):
        return k.isoformat()
    try:
        f = float(k)
        return f"{f:.10g}"
    except (TypeError, ValueError):
        return str(k).strip()


def valores_pandas(op: str, df: pd.DataFrame, col: str | None, por: list[str], escala: int = 5) -> dict[str, float]:
    def um(sub: pd.DataFrame) -> float:
        if op == "contagem":
            return float(len(sub))
        s = pd.to_numeric(sub[col], errors="coerce").dropna() if col else None
        if op == "distintos":
            return float(sub[col].nunique())
        if op == "nps":
            return 100 * ((s >= 9).sum() - (s <= 6).sum()) / len(s) if len(s) else float("nan")
        if op == "csat":
            return 100 * (s >= escala - 1).sum() / len(s) if len(s) else float("nan")
        return float({"soma": s.sum, "media": s.mean, "minimo": s.min, "maximo": s.max, "mediana": s.median,
                      "desvio": s.std}[op]())
    if not por:
        return {"(total)": um(df)}
    return {_chave(k): um(sub) for k, sub in df.groupby(por, dropna=True, observed=True)}


# --------------------------------------------------------------------------- SQL e Excel equivalentes
def montar_sql(op: str, col: str | None, por: list[str], filtros: list, tabela: str, arquivo: str | None,
               escala: int = 5, col_arquivo: str = '"_arquivo_origem"') -> str | None:
    if op == "csat":
        expr = f"100.0 * COUNT(*) FILTER (WHERE {_ident(col)} >= {escala - 1}) / NULLIF(COUNT({_ident(col)}), 0)"
    elif op in OPS_SQL:
        expr = OPS_SQL[op].format(c=_ident(col) if col else "*")
    else:
        return None
    onde = [f"{_ident(c)} {s} {_literal(v)}" for c, s, v in filtros]
    if arquivo:
        onde.append(f'{col_arquivo} = {_literal(arquivo)}')
    onde += [f"{_ident(c)} IS NOT NULL" for c in por]
    grupo = ", ".join(_ident(c) for c in por)
    sql = f"SELECT {grupo + ', ' if grupo else ''}{expr} AS valor, COUNT(*) AS linhas\nFROM {tabela}"
    if onde:
        sql += "\nWHERE " + "\n  AND ".join(onde)
    if grupo:
        sql += f"\nGROUP BY {grupo}\nORDER BY linhas DESC"
    return sql


def montar_excel(op: str, col: str | None, por: list[str], filtros: list, exemplo_grupo: dict | None,
                 escala: int = 5, dimensoes: dict | None = None) -> str:
    """Passo a passo no Excel (com a base como Tabela, Ctrl+T, chamada 'Tabela').
    dimensoes = {coluna: (aba_da_dimensao, chave)} para colunas que não estão na aba da fato."""
    crit = []
    for c, s, v in filtros:
        valor = f'"{s if s != "=" else ""}{v}"' if s != "=" or not isinstance(v, (int, float)) else str(v)
        crit.append(f"Tabela[{c}];{valor}")
    for c, v in (exemplo_grupo or {}).items():
        crit.append(f'Tabela[{c}];"{v}"')
    filtros_excel = ";".join(crit)
    passos = ["1. Abra a planilha publicada, clique numa célula da base e aperte Ctrl+T (vira uma Tabela); "
              "na guia Design da Tabela, dê o nome 'Tabela'."]
    for c, (aba, chave) in (dimensoes or {}).items():
        passos.append(f"1b. A coluna {c} está na aba {aba}. Faça dela uma Tabela chamada {aba} e, na Tabela principal, crie a "
                      f"coluna {c} com: =PROCX([@{chave}];{aba}[{chave}];{aba}[{c}])")
    if op == "nps":
        f = ";" + filtros_excel if filtros_excel else ""
        passos += [f"2. Promotores: =CONT.SES(Tabela[{col}];\">=9\"{f})",
                   f"3. Detratores: =CONT.SES(Tabela[{col}];\"<=6\"{f})",
                   f"4. Respondentes: =CONT.SES(Tabela[{col}];\">=0\"{f})",
                   "5. NPS: =ARRED((Promotores − Detratores) / Respondentes * 100; 0) — use as células dos passos 2 a 4."]
    elif op == "csat":
        f = ";" + filtros_excel if filtros_excel else ""
        passos += [f"2. Satisfeitos: =CONT.SES(Tabela[{col}];\">={escala - 1}\"{f})",
                   f"3. Respondentes: =CONT.SES(Tabela[{col}];\">=0\"{f})",
                   "4. CSAT %: =Satisfeitos / Respondentes (formate como porcentagem)."]
    elif op in EXCEL:
        nome, modelo = EXCEL[op]
        if op == "contagem" and not filtros_excel:
            formula = "=LINS(Tabela)"
        else:
            formula = modelo.format(c=col, sep=";" if filtros_excel else "", filtros_excel=filtros_excel)
        passos.append(f"2. {nome}: {formula}")
    elif op == "mediana":
        passos.append(f"2. =MED(Tabela[{col}]) — com filtro ou grupo: =MED(FILTRO(Tabela[{col}]; Tabela[coluna]=\"valor\"))")
    elif op == "desvio":
        passos.append(f"2. =DESVPAD.A(Tabela[{col}]) (com filtros, use FILTRO como na mediana)")
    elif op == "distintos":
        passos.append(f"2. =CONT.VALORES(ÚNICO(Tabela[{col}]))")
    else:
        passos.append("2. Use uma Tabela Dinâmica: Inserir → Tabela Dinâmica → arraste a coluna para Linhas e para Valores.")
    if por:
        ex = ", ".join(f"{k} = {v}" for k, v in (exemplo_grupo or {}).items()) or "o primeiro grupo"
        passos.append(f"Para cada grupo de {', '.join(por)}: troque o valor do grupo na fórmula (o exemplo acima usa "
                      f"{ex}), ou faça uma Tabela Dinâmica com {', '.join(por)} em Linhas.")
    return "\n".join(passos)


# --------------------------------------------------------------------------- conferência no PostgreSQL
def _colunas(conn, schema: str, tabela: str) -> list[str]:
    return [r[0] for r in conn.execute("SELECT column_name FROM information_schema.columns WHERE table_schema=%s "
                                       "AND table_name=%s ORDER BY ordinal_position", (schema, tabela))]


def _tabela_no_banco(cfg, banco, base: str, arquivo: str, colunas: list[str]) -> tuple[str | None, str, str]:
    """FROM (fato, com JOIN nas dimensões quando preciso) que recebeu este arquivo e tem todas as colunas usadas.
    Devolve (from_sql, coluna_do_arquivo, motivo_se_não_achou)."""
    from robo import historico
    exe = None
    for h in historico.ler(cfg.logs, 3000):
        if h.get("arquivo") == arquivo and h.get("status") == "sucesso":
            exe = historico.ler_execucao(cfg.logs, h.get("id") or "")
            break
    schema_t = banco.schemas["tratado"]
    tabelas = [str(t).split(".", 1)[1] for t in (exe or {}).get("tabelas_banco") or [] if str(t).startswith(schema_t + ".")]
    if not tabelas:
        return None, "", "este arquivo não foi carregado no PostgreSQL"
    with banco.conectar() as conn:
        cols = {t: _colunas(conn, schema_t, t) for t in tabelas}
    fatos = [t for t in tabelas if "_arquivo_origem" in cols[t] and not t.startswith("dim_")]
    for fato in fatos:
        faltam = [c for c in colunas if c not in cols[fato]]
        juncoes = []
        for dim in [t for t in tabelas if t.startswith("dim_") and t != "dim_calendario"]:
            usa = [c for c in faltam if c in cols[dim]]
            chave = next((c for c in cols[dim] if c in cols[fato] and not c.startswith("_")), None)
            if usa and chave:
                juncoes.append(f"JOIN {_ident(schema_t)}.{_ident(dim)} d{len(juncoes)} USING ({_ident(chave)})")
                faltam = [c for c in faltam if c not in usa]
        if not faltam:
            return (f"{_ident(schema_t)}.{_ident(fato)} f" + "".join("\n" + j for j in juncoes), 'f."_arquivo_origem"', "")
    return None, "", ("as colunas usadas (" + ", ".join(colunas) + ") não estão na fato nem nas dimensões do banco")


def dimensoes_do_arquivo(cfg, base: str, colunas: list[str]) -> dict:
    """{coluna: (aba_dimensão, chave)} das colunas que, na planilha publicada, estão numa aba dim_*."""
    try:
        from openpyxl import load_workbook

        from robo import calculos
        saida = calculos.bases_publicadas(cfg)[base]["saidas"][0]
        if not saida.lower().endswith(".xlsx"):
            return {}
        wb = load_workbook(saida, read_only=True)
        try:
            cab = {ws.title: [c.value for c in next(ws.iter_rows(max_row=1))] for ws in wb.worksheets}
        finally:
            wb.close()
    except Exception:
        return {}
    fato = next((n for n in cab if not n.startswith("dim_") and n not in ("dicionario_dados", "log_qualidade")), None)
    if not fato:
        return {}
    saida_d = {}
    for c in colunas:
        if c in cab[fato]:
            continue
        for aba, cols in cab.items():
            if aba.startswith("dim_") and c in cols and cols and cols[0] in cab[fato]:
                saida_d[c] = (aba, cols[0])
                break
    return saida_d


def conferir(cfg, calc: dict) -> dict:
    """Refaz o cálculo no PostgreSQL e compara com o resultado do pandas. Nunca diz 'conferido' sem ter conferido."""
    op, por, filtros = calc["op"], list(calc["por"] or []), list(calc["filtros"] or [])
    col = (calc["metricas"] or [None])[0]
    dados = calc["dados"]
    escala = _escala_csat(dados[col]) if op == "csat" and col else 5
    exemplo = None
    if por and len(dados):
        linha = dados.dropna(subset=por)
        exemplo = {c: linha.iloc[0][c] for c in por} if len(linha) else None
    usadas = [c for c in [col, *por, *[f[0] for f in filtros]] if c]
    dims = dimensoes_do_arquivo(cfg, calc["base"], usadas) if calc.get("fonte") == "publicada" else {}
    excel = montar_excel(op, col, por, filtros, exemplo, escala, dims)
    base = {"excel": excel, "sql": montar_sql(op, col, por, filtros, "tratado.<tabela_da_base>", calc.get("arquivo"), escala)}

    def nao(motivo):
        return {**base, "status": "nao_aplicavel", "motivo": motivo}
    if op in ("percentual", "correlacao") or base["sql"] is None:
        return nao("este tipo de cálculo ainda não tem conferência no banco")
    if calc.get("fonte") != "publicada":
        return nao("o cálculo foi feito sobre um arquivo anexado, que não está no banco")
    if not cfg.banco.get("habilitado"):
        return nao("o PostgreSQL não está ligado nas Configurações")
    if any(not isinstance(v, (int, float)) for _, _, v in filtros):
        return nao("há filtro de texto: o banco compara acento e maiúscula de forma diferente da calculadora")
    from robo.banco import Banco
    try:
        banco = Banco(cfg.banco, cfg.base)
        tabela, col_arq, motivo = _tabela_no_banco(cfg, banco, calc["base"], calc["arquivo"], usadas)
        if not tabela:
            return nao(motivo)
        sql = montar_sql(op, col, por, filtros, tabela, calc["arquivo"], escala, col_arq)
        with banco.conectar() as conn:
            with conn.transaction():
                conn.execute("SET TRANSACTION READ ONLY")
                linhas = conn.execute(sql).fetchall()
            conn.rollback()
    except Exception as e:
        return nao(f"o banco não respondeu ({type(e).__name__}: {e})")
    no_banco = {(_chave(tuple(r[:-2])) if len(por) > 1 else _chave(r[0])) if por else "(total)": (float(r[-2]) if r[-2] is not None else float("nan"))
                for r in linhas}
    local = valores_pandas(op, dados, col, por, escala)
    diferencas = []
    for k in sorted(set(local) | set(no_banco)):
        a, b = local.get(k), no_banco.get(k)
        if a is None or b is None:
            diferencas.append(f"{k}: só existe {'na calculadora' if b is None else 'no banco'}")
        elif not (math.isnan(a) and math.isnan(b)) and not math.isclose(a, b, rel_tol=1e-6, abs_tol=1e-6):
            diferencas.append(f"{k}: calculadora {a:.4f} × banco {b:.4f}")
    return {**base, "sql": sql, "status": "divergente" if diferencas else "conferido", "grupos": len(local),
            "diferencas": diferencas[:20]}


def texto(resultado: dict) -> str:
    st = resultado["status"]
    if st == "conferido":
        cab = (f"✅ Conferido no PostgreSQL: a mesma conta refeita no banco, por outro caminho (SQL), deu o mesmo "
               f"resultado nos {resultado['grupos']} grupo(s).")
    elif st == "divergente":
        cab = ("⛔ DIVERGÊNCIA: a calculadora e o banco não bateram. Não use este número até a Dora (DBA) explicar. "
               + "; ".join(resultado["diferencas"][:5]))
    else:
        cab = f"ℹ Sem conferência no banco: {resultado['motivo']}. O número vem só da calculadora."
    return (f"{cab}\n\nPara fazer você mesmo no Excel:\n{resultado['excel']}\n\n"
            f"A mesma conta em SQL (PostgreSQL):\n{resultado['sql']}")
