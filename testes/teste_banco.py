"""Testes contra um PostgreSQL REAL. Rodam só se a variável ROBO_TESTE_PG estiver definida, ex.:
    set ROBO_TESTE_PG=host=localhost port=5432 user=postgres password=suasenha
"""

import os
import uuid

import pandas as pd
import pytest
from pathlib import Path
from openpyxl import Workbook

from conftest import REGRA_NPS
from robo.pipeline import ADIADO, REPROVADO, SUCESSO, processar_arquivo

DSN = os.environ.get("ROBO_TESTE_PG")
pytestmark = pytest.mark.skipif(not DSN, reason="defina ROBO_TESTE_PG para testar com PostgreSQL")


def _dsn_dict():
    return dict(p.split("=", 1) for p in DSN.split())


@pytest.fixture
def banco_cfg():
    import psycopg
    d = _dsn_dict()
    nome = f"robo_teste_{uuid.uuid4().hex[:8]}"
    cfg = {"habilitado": True, "host": d.get("host", "localhost"), "porta": int(d.get("port", 5432)),
           "usuario": d.get("user", "postgres"), "senha": d.get("password", ""), "banco": nome}
    yield cfg
    with psycopg.connect(host=cfg["host"], port=cfg["porta"], user=cfg["usuario"], password=cfg["senha"],
                         dbname="postgres", autocommit=True) as c:
        c.execute(f'DROP DATABASE IF EXISTS "{nome}" WITH (FORCE)')


def _q(cfg, sql, params=()):
    import psycopg
    with psycopg.connect(host=cfg["host"], port=cfg["porta"], user=cfg["usuario"], password=cfg["senha"],
                         dbname=cfg["banco"]) as c:
        return c.execute(sql, params).fetchall()


def _planilha(pasta, nome, linhas, cabecalho=("ID Resposta", "Data Resposta", "Nota NPS", "Curso", "Modalidade", "UF")):
    wb = Workbook()
    ws = wb.active
    ws.append(list(cabecalho))
    for l in linhas:
        ws.append(list(l))
    caminho = pasta / "Bases-Limpeza" / nome
    wb.save(caminho)
    return caminho


def teste_elt_completo_bruto_tratado_analitico(planilha_baguncada, cfg_factory, banco_cfg, tmp_path):
    modelos = tmp_path / "modelos"
    modelos.mkdir()
    (modelos / "01_nps_por_curso.sql").write_text("""
        CREATE OR REPLACE VIEW analitico.vw_nps_por_curso AS
        SELECT c.curso, count(*) AS respostas,
               round(100.0 * (count(*) FILTER (WHERE f.nota_nps >= 9) - count(*) FILTER (WHERE f.nota_nps <= 6)) / count(*), 1) AS nps
        FROM tratado.fato_nps f JOIN tratado.dim_curso c USING (id_curso) GROUP BY c.curso;""", encoding="utf-8")
    (modelos / "_rascunho.sql").write_text("ISSO NAO E SQL", encoding="utf-8")  # ignorado (começa com _)

    regra = dict(REGRA_NPS, modelo={"fato": "fato_nps", "calendario": "data_resposta",
                                    "dimensoes": {"dim_curso": {"colunas": ["curso", "modalidade"], "chave": "id_curso"}}})
    cfg = cfg_factory([regra], banco={**banco_cfg, "pasta_modelos": str(modelos)})
    r = processar_arquivo(planilha_baguncada, cfg)
    assert r.status == SUCESSO, r.relatorio.texto()

    # bruto: exatamente o que veio (6 linhas de dados incluindo duplicada e Total), tudo texto, com linha de origem
    bruto = _q(banco_cfg, "SELECT _linha_origem, valor_mensalidade FROM bruto.nps_outubro ORDER BY 1")
    assert len(bruto) == 6
    assert bruto[0] == (4, "R$ 1.234,56")
    # tratado: tipado e limpo
    fato = _q(banco_cfg, "SELECT id_resposta, nota_nps, valor_mensalidade, data_resposta FROM tratado.fato_nps ORDER BY 1")
    assert len(fato) == 4
    assert fato[0][1] == 10 and fato[0][2] == 1234.56
    tipos = dict(_q(banco_cfg, "SELECT column_name, data_type FROM information_schema.columns "
                               "WHERE table_schema='tratado' AND table_name='fato_nps'"))
    assert tipos["nota_nps"] == "bigint" and tipos["data_resposta"] == "date"
    # analítico: view criada pelo modelo SQL
    nps = dict(_q(banco_cfg, "SELECT curso, nps FROM analitico.vw_nps_por_curso"))
    assert float(nps["Direito"]) == pytest.approx(33.3)
    # controle
    lotes = _q(banco_cfg, "SELECT status, linhas_saida FROM robo.lotes")
    assert lotes == [("sucesso", 4)]
    assert _q(banco_cfg, "SELECT count(*) FROM tratado.dim_calendario")[0][0] == 365
    # 2º andar: análise gravada no schema analise, ligada ao lote
    nps = _q(banco_cfg, "SELECT valor, n FROM analise.indicadores WHERE indicador = 'NPS' AND segmento = 'Total'")
    assert nps and nps[0][1] == 4
    assert _q(banco_cfg, "SELECT count(*) FROM analise.resumo")[0][0] >= 2
    assert _q(banco_cfg, "SELECT DISTINCT lote_id FROM analise.perfil") == [(r.id,)]


def teste_reprocessar_mesmo_arquivo_nao_duplica(pastas, cfg_factory, banco_cfg):
    cfg = cfg_factory([], banco=banco_cfg)
    for _ in range(2):
        c = _planilha(pastas, "vendas.xlsx", [(1, "01/10/2026", 9, "Direito", "EAD", "SP")])
        assert processar_arquivo(c, cfg).status == SUCESSO
    assert _q(banco_cfg, "SELECT count(*) FROM tratado.vendas")[0][0] == 1
    assert _q(banco_cfg, "SELECT count(*) FROM bruto.vendas")[0][0] == 1
    assert _q(banco_cfg, "SELECT count(*) FROM robo.lotes")[0][0] == 2


def teste_dimensao_mantem_chaves_entre_arquivos(pastas, cfg_factory, banco_cfg):
    regra = {"padrao": "nps_*", "destino": {"tabela": "nps"},
             "modelo": {"fato": "fato_nps", "dimensoes": {"dim_curso": {"colunas": ["curso"], "chave": "id_curso"}}}}
    cfg = cfg_factory([regra], banco=banco_cfg)
    c1 = _planilha(pastas, "nps_set.xlsx", [(1, "01/09/2026", 9, "Medicina", "EAD", "SP"),
                                            (2, "02/09/2026", 5, "Direito", "EAD", "SP")])
    assert processar_arquivo(c1, cfg).status == SUCESSO
    c2 = _planilha(pastas, "nps_out.xlsx", [(3, "01/10/2026", 10, "Direito", "EAD", "SP"),
                                            (4, "02/10/2026", 8, "Psicologia", "EAD", "SP")])
    assert processar_arquivo(c2, cfg).status == SUCESSO
    dim = dict(_q(banco_cfg, "SELECT curso, id_curso FROM tratado.dim_curso"))
    assert len(dim) == 3
    juntos = _q(banco_cfg, "SELECT f.id_resposta, d.curso FROM tratado.fato_nps f JOIN tratado.dim_curso d "
                           "USING (id_curso) ORDER BY 1")
    assert [c for _, c in juntos] == ["Medicina", "Direito", "Direito", "Psicologia"]


def teste_mesclar_por_chave(pastas, cfg_factory, banco_cfg):
    regra = {"padrao": "clientes*", "destino": {"tabela": "clientes", "modo": "mesclar", "chave": ["id_cliente"]}}
    cfg = cfg_factory([regra], banco=banco_cfg)
    cab = ("ID Cliente", "Nome", "Plano")
    assert processar_arquivo(_planilha(pastas, "clientes_1.xlsx", [(1, "Ana", "A"), (2, "Bia", "B")], cab), cfg).status == SUCESSO
    assert processar_arquivo(_planilha(pastas, "clientes_2.xlsx", [(2, "Bia", "C"), (3, "Caio", "A")], cab), cfg).status == SUCESSO
    assert _q(banco_cfg, "SELECT id_cliente, plano FROM tratado.clientes ORDER BY 1") == [("1", "A"), ("2", "C"), ("3", "A")]


def teste_conflito_de_tipo_reprova_e_nao_publica(pastas, cfg_factory, banco_cfg):
    cfg = cfg_factory([{"padrao": "metas*", "destino": {"tabela": "metas"}}], banco=banco_cfg)
    cab = ("Mes", "Meta")
    assert processar_arquivo(_planilha(pastas, "metas_1.xlsx", [("jan", 10), ("fev", 20)], cab), cfg).status == SUCESSO
    r = processar_arquivo(_planilha(pastas, "metas_2.xlsx", [("mar", "alta"), ("abr", "baixa")], cab), cfg)
    assert r.status == REPROVADO
    assert "Conflito de tipo" in r.relatorio.texto()
    assert not (cfg.tratadas / "metas_2.xlsx").exists()
    assert _q(banco_cfg, "SELECT count(*) FROM tratado.metas")[0][0] == 2
    # bruto guarda tudo o que chegou (auditoria), inclusive o reprovado; tratado não
    assert _q(banco_cfg, "SELECT count(*) FROM bruto.metas WHERE _arquivo_origem = 'metas_2.xlsx'")[0][0] == 2
    assert ("reprovado",) in _q(banco_cfg, "SELECT status FROM robo.lotes")
    erros = _q(banco_cfg, "SELECT arquivo, nivel, mensagem FROM robo.ocorrencias WHERE nivel = 'ERRO'")
    assert erros and erros[0][0] == "metas_2.xlsx" and "Conflito de tipo" in erros[0][2]


def teste_alargamento_de_tipo_e_coluna_nova(pastas, cfg_factory, banco_cfg):
    cfg = cfg_factory([{"padrao": "m*", "destino": {"tabela": "m"}}], banco=banco_cfg)
    assert processar_arquivo(_planilha(pastas, "m1.xlsx", [("a", 10)], ("Item", "Valor")), cfg).status == SUCESSO
    assert processar_arquivo(_planilha(pastas, "m2.xlsx", [("b", 2.5, "x")], ("Item", "Valor", "Extra")), cfg).status == SUCESSO
    tipos = dict(_q(banco_cfg, "SELECT column_name, data_type FROM information_schema.columns "
                               "WHERE table_schema='tratado' AND table_name='m'"))
    assert tipos["valor"] == "double precision" and tipos["extra"] == "text"


def teste_modelo_sql_com_erro_reprova(planilha_baguncada, cfg_factory, banco_cfg, tmp_path):
    (tmp_path / "ruim.sql").write_text("SELECT * FROM tabela_que_nao_existe;", encoding="utf-8")
    cfg = cfg_factory([REGRA_NPS], banco={**banco_cfg, "pasta_modelos": str(tmp_path)})
    r = processar_arquivo(planilha_baguncada, cfg)
    assert r.status == REPROVADO
    assert "ruim.sql" in r.relatorio.texto()


def teste_banco_fora_do_ar_adia_sem_mexer_no_arquivo(planilha_baguncada, cfg_factory, banco_cfg):
    cfg = cfg_factory([REGRA_NPS], banco={**banco_cfg, "porta": 1, "tempo_limite_conexao": 2})
    r = processar_arquivo(planilha_baguncada, cfg)
    assert r.status == ADIADO
    assert planilha_baguncada.exists()
    assert not (cfg.tratadas / "nps_outubro.xlsx").exists()


def teste_abas_empilhadas_carregam_no_banco_com_aba_e_linha(pastas, cfg_factory, banco_cfg):
    """Várias abas com o mesmo cabeçalho (uma por mês) + PostgreSQL: não pode travar, e cada linha guarda aba e linha."""
    from openpyxl import Workbook
    wb = Workbook()
    for i, mes in enumerate(["Jan", "Fev"]):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = mes
        ws.append(["ID", "_Interna", "Nota"])
        for j in range(1, 4):
            ws.append([i * 10 + j, "x", j])
    caminho = pastas / "Bases-Limpeza" / "mensal.xlsx"
    wb.save(caminho)
    r = processar_arquivo(caminho, cfg_factory([], banco=banco_cfg))
    assert r.status == SUCESSO, [o.mensagem for o in r.relatorio.erros]
    linhas = _q(banco_cfg, "SELECT _aba, _linha_origem FROM bruto.mensal ORDER BY _aba, _linha_origem")
    assert linhas == [("Fev", 2), ("Fev", 3), ("Fev", 4), ("Jan", 2), ("Jan", 3), ("Jan", 4)]
    assert _q(banco_cfg, "SELECT count(*) FROM tratado.mensal")[0][0] == 6


def teste_memoria_dos_funcionarios_vai_para_o_postgres(planilha_baguncada, cfg_factory, banco_cfg):
    from robo import memoria
    cfg = cfg_factory(regras=[{**REGRA_NPS, "valores_permitidos": {"uf": ["SP"]}}], banco=banco_cfg)
    r = processar_arquivo(planilha_baguncada, cfg)  # reprovado: gera erro, acerto de quem detectou e handover
    memoria.registrar_execucao(cfg, r)
    s = memoria.sincronizar_postgres(cfg)
    assert s["ok"] and s["experiencias"] > 0 and s["handovers"] > 0
    assert _q(banco_cfg, "SELECT count(*) FROM robo.memoria_experiencias")[0][0] == s["experiencias"]
    memoria.registrar_execucao(cfg, r)  # de novo: recorrência atualiza a mesma linha no PostgreSQL
    memoria.sincronizar_postgres(cfg)
    assert _q(banco_cfg, "SELECT max(recorrencias) FROM robo.memoria_handovers")[0][0] == 2
    # caderno de lições e cadeia de registros também vão para o banco (e a cópia bate com o SQLite)
    from robo import licoes
    licoes.rodada(cfg)
    s = memoria.sincronizar_postgres(cfg)
    assert s["licoes"] > 0 and s["licoes_funcionario"] > 0
    assert _q(banco_cfg, "SELECT count(*) FROM robo.memoria_registro_conhecimento")[0][0] == s["registro_conhecimento"]
    ultimo = _q(banco_cfg, "SELECT hash FROM robo.memoria_registro_conhecimento ORDER BY id DESC LIMIT 1")[0][0]
    assert ultimo == licoes.verificar_integridade(cfg)["ultimo_hash"]
    # histórico de cada erro e uso das ferramentas também vão para o banco
    assert s["handover_eventos"] > 0 and _q(banco_cfg, "SELECT count(*) FROM robo.memoria_handover_eventos")[0][0] == s["handover_eventos"]
    # registro imutável: adulterar o SQLite não sobrescreve a cópia do PostgreSQL
    import sqlite3
    c = sqlite3.connect(cfg.logs / "trabalho.db")
    c.execute("UPDATE experiencias SET resumo='ADULTERADO' WHERE id=1")
    c.commit()
    c.close()
    memoria.sincronizar_postgres(cfg)
    assert _q(banco_cfg, "SELECT resumo FROM robo.memoria_experiencias WHERE id=1")[0][0] != "ADULTERADO"


def teste_espelho_acompanha_coluna_nova_em_tabela_antiga_do_postgres(planilha_baguncada, cfg_factory, banco_cfg):
    """Banco criado por uma versão anterior (sem a coluna 'nivel' do laboratório): a cópia não pode falhar."""
    from robo import laboratorio, memoria
    cfg = cfg_factory([], banco=banco_cfg)
    laboratorio.exercicio(cfg, "varejo", 9, nivel=2)
    import psycopg
    with psycopg.connect(host=banco_cfg["host"], port=banco_cfg["porta"], user=banco_cfg["usuario"],
                         password=banco_cfg["senha"], dbname=banco_cfg["banco"]) as c:
        c.execute("CREATE SCHEMA IF NOT EXISTS robo")
        c.execute("CREATE TABLE robo.memoria_laboratorio_exercicios (id BIGINT PRIMARY KEY, rodada TEXT)")
    s = memoria.sincronizar_postgres(cfg)
    assert s["ok"] and s["laboratorio_exercicios"] > 0
    assert _q(banco_cfg, "SELECT DISTINCT nivel FROM robo.memoria_laboratorio_exercicios") == [(2,)]


def teste_correcao_aprovada_aparece_na_resposta_conferida_no_postgres(planilha_baguncada, cfg_factory, banco_cfg):
    from robo import aprendizado, ferramentas
    cfg = cfg_factory([REGRA_NPS], banco=banco_cfg)
    processar_arquivo(planilha_baguncada, cfg)
    pid = aprendizado.propor_apelido(cfg, "recomendacao", "nota_nps", "analise_estatistica")
    aprendizado.decidir(cfg, pid, True)
    r = ferramentas.executar("calcular", cfg, pergunta="média de recomendacao por curso", quem="analise_estatistica")
    assert r["dados"]["conferencia"] == "conferido" and r["texto"].count("Usei a correção que você aprovou") == 1
    assert "Usei a correção" in r["texto"].split("## ")[0]  # na parte visível, não numa seção recolhida


def teste_calculo_conferido_no_postgres_e_modo_manual(planilha_baguncada, cfg_factory, banco_cfg):
    from robo import conferencia, ferramentas
    cfg = cfg_factory([REGRA_NPS], banco=banco_cfg)
    assert processar_arquivo(planilha_baguncada, cfg).status == SUCESSO
    r = ferramentas.executar("calcular", cfg, pergunta="nps por curso")
    assert r["dados"]["conferencia"] == "conferido", r["texto"]
    assert "✅ Conferi o mesmo número no PostgreSQL" in r["texto"] and "CONT.SES(Tabela[nota_nps];\">=9\"" in r["texto"]
    assert "GROUP BY \"curso\"" in r["dados"]["sql"]
    for pergunta in ("média de valor_mensalidade por uf", "quantos por modalidade", "soma de valor_mensalidade",
                     "mediana de nota_nps"):
        assert ferramentas.executar("calcular", cfg, pergunta=pergunta)["dados"]["conferencia"] == "conferido", pergunta
    # divergência de verdade: alguém muda o banco por fora → o robô percebe e não chama o número de confiável
    import re
    tabela = re.search(r"FROM (\S+)", r["dados"]["sql"]).group(1)
    _q(banco_cfg, f"UPDATE {tabela} SET nota_nps = 0 WHERE curso = 'Direito' RETURNING 1")
    r = ferramentas.executar("calcular", cfg, pergunta="nps por curso")
    assert r["dados"]["conferencia"] == "divergente" and r["natureza"] == "hipotese" and "DIVERGÊNCIA" in r["texto"]
    from robo import memoria
    assert any(h["responsavel"] == "dba" for h in memoria.handovers(cfg))


def teste_dora_roda_select_livre_somente_leitura(planilha_baguncada, cfg_factory, banco_cfg):
    from robo import atendimento
    cfg = cfg_factory([REGRA_NPS], banco=banco_cfg)
    processar_arquivo(planilha_baguncada, cfg)
    r = atendimento.conversar(cfg, "dba", "SELECT count(*) AS total FROM robo.lotes")
    assert r["respostas"][-1]["natureza"] == "execucao" and "total" in r["respostas"][-1]["texto"]
    r = atendimento.conversar(cfg, "dba", "WITH x AS (DELETE FROM robo.lotes RETURNING 1) SELECT count(*) FROM x")
    assert r["respostas"][-1]["natureza"] == "limitacao"  # o próprio banco recusa: transação somente leitura
    assert _q(banco_cfg, "SELECT count(*) FROM robo.lotes")[0][0] >= 1
    r = atendimento.conversar(cfg, "infra_pleno", "SELECT 1")  # Rodrigo (TI) não tem SQL no contrato: encaminha
    assert r["respostas"][0]["natureza"] == "limitacao" and r["respostas"][1]["autor"] == "dba"
    assert r["respostas"][1]["natureza"] == "execucao"


def teste_conferencia_com_modelo_estrela_junta_dimensao(pastas, cfg_factory, banco_cfg):
    from robo import ferramentas
    regra = dict(REGRA_NPS, modelo={"fato": "fato_nps", "dimensoes": {"dim_curso": {"colunas": ["curso", "modalidade"],
                                                                                     "chave": "id_curso"}}})
    cfg = cfg_factory([regra], banco=banco_cfg)
    linhas = [(i, "01/10/2026", i % 11, ["Direito", "Medicina"][i % 2], ["EAD", "Presencial"][i % 2], "SP") for i in range(1, 80)]
    assert processar_arquivo(_planilha(pastas, "nps_estrela.xlsx", linhas), cfg).status == SUCESSO
    r = ferramentas.executar("calcular", cfg, pergunta="nps por modalidade")
    assert r["dados"]["conferencia"] == "conferido", r["texto"]
    assert "JOIN" in r["dados"]["sql"] and "PROCX([@id_curso];dim_curso[id_curso];dim_curso[modalidade])" in r["texto"]


def teste_laboratorio_no_postgres_em_schema_separado(cfg_factory, banco_cfg):
    from robo import laboratorio
    cfg = cfg_factory([], banco=banco_cfg)
    r = laboratorio.exercicio(cfg, "varejo", 5)
    assert r["falhas"] == [] and r["etapas"] == 6, r  # cabeçalho, limpeza, datas, cálculo, SQL e biblioteca
    etapas = laboratorio.resumo(cfg)["ultimas"]
    assert any(e["responsavel"] == "dba" and e["ok"] for e in etapas)
    # a tabela de treino é apagada depois; nada vai para o schema de produção
    assert _q(banco_cfg, "SELECT count(*) FROM information_schema.tables WHERE table_schema='laboratorio'")[0][0] == 0
    assert _q(banco_cfg, "SELECT count(*) FROM information_schema.tables WHERE table_schema='tratado'")[0][0] == 0


def teste_sql_livre_nao_escapa_do_somente_leitura(banco_cfg):
    """Achados da auditoria: ';' com SET + DELETE apagava dados; funções administrativas passavam."""
    from robo import consultas
    from robo.banco import Banco
    b = Banco(banco_cfg, Path("."))
    b.criar_banco_se_preciso()
    with b.conectar() as c:
        c.execute("CREATE TABLE alvo (x int)")
        c.execute("INSERT INTO alvo VALUES (1), (2)")
    for sql in ["SET transaction_read_only = off; DELETE FROM alvo RETURNING x",
                "SELECT pg_terminate_backend(pid) FROM pg_stat_activity", "SELECT pg_read_file('/etc/hostname')",
                "SELECT pg_reload_conf()", "SELECT set_config('role','postgres',true)", "COPY alvo TO STDOUT",
                "WITH x AS (DELETE FROM alvo RETURNING 1) SELECT * FROM x", "DELETE FROM alvo"]:
        with pytest.raises(Exception):
            consultas.executar_texto(b, sql)
    assert consultas.executar_texto(b, "SELECT 'a;b' AS t, count(*) FROM alvo -- comentário; ok")["linhas"] == [("a;b", 2)]
    assert _q(banco_cfg, "SELECT count(*) FROM alvo")[0][0] == 2


def teste_ti_no_computador_novo_nao_acusa_banco_ainda_nao_criado(banco_cfg, cfg_factory):
    """Achado no teste da operação: a ronda de TI da abertura rodava antes da 1ª carga criar o banco e dava vermelho."""
    from robo import ti
    cfg = cfg_factory(banco={**banco_cfg, "banco": "robo_ainda_nao_existe_x"})
    assert ti.chk_cloud(cfg) == []


def teste_primeiro_dia_rotinas_antes_da_primeira_carga(cfg_factory, banco_cfg):
    """Achado na demonstração zerada: ao abrir, backup e cópia da memória rodavam ANTES de a 1ª planilha criar o banco
    e viravam erro de sistema. Agora: backup diz que não há o que copiar (sem erro); a cópia cria o banco e copia."""
    from robo import backup, erros, memoria
    cfg = cfg_factory([], banco=banco_cfg)
    r = backup.fazer_backup(cfg)
    assert r.get("pulado") and "ainda não foi criado" in r["motivo"]
    s = memoria.sincronizar_postgres(cfg)
    assert s["ok"] and _q(banco_cfg, "SELECT 1")[0][0] == 1
    assert erros.listar(cfg) == []


def teste_treino_de_sql_nos_5_niveis_no_postgres(cfg_factory, banco_cfg):
    """Competência de SQL: a tabela sintética entra no schema 'treino', o SQL roda SOMENTE LEITURA e bate com a verdade
    (inclusive filtro com apóstrofo no nível 4); a tabela sai depois (é refeita pela semente)."""
    from robo import treino
    cfg = cfg_factory([], banco=banco_cfg)
    cfg.criar_pastas()
    assert "sql" in treino.competencias_de("dba", cfg)
    for nivel in range(1, 6):
        caso = treino.criar_caso(cfg, "sql", nivel, semente=300 + nivel)
        assert caso["estado"] == "validado", caso["revisao"]
        r = treino.tentar(cfg, "dba", "sql", caso, "T")
        assert r["acertou"], (nivel, r["detalhe"])
    assert _q(banco_cfg, "SELECT count(*) FROM information_schema.tables WHERE table_schema='treino'")[0][0] == 0


def teste_instalar_banco_real_com_usuario_so_leitura(tmp_path, cfg_factory, banco_cfg):
    """O instalador cria usuários e permissões de verdade; o de leitura não consegue escrever; nenhuma senha vai para o
    config.yaml; a memória inteira é copiada."""
    import psycopg
    import yaml
    from robo import instalacao_banco
    sufixo = banco_cfg["banco"].split("_")[-1]
    app, leitura = f"robo_app_{sufixo}", f"robo_leitura_{sufixo}"
    demo = f"{banco_cfg['banco']}_demo"
    base = cfg_factory()
    base.criar_pastas()
    caminho = tmp_path / "config.yaml"
    caminho.write_text(yaml.safe_dump(base.dados, allow_unicode=True), encoding="utf-8")
    try:
        r = instalacao_banco.instalar(caminho, banco_cfg["usuario"], banco_cfg["senha"], banco_cfg["host"], banco_cfg["porta"],
                                      banco_cfg["banco"], app, leitura, gravar_variaveis=False, banco_demonstracao=demo)
        assert r["ok"] and any("recusou" in p for p in r["passos"])
        b = yaml.safe_load(caminho.read_text(encoding="utf-8"))["banco"]
        assert b["habilitado"] and b["usuario"] == app and b["usuario_leitura"] == leitura and not b.get("senha")
        assert r["senhas"][app] not in caminho.read_text(encoding="utf-8")
        with psycopg.connect(host=banco_cfg["host"], port=banco_cfg["porta"], user=leitura, password=r["senhas"][leitura],
                             dbname=banco_cfg["banco"]) as c:
            assert c.execute("SELECT count(*) FROM robo.memoria_experiencias").fetchone()[0] >= 0  # lê
        with psycopg.connect(host=banco_cfg["host"], port=banco_cfg["porta"], user=app, password=r["senhas"][app],
                             dbname=demo, autocommit=True) as c:  # o modo demonstração tem banco próprio e o robo_app grava nele
            c.execute("CREATE TABLE robo.prova_demo (x int)")
        instalacao_banco.instalar(caminho, banco_cfg["usuario"], banco_cfg["senha"], banco_cfg["host"], banco_cfg["porta"],
                                  banco_cfg["banco"], app, leitura, gravar_variaveis=False, banco_demonstracao=demo)  # rodar de novo não estraga
    finally:
        with psycopg.connect(host=banco_cfg["host"], port=banco_cfg["porta"], user=banco_cfg["usuario"],
                             password=banco_cfg["senha"], dbname="postgres", autocommit=True) as c:
            for nome in (banco_cfg["banco"], demo):
                c.execute(f'DROP DATABASE IF EXISTS "{nome}" WITH (FORCE)')
            for papel in (app, leitura):
                c.execute(f'DROP ROLE IF EXISTS "{papel}"')


def teste_turnos_na_nuvem_continuam_de_onde_pararam(banco_cfg, tmp_path, monkeypatch):
    """Dois turnos em máquinas novas: o 2º restaura o instantâneo do 1º, continua a numeração e nada se perde."""
    import psycopg
    import sqlite3
    from robo import nuvem
    with psycopg.connect(host=banco_cfg["host"], port=banco_cfg["porta"], user=banco_cfg["usuario"],
                         password=banco_cfg["senha"], dbname="postgres", autocommit=True) as c:
        c.execute(f'CREATE DATABASE "{banco_cfg["banco"]}"')
    url = f"postgresql://{banco_cfg['usuario']}@{banco_cfg['host']}:{banco_cfg['porta']}/{banco_cfg['banco']}"
    monkeypatch.setattr(nuvem, "MEMORIA_INICIAL", tmp_path / "nao_existe.db.gz")
    cfg1 = nuvem.config_da_nuvem(tmp_path / "turno1", url)
    r1 = nuvem.turno(cfg1, minutos=0.02, pausa_s=0)
    assert r1["origem"] == "vazia" and r1["memoria"]["linhas"] > 0
    t1 = sqlite3.connect(cfg1.logs / "trabalho.db").execute("SELECT COUNT(*), MAX(id) FROM treino_tentativas").fetchone()
    cfg2 = nuvem.config_da_nuvem(tmp_path / "turno2", url)          # máquina nova: só o banco sobrevive
    r2 = nuvem.turno(cfg2, minutos=0.02, pausa_s=0)
    assert r2["origem"].startswith("instantaneo ")
    c2 = sqlite3.connect(cfg2.logs / "trabalho.db")
    t2 = c2.execute("SELECT COUNT(*), MAX(id) FROM treino_tentativas").fetchone()
    assert t2[0] >= t1[0] and c2.execute("SELECT COUNT(*) - COUNT(DISTINCT id) FROM treino_tentativas").fetchone()[0] == 0
    assert _q(banco_cfg, "SELECT COUNT(*) FROM robo.memoria_treino_tentativas")[0][0] == t2[0]   # a nuvem tem tudo
    assert _q(banco_cfg, "SELECT COUNT(*) FROM robo.instantaneos_sqlite")[0][0] == 2
