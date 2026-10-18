"""Regressão dos achados da auditoria de outubro/2026 (dados e ETL). Cada teste reproduz o defeito original."""

import shutil

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from robo import calculos, ciencia
from robo.leitura import ler_arquivo
from robo.pipeline import processar_arquivo


@pytest.fixture
def cfg(cfg_factory):
    c = cfg_factory()
    c.criar_pastas()
    return c


def _xlsx(caminho, linhas):
    wb = Workbook()
    for l in linhas:
        wb.active.append(l)
    wb.save(caminho)
    return caminho


def teste_nps_ignora_codigos_fora_da_escala():
    df = pd.DataFrame({"nota_nps": [10, 9, 8, 7, 6, 99, 99, 99, 0, 5, 9, 10, 10, 10, 9]})
    t = calculos.calcular(None, "nps", fontes={"b": (df, "b.xlsx")})["texto"]
    # válidas: 12 (sem os três 99): 7 promotores, 3 detratores → (7−3)/12 = 33,3
    assert "NPS 33 = 58,3% promotores (7) − 25,0% detratores (3)" in t and "3 valor(es) fora de 0–10" in t


def teste_tempo_de_atendimento_nao_e_csat():
    df = pd.DataFrame({"tempo_atendimento": [3, 4, 5, 12, 15, 9, 10, 2, 1, 30]})
    assert "Nenhuma base" in calculos.calcular(None, "csat", fontes={"b": (df, "b.xlsx")})["texto"]


def teste_pesquisa_sem_id_nao_perde_respostas_iguais(cfg):
    _xlsx(cfg.entrada / "pesquisa.xlsx", [["Curso", "Nota NPS"]] + [["A", n] for n in [10, 10, 10, 9, 6, 6, 8, 10]])
    r = processar_arquivo(cfg.entrada / "pesquisa.xlsx", cfg)
    assert r.status == "sucesso" and r.linhas_saida == 8
    assert any("MANTIDAS" in o.mensagem for o in r.relatorio.avisos)


def teste_total_como_categoria_nao_e_removido(cfg):
    _xlsx(cfg.entrada / "cob.xlsx", [["Cobertura", "Cliente", "Valor"]] +
          [[c, f"cli{i}", i * 10] for i, c in enumerate(["Total", "Parcial", "Total", "Parcial", "Total", "Parcial"])])
    r = processar_arquivo(cfg.entrada / "cob.xlsx", cfg)
    assert r.status == "sucesso" and r.linhas_saida == 6


def teste_data_impossivel_nao_derruba_o_arquivo(cfg):
    (cfg.entrada / "datas.csv").write_text("id;data;nota\n" + "".join(f"{i};01/02/2025;5\n" for i in range(40))
                                           + "99;01/02/0202;5\n", encoding="utf-8")
    r = processar_arquivo(cfg.entrada / "datas.csv", cfg)
    assert not any("OutOfBounds" in o.mensagem or "Falha inesperada" in o.mensagem for o in r.relatorio.ocorrencias)


def teste_cabecalho_com_anos_nao_vira_dado(tmp_path):
    arq = _xlsx(tmp_path / "v.xlsx", [["Relatório de vendas 2025"], ["Loja", 2023, 2024, 2025], ["A", 1, 2, 3], ["B", 4, 5, 6]])
    bruto = list(ler_arquivo(arq).values())[0] if isinstance(ler_arquivo(arq), dict) else None
    from robo.leitura import _localizar_cabecalho
    df = pd.DataFrame([["Relatório de vendas 2025", None, None, None], ["Loja", 2023, 2024, 2025], ["A", 1, 2, 3]])
    assert _localizar_cabecalho(df, 20) == 1
    assert bruto is not None


def teste_csv_com_titulo_e_utf16_sao_lidos(cfg):
    (cfg.entrada / "titulo.csv").write_text("Relatorio de vendas\nnome;valor;qtd\nA;1,5;2\nB;2,5;3\n", encoding="utf-8")
    assert processar_arquivo(cfg.entrada / "titulo.csv", cfg).status == "sucesso"
    (cfg.entrada / "unicode.txt").write_text("nome\tvalor\nAna\t10\nBia\t20\n", encoding="utf-16")
    r = processar_arquivo(cfg.entrada / "unicode.txt", cfg)
    assert r.status == "sucesso", [o.mensagem for o in r.relatorio.erros]


def teste_csv_de_saida_nao_leva_formula_ativa(cfg_factory):
    c = cfg_factory(saida={"formato": "csv"})
    c.criar_pastas()
    (c.entrada / "f.csv").write_text('Nome;Link;Fone\na;=HYPERLINK("http://x");+55 11 9999\nb;ok;-5\n', encoding="utf-8")
    r = processar_arquivo(c.entrada / "f.csv", c)
    texto = r.saidas[0].read_text(encoding="utf-8-sig")
    assert "'=HYPERLINK" in texto and "+55 11 9999" in texto and "'+55" not in texto


def teste_relatorio_de_qualidade_nao_leva_formula_ativa(cfg):
    (cfg.entrada / "q.csv").write_text("Nome;Obs\na;=1+1\nb;x\n", encoding="utf-8")
    processar_arquivo(cfg.entrada / "q.csv", cfg)
    wb = load_workbook(next(cfg.relatorios.glob("q*_qualidade.xlsx")))
    assert not any(c.data_type == "f" for ws in wb for row in ws.iter_rows() for c in row)


def teste_original_travado_nao_e_reprocessado(cfg, monkeypatch):
    import robo.pipeline as pl
    arq = _xlsx(cfg.entrada / "trava.xlsx", [["Nome", "Valor"], ["a", 1], ["b", 2]])
    real = shutil.move

    def travado(o, d):
        if str(o).endswith("trava.xlsx"):
            raise PermissionError("em uso pelo OneDrive")
        return real(o, d)
    monkeypatch.setattr(pl.shutil, "move", travado)
    monkeypatch.setattr(pl.time, "sleep", lambda s: None)
    r = processar_arquivo(arq, cfg)
    assert r.status == "sucesso" and r.destino_original.exists() and arq.exists()
    monkeypatch.setattr(pl.shutil, "move", real)
    assert pl.ja_publicado(cfg, arq) and not arq.exists()  # liberado: sai da entrada sem reprocessar


def teste_segmentacao_nao_estoura_memoria():
    import tracemalloc
    df = pd.DataFrame(np.random.default_rng(1).integers(0, 10, (3000, 20)), columns=[f"c{i}" for i in range(20)])
    tracemalloc.start()
    assert ciencia.segmentar(df)["ok"]
    pico = tracemalloc.get_traced_memory()[1]
    tracemalloc.stop()
    assert pico < 200e6
