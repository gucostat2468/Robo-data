"""Modelo estrela da empresa para o Power BI — gerado a partir da memória da empresa (SQLite) e do código.

    python powerbi/exportar_estrela.py [caminho/do/trabalho.db]

Produz, pelas MESMAS regras (para os dois baterem linha a linha):
  powerbi/dados/*.csv              uma tabela por arquivo (UTF-8 com BOM, separador ";", vírgula decimal → abre certo
                                   no Excel em português)
  powerbi/dados/modelo_estrela.xlsx  as mesmas tabelas, uma aba cada (para "Obter dados → Pasta de trabalho do Excel")
  powerbi/estrela_neon.sql         as views do schema bi no PostgreSQL (Neon) + a carga das dimensões que vivem no código

Por que o organograma e o catálogo de competências vão como tabela carregada: eles estão no CÓDIGO do robô
(robo/organizacao.py, robo/rh.py, robo/treino_casos.py), não no banco. As fatos são VIEWS: sempre ao vivo.
"""

from __future__ import annotations

import csv
import sqlite3
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

PASTA = Path(__file__).resolve().parent
DADOS = PASTA / "dados"

ORDEM_CARGO = {"estagiario": 1, "junior": 2, "pleno": 3, "senior": 4, "especialista": 5, "coordenador": 6, "gerente": 7,
               "head": 8, "diretor": 9}
# colunas das fatos que podem vir sem linha nenhuma (iguais às da view correspondente em estrela_neon.sql)
COLUNAS_SE_VAZIA = {"fato_rh_movimentos": ["movimento_id", "data_hora", "data", "funcionario_id", "tipo", "de_nivel", "para_nivel",
                                           "competencia_id", "situacao", "proposto_por", "validado_por", "motivo"]}


# =============================================================================== dimensões que vivem no código
def dimensoes_do_codigo(cfg) -> dict[str, list[dict]]:
    from robo import organizacao, rh
    from robo.treino_casos import CATALOGO, NIVEIS
    gente = organizacao.lista(cfg)
    nome = {p["id"]: p["nome"] for p in gente}
    funcionarios = []
    for p in gente:
        nivel = organizacao.nivel_de(p["id"], p) or ""
        gestor = organizacao.gestor_de(p["id"]) or ""
        funcionarios.append({
            "funcionario_id": p["id"], "nome": p["nome"], "cargo": (p.get("papel") or "").split(" — ")[0],
            "nivel_cargo": rh.NOME_NIVEL.get(nivel, nivel), "nivel_cargo_ordem": ORDEM_CARGO.get(nivel, 0),
            "andar": int(p.get("andar") or 0), "area": p.get("grupo") or "",
            "gestor_id": gestor, "gestor_nome": "Você (dono)" if gestor == "dono" else nome.get(gestor, gestor),
            "no_quadro": 1})
    competencias = [{"competencia_id": c.id, "competencia": c.nome, "ferramenta": c.ferramenta,
                     "designer_id": c.designer, "designer_nome": nome.get(c.designer, c.designer),
                     "descricao": c.descricao} for c in CATALOGO.values()]
    andares = [{"andar": a, "andar_nome": n, "andar_rotulo": f"{a}º · {n}"} for a, n in sorted(organizacao.NOME_ANDAR.items())]
    # andar 0 recebe quem aparece nas fatos mas saiu do organograma: sem ele, o Power BI mostra uma linha "(Em branco)"
    andares.append({"andar": 0, "andar_nome": "Fora do quadro atual", "andar_rotulo": "— · Fora do quadro atual"})
    niveis = [{"nivel": n, "nivel_nome": f"{n} · {d.capitalize()}"} for n, d in sorted(NIVEIS.items())]
    return {"dim_funcionario_base": funcionarios, "dim_competencia_base": competencias, "dim_andar": andares,
            "dim_nivel": niveis}


# =============================================================================== fatos (mesmas regras do SQL)
def _iso(t):
    return datetime.fromisoformat(str(t)) if t else None


def fatos(c: sqlite3.Connection) -> dict[str, list[dict]]:
    tabelas = {r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    out: dict[str, list[dict]] = {}
    out["fato_treino"] = [{
        "tentativa_id": r["id"], "data_hora": _iso(r["data_hora"]), "data": _iso(r["data_hora"]).date(),
        "hora": _iso(r["data_hora"]).hour, "funcionario_id": r["funcionario"], "andar": r["andar"],
        "cargo_nivel": r["cargo_nivel"], "competencia_id": r["competencia"], "nivel": r["nivel"],
        "acertou": int(r["acertou"] or 0), "errou": 1 - int(r["acertou"] or 0), "nota": r["nota"], "segundos": r["segundos"],
        "tipo_erro": r["tipo_erro"] or ("(sem erro)" if r["acertou"] else "(não informado)"), "rodada": r["rodada"],
        "versao_codigo": r["versao_codigo"]}
        for r in c.execute("SELECT * FROM treino_tentativas ORDER BY id")]
    out["fato_interacoes"] = [{
        "interacao_id": r["id"], "data_hora": _iso(r["data_hora"]), "data": _iso(r["data_hora"]).date(),
        "de_id": r["de"], "para_id": r["para"] or "", "tipo": r["tipo"], "texto": r["texto"], "ref": r["ref"] or ""}
        for r in c.execute("SELECT * FROM interacoes ORDER BY id")] if "interacoes" in tabelas else []
    out["fato_mentorias"] = [{
        "mentoria_id": r["id"], "data_hora": _iso(r["data_hora"]), "data": _iso(r["data_hora"]).date(),
        "mentor_id": r["mentor"], "aprendiz_id": r["aprendiz"], "competencia_id": r["competencia"], "nivel": r["nivel"],
        "situacao": r["situacao"], "ajudou": 1 if r["situacao"] == "ajudou" else 0,
        "fechada": 0 if r["situacao"] == "aberta" else 1, "fechada_em": _iso(r["fechada_em"])}
        for r in c.execute("SELECT * FROM mentorias ORDER BY id")] if "mentorias" in tabelas else []
    out["fato_rh_movimentos"] = [{
        "movimento_id": r["id"], "data_hora": _iso(r["data_hora"]), "data": _iso(r["data_hora"]).date(),
        "funcionario_id": r["funcionario"], "tipo": r["tipo"], "de_nivel": r["de_nivel"], "para_nivel": r["para_nivel"],
        "competencia_id": r["competencia"], "situacao": r["situacao"], "proposto_por": r["proposto_por"],
        "validado_por": r["validado_por"], "motivo": r["motivo"]}
        for r in c.execute("SELECT * FROM rh_movimentos ORDER BY id")] if "rh_movimentos" in tabelas else []
    out["fato_rotinas"] = [{
        "execucao_id": r["id"], "data_hora": _iso(r["inicio"]), "data": _iso(r["inicio"]).date(), "rotina": r["nome"],
        "status": r["status"], "com_erro": 1 if r["status"] == "erro" else 0, "segundos": r["segundos"]}
        for r in c.execute("SELECT * FROM rotinas_execucoes ORDER BY id")]
    out["fato_evolucao_nivel"] = [{
        "evento_id": r["id"], "data_hora": _iso(r["data_hora"]), "data": _iso(r["data_hora"]).date(),
        "funcionario_id": r["funcionario"], "competencia_id": r["competencia"], "nivel_de": r["de"], "nivel_para": r["para"],
        "variacao": (r["para"] - r["de"]) if r["de"] is not None else 0,
        "tipo_evento": "inicial" if r["de"] is None else ("subiu" if r["para"] > r["de"] else "desceu"), "motivo": r["motivo"]}
        for r in c.execute("SELECT * FROM treino_niveis_eventos ORDER BY id")]
    out["fato_nivel_atual"] = [{"funcionario_id": r["funcionario"], "competencia_id": r["competencia"], "nivel": r["nivel"],
                                "desde": _iso(r["desde"])}
                               for r in c.execute("SELECT * FROM treino_niveis ORDER BY funcionario, competencia")]
    return out


def completar_dimensoes(dims: dict, fat: dict) -> dict:
    """Pessoas/competências que aparecem nas fatos mas não estão no código (ex.: contratados depois) entram marcadas."""
    base = {f["funcionario_id"] for f in dims["dim_funcionario_base"]}
    vistos = set()
    for nome, chaves in (("fato_treino", ["funcionario_id"]), ("fato_interacoes", ["de_id", "para_id"]),
                         ("fato_mentorias", ["mentor_id", "aprendiz_id"]), ("fato_rh_movimentos", ["funcionario_id"]),
                         ("fato_evolucao_nivel", ["funcionario_id"]), ("fato_nivel_atual", ["funcionario_id"])):
        for r in fat[nome]:
            vistos |= {r[k] for k in chaves if r.get(k)}
    extra = [{"funcionario_id": x, "nome": x, "cargo": "(fora do quadro atual)", "nivel_cargo": "", "nivel_cargo_ordem": 0,
              "andar": 0, "area": "", "gestor_id": "", "gestor_nome": "", "no_quadro": 0} for x in sorted(vistos - base)]
    comp_base = {c["competencia_id"] for c in dims["dim_competencia_base"]}
    comps = {r["competencia_id"] for n in ("fato_treino", "fato_mentorias", "fato_evolucao_nivel", "fato_nivel_atual")
             for r in fat[n] if r.get("competencia_id")}
    comp_extra = [{"competencia_id": x, "competencia": x, "ferramenta": "", "designer_id": "", "designer_nome": "",
                   "descricao": "(fora do catálogo atual)"} for x in sorted(comps - comp_base)]
    datas = [r["data"] for n in fat for r in fat[n] if r.get("data")]
    ini, fim = min(datas), max(datas)
    dias = [ini + timedelta(days=i) for i in range((fim - ini).days + 1)]
    meses = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]
    semana = ["seg", "ter", "qua", "qui", "sex", "sáb", "dom"]
    dim_data = [{"data": d, "ano": d.year, "mes": d.month, "mes_nome": meses[d.month - 1], "ano_mes": f"{d:%Y-%m}",
                 "dia": d.day, "dia_semana": d.isoweekday(), "dia_semana_nome": semana[d.weekday()],
                 "semana_iso": d.isocalendar()[1], "fim_de_semana": 1 if d.weekday() >= 5 else 0} for d in dias]
    return {"dim_funcionario": dims["dim_funcionario_base"] + extra, "dim_competencia": dims["dim_competencia_base"] + comp_extra,
            "dim_andar": dims["dim_andar"], "dim_nivel": dims["dim_nivel"], "dim_data": dim_data}


# =============================================================================== saída
def _txt(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return f"{v:.6f}".rstrip("0").rstrip(".").replace(".", ",")
    return str(v)


def gravar_csv_xlsx(tabelas: dict[str, list[dict]], destino: Path = DADOS) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    destino.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for nome, linhas in tabelas.items():
        # tabela vazia ainda leva o cabeçalho certo: o Power Query precisa das colunas para não quebrar a consulta
        cols = list(linhas[0].keys()) if linhas else COLUNAS_SE_VAZIA.get(nome, ["(vazia)"])
        with open(destino / f"{nome}.csv", "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(cols)
            for r in linhas:
                w.writerow([_txt(r.get(k)) for k in cols])
        ws = wb.create_sheet(nome[:31])
        ws.append(cols)
        for cel in ws[1]:
            cel.font = Font(bold=True, color="FFFFFF")
            cel.fill = PatternFill("solid", fgColor="461F84")
        for r in linhas:
            ws.append([r.get(k) for k in cols])
        ws.freeze_panes = "A2"
        for i, col in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(48, max(12, len(col) + 4))
    wb.save(destino / "modelo_estrela.xlsx")


def _sql_lit(v):
    if v is None:
        return "NULL"
    if isinstance(v, (int, float)):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


VIEWS_SQL = r"""
-- ============================================================ FATOS (views: sempre ao vivo, só leem o schema robo)
CREATE OR REPLACE VIEW bi.fato_treino AS
SELECT t.id AS tentativa_id, t.data_hora::timestamp AS data_hora, t.data_hora::timestamp::date AS data,
       EXTRACT(HOUR FROM t.data_hora::timestamp)::int AS hora, t.funcionario AS funcionario_id, t.andar::int AS andar,
       t.cargo_nivel, t.competencia AS competencia_id, t.nivel::int AS nivel,
       COALESCE(t.acertou, 0)::int AS acertou, (1 - COALESCE(t.acertou, 0))::int AS errou, t.nota, t.segundos,
       COALESCE(NULLIF(t.tipo_erro, ''), CASE WHEN t.acertou = 1 THEN '(sem erro)' ELSE '(não informado)' END) AS tipo_erro,
       t.rodada, t.versao_codigo
FROM robo.memoria_treino_tentativas t;

CREATE OR REPLACE VIEW bi.fato_interacoes AS
SELECT i.id AS interacao_id, i.data_hora::timestamp AS data_hora, i.data_hora::timestamp::date AS data,
       i.de AS de_id, COALESCE(i.para, '') AS para_id, i.tipo, i.texto, COALESCE(i.ref, '') AS ref
FROM robo.memoria_interacoes i;

CREATE OR REPLACE VIEW bi.fato_mentorias AS
SELECT m.id AS mentoria_id, m.data_hora::timestamp AS data_hora, m.data_hora::timestamp::date AS data,
       m.mentor AS mentor_id, m.aprendiz AS aprendiz_id, m.competencia AS competencia_id, m.nivel::int AS nivel, m.situacao,
       CASE WHEN m.situacao = 'ajudou' THEN 1 ELSE 0 END AS ajudou,
       CASE WHEN m.situacao = 'aberta' THEN 0 ELSE 1 END AS fechada, NULLIF(m.fechada_em, '')::timestamp AS fechada_em
FROM robo.memoria_mentorias m;

CREATE OR REPLACE VIEW bi.fato_rh_movimentos AS
SELECT r.id AS movimento_id, r.data_hora::timestamp AS data_hora, r.data_hora::timestamp::date AS data,
       r.funcionario AS funcionario_id, r.tipo, r.de_nivel, r.para_nivel, r.competencia AS competencia_id, r.situacao,
       r.proposto_por, r.validado_por, r.motivo
FROM robo.memoria_rh_movimentos r;

CREATE OR REPLACE VIEW bi.fato_rotinas AS
SELECT e.id AS execucao_id, e.inicio::timestamp AS data_hora, e.inicio::timestamp::date AS data, e.nome AS rotina,
       e.status, CASE WHEN e.status = 'erro' THEN 1 ELSE 0 END AS com_erro, e.segundos
FROM robo.memoria_rotinas_execucoes e;

CREATE OR REPLACE VIEW bi.fato_evolucao_nivel AS
SELECT v.id AS evento_id, v.data_hora::timestamp AS data_hora, v.data_hora::timestamp::date AS data,
       v.funcionario AS funcionario_id, v.competencia AS competencia_id, v.de::int AS nivel_de, v.para::int AS nivel_para,
       COALESCE(v.para - v.de, 0)::int AS variacao,
       CASE WHEN v.de IS NULL THEN 'inicial' WHEN v.para > v.de THEN 'subiu' ELSE 'desceu' END AS tipo_evento, v.motivo
FROM robo.memoria_treino_niveis_eventos v;

CREATE OR REPLACE VIEW bi.fato_nivel_atual AS
SELECT n.funcionario AS funcionario_id, n.competencia AS competencia_id, n.nivel::int AS nivel,
       n.desde::timestamp AS desde
FROM robo.memoria_treino_niveis n;

-- ============================================================ DIMENSÕES
CREATE OR REPLACE VIEW bi.dim_funcionario AS
SELECT funcionario_id, nome, cargo, nivel_cargo, nivel_cargo_ordem, andar, area, gestor_id, gestor_nome, no_quadro
FROM bi.dim_funcionario_base
UNION ALL
SELECT x.id, x.id, '(fora do quadro atual)', '', 0, 0, '', '', '', 0
FROM (SELECT funcionario AS id FROM robo.memoria_treino_tentativas
      UNION SELECT de FROM robo.memoria_interacoes UNION SELECT NULLIF(para, '') FROM robo.memoria_interacoes
      UNION SELECT mentor FROM robo.memoria_mentorias UNION SELECT aprendiz FROM robo.memoria_mentorias
      UNION SELECT funcionario FROM robo.memoria_rh_movimentos UNION SELECT funcionario FROM robo.memoria_treino_niveis_eventos
      UNION SELECT funcionario FROM robo.memoria_treino_niveis) x
WHERE x.id IS NOT NULL AND x.id NOT IN (SELECT funcionario_id FROM bi.dim_funcionario_base);

CREATE OR REPLACE VIEW bi.dim_competencia AS
SELECT competencia_id, competencia, ferramenta, designer_id, designer_nome, descricao FROM bi.dim_competencia_base
UNION ALL
SELECT x.id, x.id, '', '', '', '(fora do catálogo atual)'
FROM (SELECT competencia AS id FROM robo.memoria_treino_tentativas UNION SELECT competencia FROM robo.memoria_mentorias
      UNION SELECT competencia FROM robo.memoria_treino_niveis_eventos UNION SELECT competencia FROM robo.memoria_treino_niveis) x
WHERE x.id IS NOT NULL AND x.id NOT IN (SELECT competencia_id FROM bi.dim_competencia_base);

CREATE OR REPLACE VIEW bi.dim_data AS
WITH limites AS (
  SELECT MIN(d) AS ini, MAX(d) AS fim FROM (
    SELECT data FROM bi.fato_treino UNION ALL SELECT data FROM bi.fato_interacoes UNION ALL SELECT data FROM bi.fato_mentorias
    UNION ALL SELECT data FROM bi.fato_rh_movimentos UNION ALL SELECT data FROM bi.fato_rotinas
    UNION ALL SELECT data FROM bi.fato_evolucao_nivel) t(d))
SELECT g::date AS data, EXTRACT(YEAR FROM g)::int AS ano, EXTRACT(MONTH FROM g)::int AS mes,
       (ARRAY['jan','fev','mar','abr','mai','jun','jul','ago','set','out','nov','dez'])[EXTRACT(MONTH FROM g)::int] AS mes_nome,
       to_char(g, 'YYYY-MM') AS ano_mes, EXTRACT(DAY FROM g)::int AS dia, EXTRACT(ISODOW FROM g)::int AS dia_semana,
       (ARRAY['seg','ter','qua','qui','sex','sáb','dom'])[EXTRACT(ISODOW FROM g)::int] AS dia_semana_nome,
       EXTRACT(WEEK FROM g)::int AS semana_iso, CASE WHEN EXTRACT(ISODOW FROM g) >= 6 THEN 1 ELSE 0 END AS fim_de_semana
FROM limites, generate_series(limites.ini, limites.fim, interval '1 day') g;

-- ============================================================ quem lê (o app do celular e o Power BI usam robo_leitura)
GRANT USAGE ON SCHEMA bi TO robo_leitura;
GRANT SELECT ON ALL TABLES IN SCHEMA bi TO robo_leitura;
ALTER DEFAULT PRIVILEGES IN SCHEMA bi GRANT SELECT ON TABLES TO robo_leitura;
"""


def gerar_sql(dims: dict) -> str:
    partes = ["""-- =====================================================================================
-- MODELO ESTRELA DA EMPRESA NO POSTGRESQL (schema bi) — gerado por powerbi/exportar_estrela.py, não edite à mão.
-- Rode no SQL Editor do Neon (ou psql) como o dono do banco. Pode rodar de novo: recria tudo sem perder nada
-- (as views só LEEM o schema robo; as tabelas *_base guardam o organograma e o catálogo, que vivem no código).
-- Os turnos da nuvem (robo/nuvem.py) só escrevem no schema robo — o schema bi nunca é tocado por eles.
-- =====================================================================================
CREATE SCHEMA IF NOT EXISTS bi;
DROP VIEW IF EXISTS bi.dim_data, bi.dim_funcionario, bi.dim_competencia, bi.fato_treino, bi.fato_interacoes,
  bi.fato_mentorias, bi.fato_rh_movimentos, bi.fato_rotinas, bi.fato_evolucao_nivel, bi.fato_nivel_atual CASCADE;
DROP TABLE IF EXISTS bi.dim_funcionario_base, bi.dim_competencia_base, bi.dim_andar, bi.dim_nivel;
CREATE TABLE bi.dim_funcionario_base (funcionario_id TEXT PRIMARY KEY, nome TEXT, cargo TEXT, nivel_cargo TEXT,
  nivel_cargo_ordem INT, andar INT, area TEXT, gestor_id TEXT, gestor_nome TEXT, no_quadro INT);
CREATE TABLE bi.dim_competencia_base (competencia_id TEXT PRIMARY KEY, competencia TEXT, ferramenta TEXT,
  designer_id TEXT, designer_nome TEXT, descricao TEXT);
CREATE TABLE bi.dim_andar (andar INT PRIMARY KEY, andar_nome TEXT, andar_rotulo TEXT);
CREATE TABLE bi.dim_nivel (nivel INT PRIMARY KEY, nivel_nome TEXT);"""]
    for tabela, cols in (("dim_funcionario_base", ["funcionario_id", "nome", "cargo", "nivel_cargo", "nivel_cargo_ordem", "andar",
                                                   "area", "gestor_id", "gestor_nome", "no_quadro"]),
                         ("dim_competencia_base", ["competencia_id", "competencia", "ferramenta", "designer_id", "designer_nome",
                                                   "descricao"]),
                         ("dim_andar", ["andar", "andar_nome", "andar_rotulo"]), ("dim_nivel", ["nivel", "nivel_nome"])):
        valores = ",\n  ".join("(" + ", ".join(_sql_lit(r[c]) for c in cols) + ")" for r in dims[tabela])
        partes.append(f"INSERT INTO bi.{tabela} ({', '.join(cols)}) VALUES\n  {valores};")
    partes.append(VIEWS_SQL.strip())
    return "\n\n".join(partes) + "\n"


def exportar(cfg, caminho_db: Path, destino: Path) -> dict:
    """Usado pela rotina dos painéis (robo/paineis.py): atualiza o modelo estrela em `destino` a cada publicação."""
    dims = dimensoes_do_codigo(cfg)
    c = sqlite3.connect(f"file:{caminho_db}?mode=ro", uri=True)
    c.row_factory = sqlite3.Row
    try:
        fat = fatos(c)
    finally:
        c.close()
    tabelas = {**completar_dimensoes(dims, fat), **fat}
    gravar_csv_xlsx(tabelas, destino)
    return {n: len(v) for n, v in tabelas.items()}


def main(caminho_db: Path) -> dict:
    from robo.config import carregar_config
    base = caminho_db.resolve().parent.parent
    cfg = carregar_config(base / "nao_existe.yaml",
                          sobrescrever={"pastas": {k: str(base / k) for k in ("entrada", "logs", "agentes", "habilidades")}})
    dims = dimensoes_do_codigo(cfg)
    c = sqlite3.connect(f"file:{caminho_db}?mode=ro", uri=True)
    c.row_factory = sqlite3.Row
    try:
        fat = fatos(c)
    finally:
        c.close()
    tabelas = {**completar_dimensoes(dims, fat), **fat}
    gravar_csv_xlsx(tabelas)
    (PASTA / "estrela_neon.sql").write_text(gerar_sql(dims), encoding="utf-8")
    return {n: len(v) for n, v in tabelas.items()}


if __name__ == "__main__":
    db = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "logs" / "trabalho.db"
    for nome, n in main(db).items():
        print(f"{nome:22s} {n:6d} linhas")
