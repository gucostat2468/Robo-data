import os
import time

import pandas as pd
from openpyxl import Workbook

from conftest import REGRA_NPS
from robo.monitor import Monitor
from robo.pipeline import REPROVADO, SUCESSO, processar_arquivo


def test_fluxo_completo_sucesso(planilha_baguncada, cfg_factory):
    cfg = cfg_factory([REGRA_NPS])
    r = processar_arquivo(planilha_baguncada, cfg)
    assert r.status == SUCESSO, r.relatorio.texto()

    saida = cfg.tratadas / "nps_outubro.xlsx"
    assert saida.exists()
    assert not planilha_baguncada.exists()                       # saiu da entrada
    assert r.destino_original.exists()                            # backup guardado
    assert (cfg.relatorios / "nps_outubro_qualidade.xlsx").exists()
    assert (cfg.logs / "historico_processamentos.csv").exists()

    df = pd.read_excel(saida, sheet_name="dados")
    assert len(df) == 4                                           # sem duplicada e sem linha Total
    assert df["uf"].tolist() == ["SP", "RJ", "MG", "SP"]
    assert df["nota_nps"].tolist() == [10, 9, 6, 0]
    assert df["curso"].tolist()[1] == "Direito"
    assert df["valor_mensalidade"].iloc[0] == 1234.56
    assert pd.isna(df["valor_mensalidade"].iloc[3])
    assert df["cep"].iloc[0] == "01310-100"
    assert "obs" not in df.columns                                # coluna vazia removida
    assert "_arquivo_origem" in df.columns
    assert pd.api.types.is_datetime64_any_dtype(df["data_resposta"])


def test_reprovado_vai_para_erros_com_relatorio(pastas, cfg_factory):
    wb = Workbook()
    ws = wb.active
    ws.append(["ID Resposta", "Nota NPS", "UF"])
    ws.append([1, 11, "SP"])        # fora da faixa
    ws.append([1, 5, "XX"])         # id repetido e UF inválida
    ws.append([2, "#DIV/0!", "RJ"])
    caminho = pastas / "Bases-Limpeza" / "nps_ruim.xlsx"
    wb.save(caminho)

    cfg = cfg_factory([{k: v for k, v in REGRA_NPS.items() if k != "tipos"}])
    r = processar_arquivo(caminho, cfg)
    assert r.status == REPROVADO
    assert not (cfg.tratadas / "nps_ruim.xlsx").exists()
    texto = r.saidas[0].read_text(encoding="utf-8-sig")
    assert "REPROVADO" in texto
    assert "fora da faixa" in texto
    assert "Chave única repetida" in texto
    assert "erro de fórmula" in texto
    assert "linha 2" in texto
    resumo = pd.read_csv(cfg.relatorios / "resumo_ocorrencias.csv", sep=";", encoding="utf-8-sig")
    assert set(resumo["nivel"]) == {"ERRO"} and (resumo["arquivo"] == "nps_ruim.xlsx").all()
    assert {"Chave única repetida (id_resposta)", "Valores fora da lista permitida"} <= set(resumo["mensagem"])


def test_modelo_estrela_e_calendario(planilha_baguncada, cfg_factory):
    regra = dict(REGRA_NPS, modelo={
        "fato": "fato_nps",
        "dimensoes": {"dim_curso": {"colunas": ["curso", "modalidade"], "chave": "id_curso"}},
        "calendario": "data_resposta",
    })
    cfg = cfg_factory([regra])
    r = processar_arquivo(planilha_baguncada, cfg)
    assert r.status == SUCESSO, r.relatorio.texto()
    abas = pd.read_excel(cfg.tratadas / "nps_outubro.xlsx", sheet_name=None)
    assert set(abas) == {"fato_nps", "dim_curso", "dim_calendario"}
    fato, dim = abas["fato_nps"], abas["dim_curso"]
    assert "curso" not in fato.columns and "id_curso" in fato.columns
    assert len(dim) == 2
    assert fato["id_curso"].notna().all()
    junto = fato.merge(dim, on="id_curso")
    assert sorted(junto["curso"].unique()) == ["Direito", "Medicina"]
    assert len(abas["dim_calendario"]) == 365


def test_diagnostico_nao_mexe_em_nada(planilha_baguncada, cfg_factory):
    cfg = cfg_factory([REGRA_NPS])
    r = processar_arquivo(planilha_baguncada, cfg, gravar=False)
    assert planilha_baguncada.exists()
    assert not (cfg.tratadas / "nps_outubro.xlsx").exists()
    assert r.tabelas["dados"].shape[0] == 4


def test_formato_nao_suportado(pastas, cfg_factory):
    caminho = pastas / "Bases-Limpeza" / "apresentacao.pdf"
    caminho.write_bytes(b"%PDF-1.4")
    r = processar_arquivo(caminho, cfg_factory())
    assert r.status == REPROVADO


def test_monitor_espera_estabilidade(planilha_baguncada, cfg_factory):
    cfg = cfg_factory([REGRA_NPS])
    m = Monitor(cfg)
    assert m.ciclo(agora=100.0) == []                 # 1ª vez: só registra
    assert m.ciclo(agora=105.0) == []                 # ainda dentro da espera (10s)
    resultados = m.ciclo(agora=111.0)
    assert [r.status for r in resultados] == [SUCESSO]
    assert (cfg.tratadas / "nps_outubro.xlsx").exists()


def test_monitor_ignora_temporarios(pastas, cfg_factory):
    (pastas / "Bases-Limpeza" / "~$nps.xlsx").write_bytes(b"lock")
    (pastas / "Bases-Limpeza" / "desktop.ini").write_text("x")
    m = Monitor(cfg_factory())
    m.ciclo(agora=0.0)
    assert m.ciclo(agora=100.0) == []


def test_csv_saida(planilha_baguncada, cfg_factory):
    cfg = cfg_factory([REGRA_NPS], saida={"formato": "csv"})
    r = processar_arquivo(planilha_baguncada, cfg)
    assert r.status == SUCESSO
    df = pd.read_csv(cfg.tratadas / "nps_outubro.csv", sep=";", decimal=",", encoding="utf-8-sig")
    assert df["valor_mensalidade"].iloc[0] == 1234.56


def test_etapas_emitidas_em_ordem(planilha_baguncada, cfg_factory):
    etapas = []
    r = processar_arquivo(planilha_baguncada, cfg_factory([REGRA_NPS]), ao_etapa=lambda e, a, **k: etapas.append((e, k)))
    assert r.status == SUCESSO
    assert [e for e, _ in etapas] == ["leitura", "limpeza", "validacao", "modelagem", "supervisor", "coord_engenharia",
                                      "publicacao", "analise:perfil", "analise:indicadores", "analise:estatistica",
                                      "analise:monitoramento", "analise:relator", "coord_analise",
                                      "excel_pesquisa", "excel_formulas", "excel_trade", "excel_dinamicas", "excel_supervisor",
                                      "concluido"]
    assert etapas[-1][1]["status"] == SUCESSO and etapas[-1][1]["linhas"] == 4


def test_texto_com_igual_nao_vira_formula_ativa(pastas, cfg_factory):
    """Segurança: célula de origem '=HYPERLINK(...)' precisa sair como texto, nunca como fórmula."""
    from openpyxl import load_workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Nome", "Obs"])
    ws.append(["Ana", '=HYPERLINK("http://malicioso","clique")'])
    ws["B2"].data_type = "s"  # na origem é texto
    ws.append(["Bia", "ok"])
    caminho = pastas / "Bases-Limpeza" / "seg.xlsx"
    wb.save(caminho)
    r = processar_arquivo(caminho, cfg_factory())
    assert r.status == SUCESSO
    celula = load_workbook(r.saidas[0])["dados"]["B2"]
    assert celula.data_type == "s" and celula.value.startswith("=HYPERLINK")
