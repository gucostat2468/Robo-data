"""Sala Excel: as fórmulas são recalculadas pelo LibreOffice e TODAS as conferências precisam fechar."""

import shutil
import subprocess

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from robo.analise import detectar_indicadores, papel_coluna
from robo.excel import Entrega, _criterio
from robo.pipeline import SUCESSO, processar_arquivo

SOFFICE = shutil.which("soffice")


def _papeis(df):
    return {c: papel_coluna(df[c], c, 30, ["id", "id_*", "*_id"]) for c in df.columns}


def _recalcular(caminho, pasta):
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(pasta), str(caminho)],
                   check=True, capture_output=True, timeout=120)
    return load_workbook(pasta / caminho.name, data_only=True)


def test_criterio_escapa_curingas():
    assert _criterio("Pós*Graduação ?") == '"Pós~*Graduação ~?"'
    assert _criterio(">10") == '"=>10"'
    assert _criterio('Diz "oi"') == '"Diz ""oi"""'


@pytest.mark.skipif(not SOFFICE, reason="LibreOffice não instalado")
@pytest.mark.parametrize("cenario", ["pesquisa", "trade"])
def test_formulas_recalculadas_batem_com_a_conferencia(tmp_path, cenario):
    rng = np.random.default_rng(4)
    if cenario == "pesquisa":
        df = pd.DataFrame({"curso": rng.choice(["Direito", "Medicina", "Pós*Graduação ?", "=Especial"], 250),
                           "uf": rng.choice(["SP", "RJ"], 250), "nota_nps": pd.array(rng.integers(0, 11, 250), dtype="Int64"),
                           "csat": pd.array(rng.integers(1, 6, 250), dtype="Int64")})
    else:
        df = pd.DataFrame({"loja": rng.choice(["Loja A", "Loja B", "Loja C"], 300), "canal": rng.choice(["Atacado", "Varejo"], 300),
                           "sku": rng.choice([f"SKU-{i}" for i in range(15)], 300),
                           "valor_venda": rng.uniform(5, 900, 300).round(2), "quantidade": rng.integers(1, 30, 300)})
    papeis = _papeis(df)
    destino = tmp_path / "entrega.xlsx"
    info = Entrega(df, "base", "base.xlsx", papeis, detectar_indicadores(df, papeis, {}), {}).gerar(destino)
    assert info["conferencias"] > 20
    wb = _recalcular(destino, tmp_path / "calc")
    status = wb["Resumo"]["B3"].value
    divergencias = [(r[1].value, r[2].value, r[3].value) for r in wb["Conferência"].iter_rows(min_row=5) if r[4].value != "OK"]
    assert not divergencias, divergencias
    assert status.startswith("✔")
    if cenario == "trade":
        assert "Trade" in wb.sheetnames
    else:
        assert {"Plano tabular", "Indicadores"} <= set(wb.sheetnames)


def test_pipeline_entrega_excel(pastas, cfg_factory):
    import sys, os
    sys.path.insert(0, os.path.dirname(__file__))
    from test_analise import _salvar_xlsx, base_nps
    cfg = cfg_factory()
    arq = pastas / "Bases-Limpeza" / "nps_set.xlsx"
    _salvar_xlsx(base_nps(n=120), arq)
    etapas = []
    r = processar_arquivo(arq, cfg, ao_etapa=lambda e, a, **k: etapas.append(e))
    assert r.status == SUCESSO and r.excel and r.excel["conferencias"] > 10
    assert (cfg.relatorios / "nps_excel.xlsx").exists()
    assert "excel_supervisor" in etapas
